"""
Set out the results of T32 (13 Oct 2026) in one file: for each reading of
READINGS_places.md, the chat and reading it came from, the PI's worth and note,
and the two markers' rates; then by chat; then by the point each reading opens,
as compare_45_46_27.md pairs them (written before any marker ran or any
worth was rated), with the wager test's plain chat on Bechky (chat 27) beside
them.

    python3 -B places_results.py <out_dir>

Reads places_packs_key.json, worth_sheet_places_filled.xlsx (the PI's sheet of 13 Oct),
places_marker_results.json, and compare_45_46_27.md, only to check by its hash that it
is the file as first written. Writes
<out_dir>/PLACES_RESULTS.md. Nothing in it changes a code, a mark or a rating.
"""
import collections
import hashlib
import json
import os
import sys

from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
COMPARE_SHA = "d5d96ea11e6c6a1dfdc79fe6f050d4085805f740dce80025fa0ee25f49ab6bea"
CHATS = [("45", "places given (chat 45)"), ("46", "places found (chat 46)")]
# The points, as compare_45_46_27.md pairs the readings: (point, chat 45's, chat 46's, chat 27's)
POINTS = [
    ("toxicology's change was a step its accreditation plan already required, brought forward; "
     "the coming challenge set its date, not its content", "W04", "W03", "R1"),
    ("firearms had already changed, from recognition on sight to a record that others can check, "
     "with no local challenge behind it", "W08", "W05", "R3, as one of three units"),
    ("narcotics kept crystal tests because a legal clock and two instruments left it no other way, "
     "more than for the values the paper names", "W07", "W01", "R2, as one of three units"),
    ("toxicology answered the comparison by presenting its existing work as continuous: the report "
     "changed, the bench need not have", "W02", None, None),
    ("the units' values are partly the residue of sorting analysts with the same degrees among the "
     "benches of one lab", None, "W06", "related, not the same: R5"),
]


def main():
    if len(sys.argv) != 2:
        raise SystemExit("usage: python3 -B places_results.py <out_dir>")
    out_dir = sys.argv[1]
    if hashlib.sha256(open(os.path.join(HERE, "compare_45_46_27.md"), "rb").read()).hexdigest() != COMPARE_SHA:
        raise SystemExit("compare_45_46_27.md is not the file as first written")
    key = json.load(open(os.path.join(HERE, "places_packs_key.json"), encoding="utf-8"))
    marks = json.load(open(os.path.join(HERE, "places_marker_results.json"), encoding="utf-8"))
    ws = load_workbook(os.path.join(HERE, "worth_sheet_places_filled.xlsx"))["Bechky"]
    worth, note = {}, {}
    for r in range(2, ws.max_row + 1):
        rid = ws[f"A{r}"].value
        if rid:
            worth[rid] = ws[f"C{r}"].value
            note[rid] = (ws[f"D{r}"].value or "").strip()
    entries = key["readings"]
    if sorted(worth) != sorted(e["reading"] for e in entries) or not all(
            isinstance(v, int) and 0 <= v <= 3 for v in worth.values()):
        raise SystemExit("the sheet does not rate every reading 0 to 3")
    by = marks["rates"]["by_reading"]

    def res(rid, m):
        x = by[rid]["rates"][m]
        return round(100 * (x["R"] + x["RI"] + x["RA"]) / x["reporting"]) if x["reporting"] else None

    out = ["# T32 on Bechky: the readings, their marks and their worth", "",
           "Written by `places_results.py` from `places_packs_key.json`, the PI's",
           "`worth_sheet_places_filled.xlsx` (13 Oct), `places_marker_results.json` and",
           "`compare_45_46_27.md` (checked by its hash to be the file as first written). Worth is the PI's,",
           "0 to 3 on the wager test's scale; resolution is the share of a reading's sentences that",
           "report something which the marker found borne out by the paper (R, RI or RA), Opus 5.5 /",
           "GPT 6 Astra; the markers agree with kappa 0.669, above the bar of 0.60.", "",
           "## By reading", "",
           "| reading | chat, reading | worth | the PI's note | resolution | U or C as warrant |",
           "|---|---|---|---|---|---|"]
    for e in entries:
        rid = e["reading"]
        x = by[rid]["rates"]
        out.append(f"| {rid} | {e['chat']}, {e['source_id']} | {worth[rid]} | {note[rid]} | "
                   f"{res(rid, 'opus')}% / {res(rid, 'astra')}% | "
                   f"{'yes' if x['opus']['warrant_U_or_C'] else 'no'} / {'yes' if x['astra']['warrant_U_or_C'] else 'no'} |")
    out += ["", "## By chat", "",
            "| | " + " | ".join(label for _, label in CHATS) + " |", "|---|---|---|"]
    ws_by = {c: [worth[e["reading"]] for e in entries if e["chat"] == c] for c, _ in CHATS}
    out.append("| readings | " + " | ".join(str(len(ws_by[c])) for c, _ in CHATS) + " |")
    out.append("| worth, mean | " + " | ".join(f"{sum(ws_by[c]) / len(ws_by[c]):.2f}" for c, _ in CHATS) + " |")
    out.append("| worth 3, 2, 1, 0 | " + " | ".join(
        ", ".join(str(collections.Counter(ws_by[c])[v]) for v in (3, 2, 1, 0)) for c, _ in CHATS) + " |")
    arm = marks["rates"]["by_arm"]
    names = {"45": "places given", "46": "places found"}

    def arm_res(c, m):
        x = arm[m][names[c]]
        return round(100 * (x["R"] + x["RI"] + x["RA"]) / x["reporting"])
    out.append("| resolution | " + " | ".join(f"{arm_res(c, 'opus')}% / {arm_res(c, 'astra')}%" for c, _ in CHATS) + " |")
    out += ["", "## By point", "",
            "The points as `compare_45_46_27.md` pairs the readings. Chat 27 is the wager test's plain",
            "chat on Bechky (11 Oct): the paper and a plain prompt, no places; the PI rated each of its",
            "five readings 2 (`stage2/worth/WORTH_RESULTS.md`).", "",
            "| the point the reading opens | at a given place (chat 45) | at a found place (chat 46) | chat 27, 11 Oct |",
            "|---|---|---|---|"]
    for point, a, b, c27 in POINTS:
        ca = f"{a}: worth {worth[a]}" if a else "none"
        cb = f"{b}: worth {worth[b]}" if b else "none"
        out.append(f"| {point} | {ca} | {cb} | {c27 + ', worth 2' if c27 else 'none'} |")
    out.append("")
    os.makedirs(out_dir, exist_ok=True)
    with open(os.path.join(out_dir, "PLACES_RESULTS.md"), "w", encoding="utf-8") as fh:
        fh.write("\n".join(out))
    print("wrote PLACES_RESULTS.md;", {c: sum(v) / len(v) for c, v in ws_by.items()})


if __name__ == "__main__":
    main()
