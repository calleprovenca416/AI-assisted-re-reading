"""
Score the matching test (T33), by MATCHING_TEST_PLAN.md. Written on 14 Oct 2026,
before any reply of the test was opened.

    python3 -B score_matching.py selftest
    python3 -B score_matching.py check --arm=A|B --key=matching_key.json REPLY
    python3 -B score_matching.py score --key=matching_key.json --blank=grouping_sheet.xlsx
        --sheet=grouping_sheet_filled.xlsx --a1=REPLY_51 --a2=REPLY_52 --b1=REPLY_53 --b2=REPLY_54
        --out=DIR

check: whether one coder's reply counts, as the plan's "Checks on the replies"
says: one JSON object (from its first "{" to its last "}") with a list
`readings` naming every reading of the key exactly once, and a list `groups`
putting every reading in exactly one group; in arm B, every `claim` one of C01
to C23 or `none`, every item of `other_claims` (at most three) one of C01 to
C23, and every item of `passages` (at most five) one of F01 to F172. Prints the
faults, if any, and exits 1 when there are some.

score: every figure of the plan, written to DIR/MATCHING_RESULTS.md and
DIR/matching_results.json. The PI's sheet counts if its readings and their
words are those of the blank sheet, in the same order, no other cell outside
the yellow ones has changed, and every reading has a group label. A label is
read as text, its ends trimmed, and labels that differ only in capitals are the
same (the sheet says so). A coder reply that fails its check is named, and every
figure that needs it is given as not decided.

The groupings: A1, A2 (chats 51, 52: their groups); B1, B2 (chats 53, 54, by
rule: the same claim, the same group; each `none` a group of its own); B1j, B2j
(chats 53, 54: their own groups); B1p, B2p (by rule, finer: within a claim's
group, readings joined when they name a passage in common, and the groups those
joins form); P (the PI's sheet).

The pair kappa of two groupings: over every pair of distinct readings, 1 if a
grouping puts the pair together and 0 if apart; Cohen's kappa on the two lists.
It is not computed (None) when either list is all 0 or all 1, as the plan says
of a grouping with no pair together or every pair together. The adjusted Rand
index is given beside it. The bootstrap: 2,000 resamples of the readings with
the seed 20261014 (random.Random, one draw of n indices per resample), pairs
formed among the drawn readings that are copies of different readings, each
kappa computed on the resample's pairs; a difference is left out of a resample
in which one of its kappas is None, and the count left out is given. The 2.5th
and 97.5th percentiles are read from the sorted values by linear interpolation
(position q(m - 1)). The cases of Q1 and the answer of Q2 are decided as the
plan fixes them; comparisons with the bars are made on values rounded to six
places.
"""
import hashlib
import itertools
import json
import os
import random
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
CLAIMS = [f"C{i:02d}" for i in range(1, 24)]
PASSAGES = [f"F{i:02d}" if i < 100 else f"F{i}" for i in range(1, 173)]
EARLIER = ("23", "26", "27", "46")
NEW = ("49", "50")
SEED = 20261014
RESAMPLES = 2000


# ---------------------------------------------------------------- reading files

def load_reply(path):
    text = open(path, encoding="utf-8").read()
    a, b = text.find("{"), text.rfind("}")
    if a < 0 or b < a:
        return None, ["no JSON object found"]
    try:
        return json.loads(text[a:b + 1]), []
    except json.JSONDecodeError as e:
        return None, [f"not one JSON object: {e}"]


def check_reply(rep, ids, arm):
    """Faults of one coder's reply; and its groups, claims and passages if it counts."""
    faults = []
    if not isinstance(rep, dict):
        return ["not a JSON object"], None
    rs, gs = rep.get("readings"), rep.get("groups")
    if not isinstance(rs, list):
        faults.append("`readings` is not a list")
        rs = []
    if not isinstance(gs, list):
        faults.append("`groups` is not a list")
        gs = []
    named = [r.get("reading") if isinstance(r, dict) else None for r in rs]
    for i in ids:
        if named.count(i) != 1:
            faults.append(f"`readings` names {i} {named.count(i)} times")
    extra = sorted({str(x) for x in named if x not in ids})
    if extra:
        faults.append(f"`readings` names readings not in the file: {extra}")
    grouped = []
    for g in gs:
        if not isinstance(g, dict) or not isinstance(g.get("readings"), list):
            faults.append("a group has no list of readings")
            continue
        grouped += g["readings"]
    for i in ids:
        if grouped.count(i) != 1:
            faults.append(f"`groups` holds {i} {grouped.count(i)} times")
    extra = sorted({str(x) for x in grouped if x not in ids})
    if extra:
        faults.append(f"`groups` holds readings not in the file: {extra}")
    claims, passages = {}, {}
    if arm == "B":
        for r in rs:
            if not isinstance(r, dict) or r.get("reading") not in ids:
                continue
            c = r.get("claim")
            if c != "none" and c not in CLAIMS:
                faults.append(f"{r['reading']}: claim {c!r} is not C01 to C23 or none")
            oc = r.get("other_claims", [])
            if not isinstance(oc, list) or len(oc) > 3 or any(x not in CLAIMS for x in oc):
                faults.append(f"{r['reading']}: other_claims {oc!r} is not a list of at most three of C01 to C23")
            ps = r.get("passages", [])
            if not isinstance(ps, list) or len(ps) > 5 or any(x not in PASSAGES for x in ps):
                faults.append(f"{r['reading']}: passages {ps!r} is not a list of at most five of F01 to F172")
            claims[r["reading"]] = c
            passages[r["reading"]] = list(ps) if isinstance(ps, list) else []
    if faults:
        return faults, None
    group_of = {}
    for k, g in enumerate(gs):
        for i in g["readings"]:
            group_of[i] = k
    return [], {"groups": group_of, "claims": claims, "passages": passages}


def read_sheet(filled, blank, ids):
    import openpyxl
    faults = []
    wf, wbk = openpyxl.load_workbook(filled), openpyxl.load_workbook(blank)
    if wf.sheetnames != wbk.sheetnames:
        return [f"sheets {wf.sheetnames} are not {wbk.sheetnames}"], None
    for name in wbk.sheetnames:
        a, b = wf[name], wbk[name]
        rows = max(a.max_row, b.max_row)
        cols = max(a.max_column, b.max_column)
        for r in range(1, rows + 1):
            for c in range(1, cols + 1):
                if name == "Bechky" and r >= 2 and c in (3, 4):
                    continue
                if a.cell(r, c).value != b.cell(r, c).value:
                    faults.append(f"{name}!{a.cell(r, c).coordinate} differs from the blank sheet")
    ws = wf["Bechky"]
    labels = {}
    for r in range(2, ws.max_row + 1):
        rid = ws.cell(r, 1).value
        if rid is None:
            continue
        v = ws.cell(r, 3).value
        lab = "" if v is None else str(v).strip()
        if not lab:
            faults.append(f"{rid} has no group label")
        labels[rid] = lab.casefold()
    if sorted(labels) != sorted(ids):
        faults.append("the sheet's readings are not the key's")
    if faults:
        return faults, None
    return [], labels


# ---------------------------------------------------------------- groupings and agreement

def rule_groups(claims):
    return {i: (c if c != "none" else f"none:{i}") for i, c in claims.items()}


def finer_groups(claims, passages):
    parent = {i: i for i in claims}

    def find(x):
        while parent[x] != x:
            parent[x] = parent[parent[x]]
            x = parent[x]
        return x
    ids = sorted(claims)
    for a, b in itertools.combinations(ids, 2):
        if claims[a] != "none" and claims[a] == claims[b] and set(passages[a]) & set(passages[b]):
            parent[find(a)] = find(b)
    return {i: find(i) for i in ids}


def pair_labels(group_of, pairs):
    return [1 if group_of[a] == group_of[b] else 0 for a, b in pairs]


def kappa(x, y):
    n = len(x)
    if n == 0 or n != len(y):
        return None
    p1, p2 = sum(x) / n, sum(y) / n
    if p1 in (0.0, 1.0) or p2 in (0.0, 1.0):
        return None
    po = sum(1 for a, b in zip(x, y) if a == b) / n
    pe = p1 * p2 + (1 - p1) * (1 - p2)
    return (po - pe) / (1 - pe)


def nominal_kappa(x, y):
    n = len(x)
    if n == 0:
        return None
    cats = sorted(set(x) | set(y))
    po = sum(1 for a, b in zip(x, y) if a == b) / n
    pe = sum((x.count(c) / n) * (y.count(c) / n) for c in cats)
    if pe >= 1.0:
        return None
    return (po - pe) / (1 - pe)


def ari(g1, g2, ids):
    def c2(k):
        return k * (k - 1) / 2
    cells, rows, cols = {}, {}, {}
    for i in ids:
        cells[(g1[i], g2[i])] = cells.get((g1[i], g2[i]), 0) + 1
        rows[g1[i]] = rows.get(g1[i], 0) + 1
        cols[g2[i]] = cols.get(g2[i], 0) + 1
    index = sum(c2(v) for v in cells.values())
    a, b, total = sum(c2(v) for v in rows.values()), sum(c2(v) for v in cols.values()), c2(len(ids))
    expected = a * b / total if total else 0.0
    top = (a + b) / 2
    if top == expected:
        return None
    return (index - expected) / (top - expected)


def percentile(sorted_vals, q):
    if not sorted_vals:
        return None
    pos = q * (len(sorted_vals) - 1)
    lo = int(pos)
    hi = min(lo + 1, len(sorted_vals) - 1)
    return sorted_vals[lo] + (sorted_vals[hi] - sorted_vals[lo]) * (pos - lo)


def mean2(a, b):
    return None if a is None or b is None else (a + b) / 2


def figures(G, pairs):
    """The plan's Q1 figures on one set of pairs. G maps a grouping's name to its groups (or None)."""
    lab = {k: (pair_labels(v, pairs) if v is not None else None) for k, v in G.items()}

    def k(a, b):
        return None if lab.get(a) is None or lab.get(b) is None else kappa(lab[a], lab[b])
    f = {"A": k("A1", "A2"), "B": k("B1", "B2"), "Bj": k("B1j", "B2j"), "Bp": k("B1p", "B2p"),
         "P_A": mean2(k("A1", "P"), k("A2", "P")), "P_B": mean2(k("B1", "P"), k("B2", "P")),
         "P_Bj": mean2(k("B1j", "P"), k("B2j", "P")), "P_Bp": mean2(k("B1p", "P"), k("B2p", "P"))}
    return f


DIFFS = [("B - A", "B", "A"), ("Bj - A", "Bj", "A"), ("Bp - A", "Bp", "A"),
         ("P_B - P_A", "P_B", "P_A"), ("P_Bj - P_A", "P_Bj", "P_A"), ("P_Bp - P_A", "P_Bp", "P_A")]


def bootstrap(G, ids, seed=SEED, resamples=RESAMPLES):
    rng = random.Random(seed)
    n = len(ids)
    vals = {d[0]: [] for d in DIFFS}
    left = {d[0]: 0 for d in DIFFS}
    for _ in range(resamples):
        draw = [ids[rng.randrange(n)] for _ in range(n)]
        pairs = [(a, b) for a, b in itertools.combinations(draw, 2) if a != b]
        f = figures(G, pairs)
        for name, x, y in DIFFS:
            if f[x] is None or f[y] is None:
                left[name] += 1
            else:
                vals[name].append(f[x] - f[y])
    out = {}
    for name, _, _ in DIFFS:
        s = sorted(vals[name])
        out[name] = {"low": percentile(s, 0.025), "high": percentile(s, 0.975), "left_out": left[name],
                     "resamples": resamples}
    return out


def r6(x):
    return None if x is None else round(x, 6)


def q1_case(A, B, PA, PB, low):
    if None in (A, B, PA, PB, low):
        return "not decided", "a kappa or the interval could not be computed"
    if r6(A) >= 0.80:
        return "repeatable without the reference", "A is at least 0.80"
    if r6(B - A) >= 0.15 and r6(low) > 0:
        if r6(PB) >= r6(PA - 0.05):
            return "the reference makes the count repeatable", "B - A at least 0.15, its interval above 0, and P_B at least P_A - 0.05"
        return "repeatable, but coarser than the PI's grouping", "B - A at least 0.15, its interval above 0, but P_B below P_A - 0.05"
    return "no clear gain", "neither case 1 nor B - A at least 0.15 with its interval above 0"


def q2_answer(found_twice, reached):
    if found_twice == 0:
        return "not decided", "no group of P holds readings of two earlier chats"
    if reached * 2 >= found_twice:
        return "recur across vendors", f"{reached} of {found_twice} groups found twice hold a new reading"
    if reached >= 1:
        return "in part", f"{reached} of {found_twice} groups found twice hold a new reading"
    return "not", "no group found twice holds a new reading"


def q2(P, chat_of):
    groups = {}
    for i, g in P.items():
        groups.setdefault(g, []).append(i)
    twice = [g for g, m in groups.items() if len({chat_of[i] for i in m if chat_of[i] in EARLIER}) >= 2]
    reached = [g for g in twice if any(chat_of[i] in NEW for i in groups[g])]
    new = [i for i in P if chat_of[i] in NEW]
    with_earlier = [i for i in new if any(chat_of[j] in EARLIER for j in groups[P[i]])]
    g49 = {P[i] for i in P if chat_of[i] == "49"}
    g49_50 = [g for g in g49 if any(chat_of[j] == "50" for j in groups[g])]
    return {"groups": len(groups), "found_twice": len(twice), "found_twice_reached": len(reached),
            "new_readings": len(new), "new_with_earlier": len(with_earlier),
            "groups_of_49": len(g49), "groups_of_49_with_50": len(g49_50),
            "answer": q2_answer(len(twice), len(reached))}


# ---------------------------------------------------------------- the run

def score(opts):
    key = json.load(open(opts["key"], encoding="utf-8"))
    ids = [r["reading"] for r in key["readings"]]
    chat_of = {r["reading"]: r["chat"] for r in key["readings"]}
    files = {k: opts[k] for k in ("key", "blank", "sheet", "a1", "a2", "b1", "b2")}
    report = {"files": {k: {"file": os.path.basename(v), "sha256": hashlib.sha256(open(v, "rb").read()).hexdigest()}
                        for k, v in files.items()}, "checks": {}}
    parsed = {}
    for name, arm in (("a1", "A"), ("a2", "A"), ("b1", "B"), ("b2", "B")):
        rep, faults = load_reply(opts[name])
        got = None
        if not faults:
            faults, got = check_reply(rep, ids, arm)
        report["checks"][name] = faults
        parsed[name] = got
    sheet_faults, labels = read_sheet(opts["sheet"], opts["blank"], ids)
    report["checks"]["sheet"] = sheet_faults
    G = {"A1": parsed["a1"]["groups"] if parsed["a1"] else None,
         "A2": parsed["a2"]["groups"] if parsed["a2"] else None,
         "B1": rule_groups(parsed["b1"]["claims"]) if parsed["b1"] else None,
         "B2": rule_groups(parsed["b2"]["claims"]) if parsed["b2"] else None,
         "B1j": parsed["b1"]["groups"] if parsed["b1"] else None,
         "B2j": parsed["b2"]["groups"] if parsed["b2"] else None,
         "B1p": finer_groups(parsed["b1"]["claims"], parsed["b1"]["passages"]) if parsed["b1"] else None,
         "B2p": finer_groups(parsed["b2"]["claims"], parsed["b2"]["passages"]) if parsed["b2"] else None,
         "P": labels}
    pairs = list(itertools.combinations(sorted(ids), 2))
    f = figures(G, pairs)
    boot = bootstrap(G, sorted(ids))
    names = list(G)
    matrix = {}
    for a, b in itertools.combinations(names, 2):
        if G[a] is None or G[b] is None:
            matrix[f"{a} {b}"] = {"kappa": None, "ari": None}
        else:
            matrix[f"{a} {b}"] = {"kappa": kappa(pair_labels(G[a], pairs), pair_labels(G[b], pairs)),
                                  "ari": ari(G[a], G[b], sorted(ids))}
    case = q1_case(f["A"], f["B"], f["P_A"], f["P_B"], boot["B - A"]["low"])
    case_j = q1_case(f["A"], f["Bj"], f["P_A"], f["P_Bj"], boot["Bj - A"]["low"])
    neither = (f["P_A"] is not None and f["P_B"] is not None and r6(f["P_A"]) < 0.40 and r6(f["P_B"]) < 0.40)
    claim_k = None
    none = {}
    if parsed["b1"] and parsed["b2"]:
        c1 = [parsed["b1"]["claims"][i] for i in sorted(ids)]
        c2 = [parsed["b2"]["claims"][i] for i in sorted(ids)]
        claim_k = nominal_kappa(c1, c2)
        none = {"53": sum(c == "none" for c in c1), "54": sum(c == "none" for c in c2),
                "both": sum(a == "none" and b == "none" for a, b in zip(c1, c2)), "of": len(ids)}
    q2f = q2(labels, chat_of) if labels else None
    voices = {}
    if labels and parsed["b1"] and parsed["b2"]:
        side = json.load(open(os.path.join(HERE, "bechky_2020_v1_voice_gloss.json"), encoding="utf-8"))
        vcode = {v["fragment_id"]: v["code"] for v in side["voice"]}
        for i in sorted(ids):
            g = labels[i]
            v = voices.setdefault(g, {"readings": [], "reported": 0, "author_insight": 0})
            v["readings"].append(i)
            for p in parsed["b1"]["passages"][i] + parsed["b2"]["passages"][i]:
                v[vcode[p]] += 1
    report.update({"figures": f, "bootstrap": boot, "matrix": matrix,
                   "q1": {"case": case, "case_judged": case_j, "neither_finds_the_PIs_points": neither},
                   "claim_kappa": claim_k, "q3_none": none, "q2": q2f, "voices": voices,
                   "groupings": {k: v for k, v in G.items()},
                   "claims": {k: parsed[k]["claims"] for k in ("b1", "b2") if parsed[k]},
                   "chat_of": chat_of})
    os.makedirs(opts["out"], exist_ok=True)
    open(os.path.join(opts["out"], "matching_results.json"), "w", encoding="utf-8").write(
        json.dumps(report, indent=1, ensure_ascii=False, sort_keys=True) + "\n")
    open(os.path.join(opts["out"], "MATCHING_RESULTS.md"), "w", encoding="utf-8").write(write_md(report, ids))
    print("Q1:", case[0], "| judged:", case_j[0], "| Q2:", q2f["answer"][0] if q2f else None)


def fmt(x, d=3):
    return "not computed" if x is None else f"{x:+.{d}f}" if x < 0 else f"{x:.{d}f}"


def write_md(rep, ids):
    f, b, q2f = rep["figures"], rep["bootstrap"], rep["q2"]
    L = ["# The matching test (T33): results", "",
         "Written by `score_matching.py` from the files below, by `MATCHING_TEST_PLAN.md`. Kappas are pair "
         "kappas (Cohen's kappa on whether each pair of readings is put together); intervals are 95% bootstrap "
         "intervals of a difference, 2,000 resamples of the readings, seed 20261014.", "",
         "## The files", "", "| what | file | sha256 opens |", "|---|---|---|"]
    for k, v in rep["files"].items():
        L.append(f"| {k} | `{v['file']}` | {v['sha256'][:16]} |")
    L += ["", "## The checks", ""]
    for k, v in rep["checks"].items():
        L.append(f"- **{k}:** " + ("counts." if not v else "does not count: " + "; ".join(v)))
    L += ["", "## Q1: does the reference make the count repeatable?", "",
          "| figure | value |", "|---|---|",
          f"| A, chats 51 and 52 | {fmt(f['A'])} |", f"| B, chats 53 and 54 by rule | {fmt(f['B'])} |",
          f"| Bj, chats 53 and 54 judged | {fmt(f['Bj'])} |", f"| Bp, chats 53 and 54 by rule, finer | {fmt(f['Bp'])} |",
          f"| P_A, arm A with the PI | {fmt(f['P_A'])} |", f"| P_B, arm B by rule with the PI | {fmt(f['P_B'])} |",
          f"| P_Bj, arm B judged with the PI | {fmt(f['P_Bj'])} |", f"| P_Bp, arm B by rule, finer, with the PI | {fmt(f['P_Bp'])} |",
          "", "| difference | value | 95% interval | resamples left out |", "|---|---|---|---|"]
    for name, x, y in DIFFS:
        v = None if f[x] is None or f[y] is None else f[x] - f[y]
        L.append(f"| {name} | {fmt(v)} | {fmt(b[name]['low'])} to {fmt(b[name]['high'])} | {b[name]['left_out']} |")
    q1 = rep["q1"]
    L += ["", f"**The answer, by rule:** {q1['case'][0]} ({q1['case'][1]}).", "",
          f"**The second reading, judged:** {q1['case_judged'][0]} ({q1['case_judged'][1]}).", "",
          "**Neither arm finds the PI's points:** " + ("yes, P_A and P_B are both below 0.40." if q1["neither_finds_the_PIs_points"] else "no."),
          "", f"**Claim agreement in arm B** (Cohen's kappa on the claim numbers, `none` a code): {fmt(rep['claim_kappa'])}.", "",
          "## Every pair of groupings", "", "| groupings | pair kappa | adjusted Rand index |", "|---|---|---|"]
    for k, v in rep["matrix"].items():
        L.append(f"| {k} | {fmt(v['kappa'])} | {fmt(v['ari'])} |")
    L += ["", "## Q2: do the openings recur across vendors?", ""]
    if q2f:
        L += [f"- Groups in the PI's grouping: {q2f['groups']}.",
              f"- Groups found twice (readings of at least two of chats 23, 26, 27 and 46): {q2f['found_twice']}; "
              f"of them holding a reading of chat 49 or 50: {q2f['found_twice_reached']}.",
              f"- New readings sharing a group with an earlier one: {q2f['new_with_earlier']} of {q2f['new_readings']}.",
              f"- Groups holding a reading of chat 49: {q2f['groups_of_49']}; of them holding a reading of chat 50: "
              f"{q2f['groups_of_49_with_50']}.", "",
              f"**The answer:** {q2f['answer'][0]} ({q2f['answer'][1]})."]
    else:
        L.append("Not decided: the PI's sheet does not count.")
    L += ["", "## Q3: readings that fit no claim of the reference", ""]
    if rep["q3_none"]:
        n = rep["q3_none"]
        L.append(f"Chat 53 gives `none` to {n['53']} of {n['of']} readings, chat 54 to {n['54']}, both to {n['both']}.")
    else:
        L.append("Not decided: an arm B reply does not count.")
    L += ["", "## The groupings, reading by reading", "",
          "| reading | chat | P | A1 | A2 | B1 claim | B2 claim | B1j | B2j |", "|---|---|---|---|---|---|---|---|---|"]
    G = rep["groupings"]
    cl = rep["claims"]

    def g(name, i):
        return "—" if G.get(name) is None else str(G[name][i])
    for i in ids:
        L.append(f"| {i} | {rep['chat_of'][i]} | {g('P', i)} | {g('A1', i)} | {g('A2', i)} | "
                 f"{cl.get('b1', {}).get(i, '—')} | {cl.get('b2', {}).get(i, '—')} | {g('B1j', i)} | {g('B2j', i)} |")
    if rep["voices"]:
        L += ["", "## The voice of the passages arm B names, by the PI's groups (described, not tested)", "",
              "| PI's group | readings | reported | author's account |", "|---|---|---|---|"]
        for k, v in rep["voices"].items():
            L.append(f"| {k} | {', '.join(v['readings'])} | {v['reported']} | {v['author_insight']} |")
    L += ["", "The limits known in advance are those of `MATCHING_TEST_PLAN.md`.", ""]
    return "\n".join(L)


# ---------------------------------------------------------------- the self-test, on made-up groupings

def selftest():
    ok = True

    def eq(name, got, want, tol=1e-9):
        nonlocal ok
        good = (got is None and want is None) or (got is not None and want is not None and abs(got - want) <= tol)
        print(("pass" if good else "FAIL"), name, got, want)
        ok = ok and good

    ids = ["a", "b", "c", "d"]
    pairs = list(itertools.combinations(ids, 2))
    g1 = {"a": 1, "b": 1, "c": 2, "d": 3}
    g2 = {"a": 1, "b": 1, "c": 2, "d": 2}
    eq("pair kappa, worked by hand (8/14)", kappa(pair_labels(g1, pairs), pair_labels(g2, pairs)), 8 / 14)
    eq("adjusted Rand index, worked by hand (4/7)", ari(g1, g2, ids), 4 / 7)
    eq("pair kappa of a grouping with itself", kappa(pair_labels(g2, pairs), pair_labels(g2, pairs)), 1.0)
    single = {i: i for i in ids}
    eq("pair kappa with an all-apart grouping is not computed", kappa(pair_labels(single, pairs), pair_labels(g2, pairs)), None)
    eq("nominal kappa, two of four alike", nominal_kappa(["C1", "C1", "C2", "C3"], ["C1", "C2", "C2", "C3"]),
       (0.75 - (2 * 1 + 1 * 2 + 1 * 1) / 16) / (1 - (2 * 1 + 1 * 2 + 1 * 1) / 16))
    rg = rule_groups({"a": "C01", "b": "C01", "c": "none", "d": "none"})
    eq("by rule: same claim together", 1.0 if rg["a"] == rg["b"] else 0.0, 1.0)
    eq("by rule: each none apart", 1.0 if rg["c"] != rg["d"] else 0.0, 1.0)
    fg = finer_groups({"a": "C01", "b": "C01", "e": "C01", "c": "C02"},
                      {"a": ["F01"], "b": ["F02"], "e": ["F02", "F03"], "c": ["F01"]})
    eq("finer: shared passage joins", 1.0 if fg["b"] == fg["e"] else 0.0, 1.0)
    eq("finer: no shared passage, apart", 1.0 if fg["a"] != fg["b"] else 0.0, 1.0)
    eq("finer: a shared passage across claims does not join", 1.0 if fg["a"] != fg["c"] else 0.0, 1.0)
    eq("percentile, linear", percentile([0.0, 1.0, 2.0, 3.0], 0.5), 1.5)
    # check_reply
    ids3 = ["M01", "M02", "M03"]
    good_a = {"readings": [{"reading": i} for i in ids3], "groups": [{"readings": ["M01", "M02"]}, {"readings": ["M03"]}]}
    eq("a good arm A reply counts", float(len(check_reply(good_a, ids3, "A")[0])), 0.0)
    bad_a = {"readings": [{"reading": "M01"}, {"reading": "M01"}, {"reading": "M03"}],
             "groups": [{"readings": ["M01", "M02"]}, {"readings": ["M02", "M03"]}]}
    eq("a bad arm A reply has faults", 1.0 if check_reply(bad_a, ids3, "A")[0] else 0.0, 1.0)
    good_b = {"readings": [{"reading": "M01", "claim": "C17", "other_claims": ["C03"], "passages": ["F98", "F101"]},
                           {"reading": "M02", "claim": "C17", "other_claims": [], "passages": ["F98"]},
                           {"reading": "M03", "claim": "none", "other_claims": [], "passages": []}],
              "groups": [{"readings": ["M01", "M02"]}, {"readings": ["M03"]}]}
    eq("a good arm B reply counts", float(len(check_reply(good_b, ids3, "B")[0])), 0.0)
    bad_b = json.loads(json.dumps(good_b))
    bad_b["readings"][0]["claim"] = "C24"
    bad_b["readings"][1]["passages"] = ["F173"]
    eq("a bad arm B reply has two faults", float(len(check_reply(bad_b, ids3, "B")[0])), 2.0)
    # decisions
    eq("case 1 at A = 0.80", 1.0 if q1_case(0.80, 0.95, 0.5, 0.5, 0.01)[0] == "repeatable without the reference" else 0.0, 1.0)
    eq("case 2 at the bars", 1.0 if q1_case(0.50, 0.65, 0.60, 0.55, 0.001)[0] == "the reference makes the count repeatable" else 0.0, 1.0)
    eq("case 3 just under P_A - 0.05", 1.0 if q1_case(0.50, 0.65, 0.60, 0.549, 0.001)[0] == "repeatable, but coarser than the PI's grouping" else 0.0, 1.0)
    eq("case 4 when the interval reaches 0", 1.0 if q1_case(0.50, 0.70, 0.60, 0.60, 0.0)[0] == "no clear gain" else 0.0, 1.0)
    eq("case 4 when the gain is short", 1.0 if q1_case(0.50, 0.649, 0.60, 0.60, 0.1)[0] == "no clear gain" else 0.0, 1.0)
    eq("not decided when a kappa is missing", 1.0 if q1_case(None, 0.7, 0.6, 0.6, 0.1)[0] == "not decided" else 0.0, 1.0)
    # Q2
    chat_of = {"m1": "23", "m2": "27", "m3": "49", "m4": "26", "m5": "46", "m6": "50", "m7": "50"}
    P = {"m1": "x", "m2": "x", "m3": "x", "m4": "y", "m5": "y", "m6": "z", "m7": "w"}
    r = q2(P, chat_of)
    eq("Q2: groups found twice", float(r["found_twice"]), 2.0)
    eq("Q2: of them reached by a new reading", float(r["found_twice_reached"]), 1.0)
    eq("Q2: new readings with an earlier one", float(r["new_with_earlier"]), 1.0)
    eq("Q2: half reached counts as recurring", 1.0 if r["answer"][0] == "recur across vendors" else 0.0, 1.0)
    eq("Q2: groups of 49 with 50", float(r["groups_of_49_with_50"]), 0.0)
    # the bootstrap: same seed, same interval; identical groupings give no spread
    G = {"A1": g1, "A2": g2, "B1": g2, "B2": g2, "B1j": g2, "B2j": g2, "B1p": g2, "B2p": g2, "P": g2}
    b1, b2 = bootstrap(G, ids, resamples=200), bootstrap(G, ids, resamples=200)
    eq("bootstrap repeats with its seed", b1["B - A"]["low"], b2["B - A"]["low"])
    print("SELFTEST", "PASSES" if ok else "FAILS")
    return ok


def main(argv):
    if not argv:
        raise SystemExit(__doc__)
    mode, opts, files = argv[0], {}, []
    for a in argv[1:]:
        if a.startswith("--"):
            k, _, v = a[2:].partition("=")
            opts[k] = v
        else:
            files.append(a)
    if mode == "selftest":
        sys.exit(0 if selftest() else 1)
    if mode == "check" and opts.get("arm") in ("A", "B") and "key" in opts and len(files) == 1:
        ids = [r["reading"] for r in json.load(open(opts["key"], encoding="utf-8"))["readings"]]
        rep, faults = load_reply(files[0])
        if not faults:
            faults, _ = check_reply(rep, ids, opts["arm"])
        print("COUNTS" if not faults else "DOES NOT COUNT:\n- " + "\n- ".join(faults))
        sys.exit(1 if faults else 0)
    if mode == "score" and {"key", "blank", "sheet", "a1", "a2", "b1", "b2", "out"} <= set(opts):
        score(opts)
        return
    raise SystemExit(__doc__)


if __name__ == "__main__":
    main(sys.argv[1:])
