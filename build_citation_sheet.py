"""
Build the PI's matching sheet of the citation-only run (T36), by CITATION_ONLY_PLAN.md.
Written on 16 Oct 2026, before any reply of the run was opened, on T33's
`build_matching_pack.py`.

    python3 -B build_citation_sheet.py check --chat=58|59|60|61 REPLY
    python3 -B build_citation_sheet.py build --seed=20261020 --out=DIR REPLY_58 REPLY_59 REPLY_60 REPLY_61

check: whether one reply counts, by the plan's "Checks on the replies": one JSON
object (from its first "{" to its last "}") with a field `what_you_know` and a list
`readings` of one to five items, each with a non-empty string `wager`. It also lists
every quoted span of more than eight words in any string of the reply, by where it
is and how many words it has, never by its words, and says whether the span lies
within the citation the chat was given; by the plan, a reply with such a span is
shown to the PI before it is filed, and the span is not filed. Prints the faults,
if any, and exits 1 when the reply does not count, 2 when it counts but holds such
a span, 0 otherwise.

A quoted span is the text between a pair of quote marks in a string of the reply:
an opening and the next closing curly double quote, guillemets, or the 1st and
2nd, 3rd and 4th (and so on) straight double quotes, an opening double quote left
unclosed running to the end of its string; or a single quote, curly or straight,
that does not follow a letter or digit and is followed by a non-space, closed by
the next one of its kind that follows a non-space and is not followed by a letter
or digit (an apostrophe neither opens nor closes; an unclosed single quote is no
span). Its words are counted by splitting it on white space.

build: from the four replies, as filed, and the readings made with the paper,
writes to DIR:

- citation_matching_sheet.xlsx, the PI's sheet: Instructions; Bechky, the
  readings of chats 58 and 59, each shown by its wager alone, shuffled with the
  seed and numbered X01 onward, with a yellow cell for the PI's match (one of the PI's ten
  labels of 14 Oct, or none; the cell offers them) and one for a note; Bechky
  groups, the PI's ten groups of 14 Oct as the PI's sheet grouping_sheet_filled.xlsx labels
  them (the grouping matching_results.json gives), in the order of their first
  reading, each reading as READINGS_matching.md gives it; Pratt, the readings of
  chats 60 and 61 in the same form, numbered on from Bechky's, with a yellow cell
  for the readings of P01 to P15 that make the same point (or none) and one for a
  note; and Pratt readings, P01 to P15 as stage2/markers/READINGS_pratt.md gives
  them, their sentences joined into one paragraph. Nothing on it says which chat
  or model wrote which reading, or which arm wrote P01 to P15.
- citation_only_key.json, for nobody: the seed; each reply's file, sha256, chat,
  model, paper and number of readings; the sha256 of every file read and of this
  script; the groups; P01 to P15 with their arms; for each X its paper, chat,
  model, place in its reply, reading_id and wager; and the sheet's sha256.

The build stops, and says why, if a reply does not count, if a wager holds a
quoted span of more than eight words, if READINGS_matching.md is not the file
matching_key.json describes, if the grouping read from the PI's sheet is not the
one of matching_results.json or a group is written in more than one way, or if
READINGS_pratt.md does not hold KEY.json's sentences of P01 to P15. The shuffle:
random.Random(seed) shuffles Bechky's list (chat 58's readings, then 59's, each in
its reply's order), then Pratt's (60's, then 61's). The workbook carries fixed
dates inside, so two builds give the same bytes.
"""
import datetime
import hashlib
import json
import os
import random
import re
import sys
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import build_worth_sheet as bws                                # noqa: E402
from openpyxl import Workbook, load_workbook                   # noqa: E402
from openpyxl.styles import Alignment                          # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation   # noqa: E402

CHATS = {"58": ("bechky", "Fable 5.1 at max"), "59": ("bechky", "GPT 6 Astra at max"),
         "60": ("pratt", "Fable 5.1 at max"), "61": ("pratt", "GPT 6 Astra at max")}
PAPERS = ("bechky", "pratt")
CORPUS = {"bechky": "bechky_2020_v1.json", "pratt": "pratt_2000_v1.json"}
READINGS_M = "READINGS_matching.md"
MATCHING_KEY = "matching_key.json"
GROUPING = "grouping_sheet_filled.xlsx"
MATCHING_RESULTS = "matching_results.json"
READINGS_P = "stage2/markers/READINGS_pratt.md"
KEY_11 = "stage2/markers/KEY.json"
SHEET = "citation_matching_sheet.xlsx"
KEY = "citation_only_key.json"
MAX_QUOTE = 8
STAMP = (2026, 10, 16, 0, 0, 0)
SHEETS = ("Instructions", "Bechky", "Bechky groups", "Pratt", "Pratt readings")
MATCH_HEAD = {"Bechky": "Match: your group, or none", "Pratt": "Match: P01 to P15, or none"}

RULE = ("Two readings make the same point when they contest the same claim of the paper with the same "
        "alternative account, whatever their wording, emphasis or evidence. Match each reading by its main point, "
        "the one its wager states first. This is the rule of your grouping of 14 Oct.")
EXAMPLE = [
    ("Groups", "Earlier readings of a made-up paper on restaurant kitchens are in three groups: a, cooks skip the "
               "tasting step because the recipe cards go missing (against claim 1, that they value speed over care); "
               "b, because the head chef forbids tasting at the pass (against claim 1); c, the head chefs who taste "
               "every dish are those who arrive before the rush (against claim 2)."),
    ("X1", "Cooks skip tasting because no one has written down what each dish should taste like. Match: a, the "
           "same alternative to claim 1 in other words."),
    ("X2", "The kitchens that skip tasting are those that lose the most staff. Match: none, since no group "
           "contests claim 1 with this alternative."),
]
RULES = [
    "Match by the point, not by worth, and each reading on its own words, whoever wrote it.",
    "Nothing on the sheet shows which chat or model wrote which reading. Please do not open "
    "citation_only_key.json until your sheet is sent.",
    "You may consult the papers. A page in a reading is a pointer, not a proof.",
    "Save the workbook as citation_matching_sheet_filled.xlsx, with no other change, and send it in our "
    "conversation or put it in replies/.",
]


def here(name):
    return os.path.join(HERE, name)


def sha(path):
    return hashlib.sha256(open(path, "rb").read()).hexdigest()


# ---------------------------------------------------------------- the replies

def load(path):
    text = open(path, encoding="utf-8").read()
    a, b = text.find("{"), text.rfind("}")
    if a < 0 or b < a:
        return None, ["no JSON object found"]
    try:
        return json.loads(text[a:b + 1]), []
    except json.JSONDecodeError as e:
        return None, [f"not one JSON object: {e}"]


def faults_of(rep):
    if not isinstance(rep, dict):
        return ["not a JSON object"]
    faults = []
    if "what_you_know" not in rep:
        faults.append("no field what_you_know")
    rs = rep.get("readings")
    if not isinstance(rs, list) or not 1 <= len(rs) <= 5:
        faults.append("`readings` is not a list of one to five items")
    else:
        for i, r in enumerate(rs, 1):
            if not isinstance(r, dict) or not isinstance(r.get("wager"), str) or not r["wager"].strip():
                faults.append(f"reading {i}: no non-empty wager")
    return faults


def spans(s):
    """(text, unclosed) for every quoted span in the string s, as the docstring defines a span."""
    out = []
    for o, c in (("\u201c", "\u201d"), ("\u00ab", "\u00bb")):
        i = 0
        while True:
            a = s.find(o, i)
            if a < 0:
                break
            b = s.find(c, a + 1)
            out.append((s[a + 1:(b if b >= 0 else len(s))], b < 0))
            if b < 0:
                break
            i = b + 1
    marks = [m.start() for m in re.finditer('"', s)]
    for k in range(0, len(marks), 2):
        b = marks[k + 1] if k + 1 < len(marks) else len(s)
        out.append((s[marks[k] + 1:b], k + 1 >= len(marks)))
    for o, c in (("\u2018", "\u2019"), ("'", "'")):
        opener = re.compile(r"(?<![^\W_])" + re.escape(o) + r"(?=\S)")
        closer = re.compile(r"(?<=\S)" + re.escape(c) + r"(?![^\W_])")
        i = 0
        while True:
            m = opener.search(s, i)
            if not m:
                break
            n = closer.search(s, m.end())
            if not n:
                break
            out.append((s[m.end():n.start()], False))
            i = n.end()
    return out


def quoted_spans(s):
    """(words, unclosed) for every quoted span in the string s."""
    return [(len(t.split()), u) for t, u in spans(s)]


def norm(s):
    s = s.translate({0x201c: '"', 0x201d: '"', 0x2018: "'", 0x2019: "'", 0x2013: "-", 0x2014: "-"})
    return " ".join(s.casefold().split())


def strings(obj, path="reply"):
    if isinstance(obj, str):
        yield path, obj
    elif isinstance(obj, dict):
        for k, v in obj.items():
            yield from strings(v, f"{path}.{k}")
    elif isinstance(obj, list):
        for i, v in enumerate(obj):
            yield from strings(v, f"{path}[{i}]")


def long_spans(rep, cite):
    """(where, words, unclosed, within the citation) for each quoted span of more than eight words; the span's
    words are compared with the citation, never returned."""
    out = []
    for where, s in strings(rep):
        for t, unclosed in spans(s):
            if len(t.split()) > MAX_QUOTE:
                inside = norm(t.strip(" .,;:!?\"'\u201c\u201d\u2018\u2019")) in norm(cite)
                out.append((where, len(t.split()), unclosed, inside))
    return out


def citation(paper):
    return json.load(open(here(CORPUS[paper]), encoding="utf-8"))["meta"]["citation"]


def check_one(path, chat):
    """Print the check of one reply; return its exit code."""
    paper, model = CHATS[chat]
    rep, faults = load(path)
    faults = faults or faults_of(rep)
    print(f"chat {chat} ({paper}, {model}): {os.path.basename(path)}, sha256 {sha(path)}")
    if faults:
        for f in faults:
            print(f"  FAULT: {f}")
        print("  DOES NOT COUNT")
        return 1
    rs = rep["readings"]
    ids = [r.get("reading_id") for r in rs]
    wyk = rep["what_you_know"]
    print(f"  paper field: {rep.get('paper')!r}")
    print(f"  what_you_know: {'a string of %d words' % len(wyk.split()) if isinstance(wyk, str) else type(wyk).__name__}")
    print(f"  readings: {len(rs)}; reading_ids {ids}" + ("" if len(set(map(str, ids))) == len(ids) else
                                                         " (repeated: readings are keyed by place)"))
    for i, r in enumerate(rs, 1):
        print(f"  reading {i}: wager {len(r['wager'].split())} words; "
              f"fields {sorted(k for k in r if k != 'wager')}")
    spans = long_spans(rep, citation(paper))
    for where, words, unclosed, inside in spans:
        print(f"  QUOTED SPAN over {MAX_QUOTE} words: {where}: {words} words"
              f"{', unclosed' if unclosed else ''}{', within the citation' if inside else ''}")
    if spans:
        print("  COUNTS; to be shown to the PI before it is filed (the plan): the spans above are not filed")
        return 2
    print("  COUNTS; no quoted span of more than eight words")
    return 0


# ---------------------------------------------------------------- the readings made with the paper

def m_readings():
    k = json.load(open(here(MATCHING_KEY), encoding="utf-8"))
    text = open(here(READINGS_M), encoding="utf-8").read()
    if hashlib.sha256(text.encode()).hexdigest() != k["readings_sha256"]:
        raise SystemExit(f"{READINGS_M} is not the file {MATCHING_KEY} describes")
    parts = re.split(r"^## Reading (M\d\d)\n", text, flags=re.M)
    got = {parts[i]: parts[i + 1].strip() for i in range(1, len(parts), 2)}
    want = {r["reading"]: r["text"] for r in k["readings"]}
    if got != want or list(got) != [f"M{i:02d}" for i in range(1, len(want) + 1)]:
        raise SystemExit(f"{READINGS_M}: its readings are not {MATCHING_KEY}'s, in order")
    return got


def his_groups(texts):
    ws = load_workbook(here(GROUPING))["Bechky"]
    written = {}
    for r in range(2, ws.max_row + 1):
        rid = ws.cell(r, 1).value
        if rid is None:
            continue
        lab = ws.cell(r, 3).value
        written[rid] = "" if lab is None else str(lab).strip()
    res = json.load(open(here(MATCHING_RESULTS), encoding="utf-8"))
    if {i: w.casefold() for i, w in written.items()} != res["groupings"]["P"]:
        raise SystemExit(f"the grouping on {GROUPING} is not the one of {MATCHING_RESULTS}")
    if sorted(written) != sorted(texts):
        raise SystemExit(f"{GROUPING}: its readings are not those of {READINGS_M}")
    groups = {}
    for i in sorted(written):
        groups.setdefault(written[i].casefold(), []).append(i)
    out = []
    for cf, members in sorted(groups.items(), key=lambda g: g[1][0]):
        forms = {written[i] for i in members}
        if len(forms) != 1:
            raise SystemExit(f"group {cf!r} is written in more than one way: {sorted(forms)}")
        out.append({"label": forms.pop(), "readings": members})
    if len(out) != 10:
        raise SystemExit(f"{len(out)} groups, not ten")
    return out


def p_readings():
    k = json.load(open(here(KEY_11), encoding="utf-8"))["papers"]["pratt"]
    if [e["reading"] for e in k] != [f"P{i:02d}" for i in range(1, 16)]:
        raise SystemExit("KEY.json: Pratt's readings are not P01 to P15 in order")
    text = open(here(READINGS_P), encoding="utf-8").read()
    parts = re.split(r"^## Reading (P\d\d)\n", text, flags=re.M)
    sections = {parts[i]: parts[i + 1] for i in range(1, len(parts), 2)}
    out = []
    for e in k:
        p = e["reading"]
        lines = re.findall(r"^" + p + r"-(\d\d)  (.+)$", sections.get(p, ""), flags=re.M)
        if [t for _, t in lines] != [u["text"] for u in e["units"]] or \
                [int(n) for n, _ in lines] != list(range(1, len(lines) + 1)):
            raise SystemExit(f"{READINGS_P}: {p} is not KEY.json's sentences in order")
        out.append({"reading": p, "arm": e["arm"], "text": " ".join(t for _, t in lines)})
    return out


# ---------------------------------------------------------------- the workbook

def steady(path):
    stamp = b"2026-10-16T00:00:00Z"
    with zipfile.ZipFile(path) as z:
        items = [(i, z.read(i.filename)) for i in z.infolist()]
    with zipfile.ZipFile(path, "w") as z:
        for i, data in items:
            if i.filename == "docProps/core.xml":
                data, n = re.subn(rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)",
                                  rb"\g<1>" + stamp + rb"\g<2>", data)
                if n != 1:
                    raise SystemExit("docProps/core.xml: no single modified date")
            j = zipfile.ZipInfo(i.filename, date_time=STAMP)
            j.compress_type = i.compress_type
            j.external_attr = i.external_attr
            z.writestr(j, data)


def instructions(wb, rows):
    ins = wb.active
    ins.title = "Instructions"
    ins.column_dimensions["A"].width = 16
    ins.column_dimensions["B"].width = 110
    for r, (label, t, bold, size) in enumerate(rows, 1):
        bws.cell(ins, f"A{r}", label, bold=bold, size=size, wrap=r != 1)
        if t is not None:
            bws.cell(ins, f"B{r}", t, bold=False, size=size)
            ins.row_dimensions[r].height = bws.height(t, 120)
    bws.printable(ins)


def answer_sheet(wb, title, entries, validate=None):
    ws = wb.create_sheet(title)
    for col, w in zip("ABCD", (10, 100, 24, 45)):
        ws.column_dimensions[col].width = w
    for i, h in enumerate(["Reading", "The reading (its wager)", MATCH_HEAD[title], "Note"]):
        bws.cell(ws, f"{'ABCD'[i]}1", h, bold=True, fill=bws.GREY).border = bws.BOX
    if validate is not None:
        ws.add_data_validation(validate)
    for r, e in enumerate(entries, start=2):
        bws.cell(ws, f"A{r}", e["x"], bold=True).border = bws.BOX
        bws.cell(ws, f"B{r}", e["wager"]).border = bws.BOX
        c = bws.cell(ws, f"C{r}", None, fill=bws.YELLOW)
        c.border = bws.BOX
        c.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
        bws.cell(ws, f"D{r}", None, fill=bws.YELLOW).border = bws.BOX
        if validate is not None:
            validate.add(f"C{r}")
        ws.row_dimensions[r].height = bws.height(e["wager"], 112)
    ws.freeze_panes = "B2"
    bws.printable(ws)


def reference_sheet(wb, title, heads, widths, rows, text_col, text_width):
    ws = wb.create_sheet(title)
    cols = "ABC"[:len(heads)]
    for col, w in zip(cols, widths):
        ws.column_dimensions[col].width = w
    for i, h in enumerate(heads):
        bws.cell(ws, f"{cols[i]}1", h, bold=True, fill=bws.GREY).border = bws.BOX
    for r, row in enumerate(rows, start=2):
        for i, v in enumerate(row):
            bws.cell(ws, f"{cols[i]}{r}", v, bold=(i < len(row) - 1)).border = bws.BOX
        ws.row_dimensions[r].height = bws.height(row[text_col], text_width)
    ws.freeze_panes = "A2"
    bws.printable(ws)


def build(out, seed, paths):
    sources, lists = {}, {p: [] for p in PAPERS}
    for chat, path in zip(sorted(CHATS), paths):
        paper, model = CHATS[chat]
        rep, faults = load(path)
        faults = faults or faults_of(rep)
        if faults:
            raise SystemExit(f"chat {chat}: the reply does not count: {'; '.join(faults)}")
        for i, r in enumerate(rep["readings"], 1):
            n = [w for w, _ in quoted_spans(r["wager"]) if w > MAX_QUOTE]
            if n:
                raise SystemExit(f"chat {chat}, reading {i}: its wager holds a quoted span of {max(n)} words")
            lists[paper].append({"paper": paper, "chat": chat, "model": model, "place": i,
                                 "reading_id": r.get("reading_id"), "wager": r["wager"].strip()})
        sources[chat] = {"file": os.path.basename(path), "sha256": sha(path), "paper": paper, "model": model,
                         "readings": len(rep["readings"])}
    rng = random.Random(seed)
    for paper in PAPERS:
        rng.shuffle(lists[paper])
    n = 0
    for paper in PAPERS:
        for e in lists[paper]:
            n += 1
            e["x"] = f"X{n:02d}"
    texts = m_readings()
    groups = his_groups(texts)
    pratt = p_readings()
    cites = {paper: citation(paper) for paper in PAPERS}

    nb, npr = len(lists["bechky"]), len(lists["pratt"])
    last_b, first_p = lists["bechky"][-1]["x"], lists["pratt"][0]["x"]
    wb = Workbook()
    rows = [("Matching the readings made from the citation alone", None, True, 13),
            ("What", f"Four chats were given a paper's citation, not the paper, and asked for alternative readings "
                     f"of it: two on Bechky (2020) and two on Pratt (2000). Their {n} readings are on the sheets "
                     f"Bechky (X01 to {last_b}) and Pratt ({first_p} to {lists['pratt'][-1]['x']}), each shown by "
                     f"its wager alone, in no particular order. For each, say which readings made with the paper "
                     f"make the same point.", True, 10),
            ("Fill in", "Only the yellow cells, Match and Note, on the sheets Bechky and Pratt. Leave every other "
                        "cell as it is.", True, 10),
            ("The rule", RULE, True, 10),
            ("Bechky", "Name the one group of yours, of 14 Oct, that makes the same point, by its label: the cell "
                       "offers your ten labels and none. Your groups and their readings are on the sheet Bechky "
                       "groups, in the words you grouped them in. Write none when no group makes the point.",
             True, 10),
            ("Pratt", "Name every reading of P01 to P15 that makes the same point, separated by commas (for example: "
                      "P03, P11), or write none when none does. They are on the sheet Pratt readings, as the wager "
                      "test's markers had them, under the same numbers.", True, 10),
            ("Example", "From a made-up paper, not one of the readings:", True, 10)]
    rows += [(a, b, False, 10) for a, b in EXAMPLE]
    rows.append(("Rules", None, True, 10))
    rows += [(str(i), t, False, 10) for i, t in enumerate(RULES, 1)]
    instructions(wb, rows)

    labels = [g["label"] for g in groups]
    formula = '"' + ",".join(labels + ["none"]) + '"'
    if len(formula) > 255 or any("," in x or '"' in x for x in labels):
        raise SystemExit("the labels do not fit a list of choices")
    dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=True,
                        errorTitle="Match", error="Choose one of your ten labels, or none.",
                        showInputMessage=True, promptTitle="Match", prompt="One of your groups, or none.")
    answer_sheet(wb, "Bechky", lists["bechky"], validate=dv)
    reference_sheet(wb, "Bechky groups", ["Your group", "Reading", "The reading"], (24, 10, 100),
                    [(g["label"], m, texts[m]) for g in groups for m in g["readings"]], 2, 112)
    answer_sheet(wb, "Pratt", lists["pratt"])
    reference_sheet(wb, "Pratt readings", ["Reading", "The reading"], (10, 100),
                    [(e["reading"], e["text"]) for e in pratt], 1, 112)
    if tuple(wb.sheetnames) != SHEETS:
        raise SystemExit(f"sheets {wb.sheetnames}")
    wb.properties.created = wb.properties.modified = datetime.datetime(*STAMP)
    wb.properties.creator = "build_citation_sheet.py"
    os.makedirs(out, exist_ok=True)
    sheet_path = os.path.join(out, SHEET)
    wb.save(sheet_path)
    steady(sheet_path)

    files = [READINGS_M, MATCHING_KEY, GROUPING, MATCHING_RESULTS, READINGS_P, KEY_11, CORPUS["bechky"],
             CORPUS["pratt"]]
    key = {"what": f"the key to {SHEET}: for nobody, and not for the PI until the PI's sheet is sent",
           "seed": seed, "sources": sources,
           "inputs": {f: sha(here(f)) for f in files},
           "builder_sha256": sha(os.path.abspath(__file__)),
           "citations": cites,
           "groups": groups,
           "pratt_readings": [{"reading": e["reading"], "arm": e["arm"]} for e in pratt],
           "readings": [{k: e[k] for k in ("x", "paper", "chat", "model", "place", "reading_id", "wager")}
                        for paper in PAPERS for e in lists[paper]],
           "sheet_sha256": sha(sheet_path)}
    open(os.path.join(out, KEY), "w", encoding="utf-8").write(json.dumps(key, indent=1, ensure_ascii=False) + "\n")
    print(f"{SHEET}: Bechky X01 to {last_b} ({nb}), Pratt {first_p} to {lists['pratt'][-1]['x']} ({npr}); "
          f"{len(groups)} groups, P01 to P15; sha256 {key['sheet_sha256']}")
    print(f"{KEY}: sha256 {sha(os.path.join(out, KEY))}")


def main(argv):
    opts, files = {}, []
    for a in argv[1:]:
        if a.startswith("--"):
            k, _, v = a[2:].partition("=")
            opts[k] = v
        else:
            files.append(a)
    mode = argv[0] if argv else None
    if mode == "check" and opts.get("chat") in CHATS and len(files) == 1:
        sys.exit(check_one(files[0], opts["chat"]))
    elif mode == "build" and {"out", "seed"} <= set(opts) and len(files) == 4:
        build(opts["out"], int(opts["seed"]), files)
    else:
        raise SystemExit(__doc__)


if __name__ == "__main__":
    main(sys.argv[1:])
