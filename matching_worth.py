"""
Set the PI's worth for the ten new readings of T33 beside the PI's earlier worth for
the readings the PI grouped with them. Written on 14 Oct 2026, after the PI's worth
sheet came in. Described, not tested: the plan of T33 fixes no figure on worth.

    python3 -I -B matching_worth.py <out_dir>

Reads matching_key.json (which reading came from which chat, and its earlier
id), grouping_sheet_filled.xlsx (the PI's groups, labels trimmed and read without
regard to capitals, as score_matching.py reads them), worth_sheet_matching_filled.xlsx
(the PI's worth for the readings of chats 49 and 50), stage2/worth/worth_sheet_filled_2.xlsx
(the PI's worth of 11 Oct for B01 to B14) and worth_sheet_places_filled.xlsx (the PI's worth
of 13 Oct for W01, W03, W05 and W06, chat 46's). Writes MATCHING_WORTH.md.
"""
import json
import os
import sys

from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))


def sheet_values(path, sheet, col=3):
    ws = load_workbook(os.path.join(HERE, path))[sheet]
    out = {}
    for r in range(2, ws.max_row + 1):
        rid = ws.cell(r, 1).value
        if rid is not None:
            out[rid] = (ws.cell(r, col).value, ws.cell(r, col + 1).value)
    return out


def worth_of(rid, v):
    # The sheets keep a worth as text ("2"); a whole number 0 to 3 is taken, nothing else.
    s = "" if v is None else str(v).strip()
    if s not in ("0", "1", "2", "3"):
        raise SystemExit(f"{rid}: worth {v!r}")
    return int(s)


def mean(xs):
    return sum(xs) / len(xs) if xs else None


def fmt(x):
    return "—" if x is None else f"{x:.2f}"


def main(out):
    key = json.load(open(os.path.join(HERE, "matching_key.json"), encoding="utf-8"))
    groups = {k: str(v[0]).strip().casefold() for k, v in sheet_values("grouping_sheet_filled.xlsx", "Bechky").items()}
    new = sheet_values("worth_sheet_matching_filled.xlsx", "Bechky")
    old_b = sheet_values("stage2/worth/worth_sheet_filled_2.xlsx", "Bechky")
    old_w = sheet_values("worth_sheet_places_filled.xlsx", "Bechky")
    if set(new) != {r["reading"] for r in key["readings"] if r["chat"] in ("49", "50")}:
        raise SystemExit(f"the worth sheet holds {sorted(new)}, not the readings of chats 49 and 50")
    if set(groups) != {r["reading"] for r in key["readings"]}:
        raise SystemExit("the grouping sheet does not hold the readings of the key")
    rows = []
    for r in key["readings"]:
        if r["chat"] in ("49", "50"):
            w, note = new[r["reading"]]
            when = "14 Oct"
        elif r["earlier_id"].startswith("B"):
            w, note = old_b[r["earlier_id"]]
            when = "11 Oct"
        else:
            w, note = old_w[r["earlier_id"]]
            when = "13 Oct"
        rows.append({"reading": r["reading"], "chat": r["chat"], "id": r["earlier_id"] or r["source_id"],
                     "group": groups[r["reading"]], "worth": worth_of(r["reading"], w), "when": when,
                     "note": (note or "").strip()})
    newrows = [x for x in rows if x["chat"] in ("49", "50")]
    L = ["# The worth of T33's new readings, beside the readings grouped with them", "",
         "Written by `matching_worth.py`. Described, not tested: the plan of T33 fixes no figure on worth. All worth "
         "is the PI's, on the wager test's scale (0 not worth pursuing, 1 some interest, 2 worth a scholar's time, 3 a "
         "strong opening), given on three occasions: B01 to B14 on the PI's sheet of 11 Oct (declared not blind: the PI had seen "
         "every reply; B06 was rated late, after the PI had been told its arm), chat 46's four on 13 Oct (on a sheet that did "
         "not say which chat wrote which), and the ten new ones on 14 Oct (not blind to their source: the sheet named "
         "chats 49 and 50; it was rated after the results of the test). What the PI knew "
         "differs between the occasions, so the means below do not compare the ways the readings were made. The groups "
         "are the PI's, from the PI's grouping sheet, made after the PI had run the four coder chats.", "",
         "## The ten new readings", "",
         "| reading | chat | its id there | the PI's group | worth | the PI's note |", "|---|---|---|---|---|---|"]
    for x in sorted(newrows, key=lambda x: (x["chat"], x["id"])):
        L.append(f"| {x['reading']} | {x['chat']} | {x['id']} | {x['group']} | {x['worth']} | {x['note']} |")
    nw = [x["worth"] for x in newrows]
    L += ["", f"Mean {fmt(mean(nw))}; rated 3: {nw.count(3)}, 2: {nw.count(2)}, 1: {nw.count(1)}, 0: {nw.count(0)}.", "",
          "## By the PI's group", "",
          "| the PI's group | earlier readings (chat, id, worth) | new readings (chat, id, worth) | earlier mean | new mean |",
          "|---|---|---|---|---|"]
    by = {}
    for x in rows:
        by.setdefault(x["group"], []).append(x)
    for g in sorted(by, key=lambda g: (-len(by[g]), g)):
        e = [x for x in by[g] if x["chat"] not in ("49", "50")]
        n = [x for x in by[g] if x["chat"] in ("49", "50")]
        es = "; ".join(f"{x['chat']} {x['id']} {x['worth']}" for x in sorted(e, key=lambda x: (x["chat"], x["id"]))) or "—"
        ns = "; ".join(f"{x['chat']} {x['id']} {x['worth']}" for x in sorted(n, key=lambda x: (x["chat"], x["id"]))) or "—"
        L.append(f"| {g} | {es} | {ns} | {fmt(mean([x['worth'] for x in e]))} | {fmt(mean([x['worth'] for x in n]))} |")
    ew = [x["worth"] for x in rows if x["chat"] not in ("49", "50")]
    L += ["", f"The 18 earlier readings: mean {fmt(mean(ew))}; rated 3: {ew.count(3)}, 2: {ew.count(2)}, 1: {ew.count(1)}, "
          f"0: {ew.count(0)}.", ""]
    os.makedirs(out, exist_ok=True)
    open(os.path.join(out, "MATCHING_WORTH.md"), "w", encoding="utf-8").write("\n".join(L))
    print(f"new mean {fmt(mean(nw))}, earlier mean {fmt(mean(ew))}; {len(by)} groups")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit(__doc__)
    main(sys.argv[1])
