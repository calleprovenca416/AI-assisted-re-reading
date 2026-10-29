"""
Build coder 2's four workbooks for one cold paper: Stages A, A2, V and B, as
the cold-paper CODING_INSTRUCTIONS.md describes them. Written 28 Sep 2026;
revised the same day.

WHAT GOES OUT, AND WHAT NEVER DOES
----------------------------------
A coder receives CODING_INSTRUCTIONS.md, the paper and one stage's workbook,
nothing else. There is no corpus file. Every tab prints only the fields listed
for it here; nothing is copied wholesale from a record, so a field cannot leak
by being forgotten.

  Stage A  (section 0)   0 outcome             claim, statement, location
  Stage A2 (section 0b)  0b asserted_direction claim, statement, the outcome
                                               recorded for it; only claims
                                               that name an outcome
  Stage V  (section V)   V voice               fragment, location, and the stub
                                               with its voice cue removed
  Stage B  (sections 1-7)
           1 direction         fragment, stub, claim, statement, outcome,
                               asserted direction (both blank where the claim
                               has no outcome, as section 1 says); rows in
                               fragment order, then claim order
           2 core              claim, statement
           3 stated_as         claim, statement, location
           4 gap kind          gap, statement, location
           5 pair drawn_by     pair, episode A with its pages, episode B with
                               its pages
           6 resolution_period the periods, each with its description
           7 gloss_relation    fragment, stub, gloss, and the fragments either
                               side of it on the page

Never printed: notes, extraction_confidence, a link's direction, basis or
confidence, core, stated_as, a gap's kind, a pair's drawn_by_authors, basis,
held_comparable or differs_in, resolution_period or resolution_basis,
data_type, voice or gloss relation, and nothing from the extraction's sidecar.

THE RULES IT ENFORCES
- V removes the three voice markers ("the author's account is that", "the paper
  explains that", "for the author,"), in any spelling of the apostrophe, in
  "author" or "authors", after markdown or stray leading characters. It refuses
  a row that still opens like a marker, and a fragment with no location. It
  also removes near-misses of the three ("The author's account:", "For the
  author the ..."), "according to the author," and an opening "the
  paper/study/author/researcher reports/argued/suggests/... that", which tilts
  the voice judgement either way; other such openings, and a cue later in a
  stub, are listed as warnings. A description left empty or still opening
  like a marker stops the build; a neutral one is then supplied in a JSON
  file ({"F12": "..."}), with the reason recorded, and passed with
  --v-descriptions, whose text is
  checked the same way. V never prints data_type, which tracks voice closely.
- B prints stubs as they stand, markers included: V is sent, and comes back,
  before B.
- B prints the period descriptions (section 6) from the extraction log's
  "Periods" section. A description that speaks to whether or where the story
  closes states the first coder's answer: the build stops. Neutral
  descriptions are then written to a JSON file ({"P1": "..."}), with the
  reason recorded, and passed with --periods; the same check runs on that file.
- B prints the fragments either side of each glossed fragment (section 7),
  in page order.
- No workbook prints notes.
- A voice mismatch noted on a V row whose opening was removed is this
  builder's doing, not stub laundering. The key lists those rows.

Revisions of 28 Sep:
- Tab 5 printed `held comparable because`, as the Carlile pack did. The brief
  tells extractors to say what the page prints, so that text says what the
  paper does with the pair ("both are used in the paper's account ...") and
  states coder 1's answer. Tab 5 now prints the two episodes and their pages
  only; an episode text that mentions the paper, the author, comparing or
  contrasting stops the build (override: --pairs, a JSON file
  {"CP01": {"episode_a": "...", "episode_b": "..."}}, checked the same way).
- Tab 1 followed the extractor's link order, which can put a fragment's
  directional links before its `illustrates` ones. Rows are now sorted.
- Plain-text returns put the code before the reason, against section 1's rule
  of reason first. The format is now: row id | ids | reason | code, where ids
  are the fragment and claim ids printed in that row (- for section 6), so
  that intake_coder2.py can tell whether a line answers the row it says.

Revision, 7 Oct 2026:
- One of the words that builder-written text may not name is no longer spelled
  in check_must_agree.py, which keeps a hash of it. The builder takes that
  test from there with the listed patterns (machinery_patterns), so it tries
  every word of what it writes, and of what it prints from the paper, against
  the hash. Where such a word is found it is reported as "a guarded word (not
  shown)", never spelled. The self-test tries this on a made-up word.

Revision, 7 Oct 2026, later the same day (version 1.3):
- THE SIDECAR IS NAMED AND RECORDED. --sidecar=<voice_gloss.json> is now
  required: the extraction's sidecar, frozen with the corpus. The builder
  checks that it is this paper's and codes every fragment, and the key records
  its name and hash beside the corpus's. Nothing from it is printed.
  agreement.py then takes as coder 1's voice and gloss relation only the
  sidecar the key records. Before this the sidecar was tied to nothing.
- A message that quotes a text of the paper, and the key's copies of the pair
  texts and the period descriptions, show a guarded word struck out (###).
  They quoted the text as it stood.
- The self-test now tries the check of the builder's own text in a build, with
  a made-up word, and not only the function that the check calls.
- After that review's second reading, the same day: THE SIDECAR IS CHECKED IN
  ONE PLACE, read_sidecar(), which agreement.py takes from here. Before this
  the builder looked at the paper's id and the voice codes only, and could
  record the hash of a sidecar that agreement.py then refused for its
  `source` line or its gloss relations. A glossed fragment with no gloss
  relation is a warning: that row cannot be compared.

Two choices made here, reversible:
- Stage A does not print `constructs` (the Carlile pack did): the outcome is
  usually one of them.
- Section 6's code list is the periods and `none`, for a paper that does not
  close; section 6 says to explain that in the reason column.

OUTPUT, in <outdir> (which must not exist yet or must be empty):
  attach/                        the four workbooks: the only files from here
                                 that go to a coder chat
  SEND_1_STAGE_A.md ... SEND_4_STAGE_B.md
                                 for the PI: what to attach and the opening
                                 message for each stage's chat, in that order
  key/<paper_id>_pack_key.json
                                 source and hashes (the corpus's and the
                                 sidecar's), row ids, the V rows whose
                                 opening was removed, period and pair texts
                                 and their source, gloss neighbours, warnings

USAGE
    python3 build_cold_pack.py <corpus_vN.json> <extraction_log.md> <outdir>
                               --sidecar=<voice_gloss.json>
                               [--periods=<file.json>] [--pairs=<file.json>]
                               [--v-descriptions=<file.json>] [--allow-unfrozen]
    python3 build_cold_pack.py --selftest

The corpus must be a frozen version (<paper_id>_v<N>.json) and validate with no
errors. --allow-unfrozen marks the pack as a TEST PACK, not to be sent.

A pack is built once. After a stage has gone out it is never built again: a
reply that comes back as text carries row ids only, and nothing in it says
which pack it answered.
"""

import hashlib
import json
import os
import re
import shutil
import sys
import tempfile
import unicodedata
from datetime import datetime, timezone

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

BUILDER_VERSION = "1.3, 7 Oct 2026"

# ---------------------------------------------------------------- the codes
CODES = {
    "V voice": ["reported", "author_insight"],
    "1 direction": ["with", "against", "illustrates", "qualifies"],
    "2 core": ["true", "false"],
    "3 stated_as": ["established", "ventured"],
    "4 gap kind": ["scope_limit", "absent_evidence", "unresolved_question",
                   "acknowledged_ambiguity", "vocabulary_mismatch"],
    "5 pair drawn_by": ["true", "false"],
    # "6 resolution_period": the paper's periods, plus none -- set per paper
    "7 gloss_relation": ["consistent", "goes_beyond", "diverges", "none"],
}
FREE_TEXT = {"0 outcome": "a short noun phrase, or none",
             "0b asserted_direction": "a short conditional, or none"}
ROW_PREFIX = {"0 outcome": "0", "0b asserted_direction": "0b", "V voice": "V",
              "1 direction": "1", "2 core": "2", "3 stated_as": "3",
              "4 gap kind": "4", "5 pair drawn_by": "5",
              "6 resolution_period": "6", "7 gloss_relation": "7"}
# the columns whose values a text line's "ids" field must repeat
ID_COLUMNS = {"0 outcome": ["claim"], "0b asserted_direction": ["claim"],
              "V voice": ["fragment"], "1 direction": ["fragment", "claim"],
              "2 core": ["claim"], "3 stated_as": ["claim"], "4 gap kind": ["gap"],
              "5 pair drawn_by": ["pair"], "6 resolution_period": [],
              "7 gloss_relation": ["fragment"]}
REASON_H = "YOUR REASON, one line"
CODE_H = "YOUR CODE"
NO_OUTCOME = {"", "none", "n/a", "na", "null", "-", "--"}
TEXT_FORMAT = "row id | ids | reason | code"

# ---------------------------------------------------------------- voice cues
APOS = dict.fromkeys(map(ord, "‘’ʼ′`´"), "'")
_VERBS = (r"reports|reported|notes|noted|records|recorded|recounts|recounted|states|stated|"
          r"observes|observed|shows|showed|documents|documented|argues|argued|suggests|"
          r"suggested|claims|claimed|concludes|concluded|proposes|proposed|infers|inferred|"
          r"interprets|interpreted|contends|contended|maintains|maintained|explains|explained|"
          r"believes|believed")
_WHO = r"the (?:paper|study|authors?|researchers?)"
# the three markers, exactly as the brief writes them: removed
MARKER = re.compile(r"^(the authors?'?s? account is that|the paper explains that|"
                    r"for the authors?,)\s*", re.I)
# near-misses of the three, and "according to the author": removed
MARKER_LOOSE = re.compile(r"^(the authors?'?s? account(?:\s+(?:is|was)(?:\s+that)?\s*[:,]?|\s*[:,])|"
                          r"the paper explains\s*[:,]|for the authors?(?!')\b\s*[,:]?|"
                          r"according to " + _WHO + r"\s*,?)\s*", re.I)
# "the paper reports that ...", "the author argued that ...": removed
CUE_THAT = re.compile(r"^" + _WHO + r"\s+(?:(?:" + _VERBS + r")\s+that|(?:describes|described)\s+how)\s+",
                      re.I)
# other such openings: listed as warnings
CUE_OTHER = re.compile(r"^" + _WHO + r"\s+(?:" + _VERBS + r"|presents?|presented|describes?|"
                       r"described|reads?)\b", re.I)
# still marker-like after removal: refused (override with --v-descriptions)
MARKER_LIKE = re.compile(r"^(the authors?'?s? account\b|the paper explains\b|for the authors?(?!')\b)",
                         re.I)
MARKER_ANYWHERE = re.compile(r"(the authors?'?s? account\b|the paper explains\b|"
                             r"for the authors?(?!')\b|according to " + _WHO + r")", re.I)

# words that let a period description state where the story closes
CLOSURE = re.compile(r"\b(resol\w*|settl\w*|clos\w*|conclu\w*|ongoing|bounded|single-period|"
                     r"no story|final\w*|end\w*|complet\w*|aftermath|outcome\w*)\b", re.I)
# words that let an episode text say what the paper does with the pair
PAIRWORDS = re.compile(r"\b(compar\w*|contrast\w*|juxtapos\w*|side by side|the paper|"
                       r"the authors?|the study)\b", re.I)

PAPER_NAME_PATTERNS = {r"\bcarlile\b", r"\bmick\b", r"\bvaughn\b", r"\bpratt",
                       r"\bbechky", r"\bamway\b",
                       r"forensic", r"present tense", r"evaluative (words|vocabulary)",
                       r"generic subject", r"the value of"}


def machinery_patterns():
    """The words the builder's own text may not name: the listed patterns, but
    for the papers' names and the voice cues, and the test for the word that
    check_must_agree.py keeps as a hash. Each has `finditer`, `search` and
    `pattern`; the last one's matches give a label and never the word."""
    from check_must_agree import FORBIDDEN_BOTH, GuardedWords
    return ([re.compile(p, re.I) for p in FORBIDDEN_BOTH if p not in PAPER_NAME_PATTERNS]
            + [GuardedWords()])


def machinery_named(texts, mach=None):
    """What the machinery patterns find in the given texts, sorted. A guarded
    word is given by its label."""
    mach = machinery_patterns() if mach is None else mach
    return sorted({m.group(0) for s in texts for p in mach for m in p.finditer(s)})


def shown(s):
    """A text of the paper as a message or the key may quote it: a guarded
    word struck out, at one length whatever the word's."""
    from check_must_agree import GUARD_STRUCK, GuardedWords
    s = str(s)
    for m in reversed(list(GuardedWords().finditer(s))):
        s = s[:m.start()] + GUARD_STRUCK + s[m.end():]
    return s


# ---------------------------------------------------------------- styling
HEAD = Font(name="Arial", bold=True, color="FFFFFF", size=11)
HEADFILL = PatternFill("solid", fgColor="44546A")
BODY = Font(name="Arial", size=10)
INPUT = PatternFill("solid", fgColor="FFF2CC")
WRAP = Alignment(wrap_text=True, vertical="top")


class Refused(Exception):
    pass


def sha(path):
    return hashlib.sha256(open(path, "rb").read()).hexdigest()


def has_outcome(c):
    return (c.get("outcome") or "").strip().lower() not in NO_OUTCOME


def read_sidecar(path, corpus, c_obj):
    """The extraction's sidecar, checked in one place for this builder and for
    agreement.py, so that the key records no sidecar that agreement.py would
    then refuse. It must be a JSON object for this paper; carry no `source`
    line, which make_voice_sidecar.py writes into a sidecar made from returned
    workbooks; code the voice of every fragment (audit.load_voice); and name,
    in each gloss relation, a fragment of the corpus, once, with a code on the
    list.

    Returns (the sidecar, {fragment: voice code}, warnings). A glossed fragment
    with no gloss relation is a warning and not a refusal: coder 1 then has no
    code for that row, and the paper's agreement log lists it as one that
    cannot be compared."""
    import audit
    name = os.path.basename(path)
    if not os.path.isfile(path):
        raise Refused(f"no such sidecar: {path}")
    try:
        side = json.load(open(path, encoding="utf-8"))
    except ValueError:
        raise Refused(f"{name} is not a JSON file")
    pid = corpus["meta"]["paper_id"]
    if not isinstance(side, dict) or side.get("paper_id") != pid:
        raise Refused(f"{name} is not the sidecar of {pid!r}")
    if "source" in side:
        raise Refused(f"{name}: the sidecar carries a `source` line, which make_voice_sidecar.py "
                      "writes into a sidecar made from returned workbooks: it is not an "
                      "extraction's own")
    try:
        voice = audit.load_voice(path, c_obj)
    except ValueError as e:
        raise Refused(f"{name}: {e}")
    gloss = side.get("gloss_relation") or []
    if not isinstance(gloss, list) or not all(isinstance(g, dict) for g in gloss):
        raise Refused(f"{name}: the sidecar's gloss relations are not a list of entries")
    gids = [g.get("fragment_id") for g in gloss]
    known = {f["fragment_id"] for f in corpus["fragments"]}
    if (not all(isinstance(x, str) for x in gids) or len(set(gids)) != len(gids)
            or set(gids) - known
            or any(g.get("code") not in CODES["7 gloss_relation"] for g in gloss)):
        raise Refused(f"{name}: the sidecar's gloss relations name a fragment twice, name a "
                      "fragment that is not in the corpus, or carry a code that is not on the list")
    bare = [f["fragment_id"] for f in corpus["fragments"]
            if (f.get("author_gloss") or "").strip() and f["fragment_id"] not in gids]
    warnings = []
    if bare:
        warnings.append(f"the sidecar gives no gloss relation for {len(bare)} glossed fragment(s): "
                        f"{', '.join(bare)}. Coder 1 has no code for those rows, so they cannot "
                        "be compared")
    return side, voice, warnings


def clean_for_match(s):
    """NFKC, one kind of apostrophe, single spaces, no markdown emphasis, no
    leading quotes or invisible characters. Used to match; printed only when an
    opening has been removed."""
    s = unicodedata.normalize("NFKC", s or "").translate(APOS)
    s = re.sub(r"[​-‏⁠﻿]", "", s)
    s = s.replace("*", "")
    s = re.sub(r"(?<![A-Za-z0-9])_|_(?![A-Za-z0-9])", "", s)
    s = re.sub(r"\s+", " ", s).strip()
    return re.sub(r"^[\s\"'“”«»>•·\-–—]+", "", s)


def v_description(stub, override=None):
    """(description, what was removed or None, warning or None). Refuses a
    description that is empty or still opens like a voice marker."""
    if override is not None:
        s, removed = override.strip(), "replaced by --v-descriptions"
    else:
        s, removed = clean_for_match(stub), None
        for pat in (MARKER, MARKER_LOOSE, CUE_THAT):
            m = pat.match(s)
            if m:
                removed = (removed + " + " if removed else "") + m.group(0).strip().lower()
                s = s[m.end():]
        if removed is None:
            s = (stub or "").strip()          # nothing removed: the stub exactly as written
        elif s[:1].isalpha():
            s = s[0].upper() + s[1:]
    c = clean_for_match(s)
    if not c:
        raise Refused("the description is empty once the voice cue is removed; "
                      "supply one with --v-descriptions")
    if MARKER.match(c) or MARKER_LOOSE.match(c) or MARKER_LIKE.match(c) or CUE_THAT.match(c):
        raise Refused(f"still opens with a voice cue: {shown(s)[:80]!r}; "
                      "supply a description with --v-descriptions")
    warn = None
    if CUE_OTHER.match(c):
        warn = f"opens with a reporting or interpreting verb: {shown(s)[:60]!r}"
    elif MARKER_ANYWHERE.search(c):
        warn = f"voice cue later in the text: {shown(s)[:60]!r}"
    return s, removed, warn


def text(v):
    """A printed value: control characters removed."""
    if v is None:
        return None
    return ILLEGAL_CHARACTERS_RE.sub("", str(v))


def page_of(loc):
    m = re.search(r"pp?\.\s*(\d+)", loc or "")
    return int(m.group(1)) if m else None


def page_order(frags):
    """Fragment indices in page order; a fragment with no page keeps the page of
    the fragment before it in the list, so it stays beside its list neighbours."""
    pages, last = [], -1
    for f in frags:
        p = page_of(f.get("location"))
        last = p if p is not None else last
        pages.append(last)
    return sorted(range(len(frags)), key=lambda i: (pages[i], i))


def periods_text(log_path, periods, override=None):
    if override:
        d = json.load(open(override, encoding="utf-8"))
        found, src = {p: str(d.get(p, "")).strip() for p in periods}, \
            f"override: {os.path.basename(override)}"
    else:
        found, inside = {}, False
        for line in open(log_path, encoding="utf-8").read().splitlines():
            if re.match(r"^#{2,4}\s*Periods\s*$", line.strip()):
                inside = True
                continue
            if inside and line.startswith("#"):
                break
            if inside:
                m = re.match(r"^\s*[-*]\s*(P\d+)\s*[—–-]+\s*(.+)$", line)
                if m:
                    found[m.group(1)] = m.group(2).strip()
        src = f"log: {os.path.basename(log_path)}, Periods section"
    missing = [p for p in periods if not found.get(p)]
    if missing:
        raise Refused(f"no description for {missing} ({src}); write them to a JSON file and "
                      "pass --periods")
    hits = {p: CLOSURE.findall(found[p]) for p in periods if CLOSURE.search(found[p])}
    if hits:
        raise Refused(f"a period description speaks to closure, which is the answer section 6 "
                      f"asks for ({src}): {hits}. Write neutral descriptions to a JSON file, "
                      "log why, and pass --periods")
    return {p: found[p] for p in periods}, src


def pair_texts(pairs, override=None):
    ov = json.load(open(override, encoding="utf-8")) if override else {}
    out, used = {}, []
    for p in pairs:
        a, b = p["episode_a"], p["episode_b"]
        if p["pair_id"] in ov:
            o = ov[p["pair_id"]]
            if not (isinstance(o, dict) and str(o.get("episode_a", "")).strip()
                    and str(o.get("episode_b", "")).strip()):
                raise Refused(f"--pairs: {p['pair_id']} needs both episode_a and episode_b")
            a, b = str(o["episode_a"]).strip(), str(o["episode_b"]).strip()
            used.append(p["pair_id"])
        hits = PAIRWORDS.findall(a) + PAIRWORDS.findall(b)
        if hits:
            raise Refused(f"{p['pair_id']}: an episode text says what the paper does with the pair "
                          f"({sorted(set(h.lower() for h in hits))}), which is section 5's answer. "
                          "Write neutral episode texts to a JSON file, log why, and pass --pairs")
        out[p["pair_id"]] = (a, b)
    return out, (f"override: {os.path.basename(override)} for {used}" if used else "corpus")


# ---------------------------------------------------------------- workbook
def add_tab(wb, title, headers, widths, rows, first=False, codes=None):
    """rows: printed values, the row id first; the coder's two columns are
    appended, shaded, reason before code (section 1: reason first)."""
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    head = headers + [REASON_H, CODE_H]
    ws.append(head)
    for c in ws[1]:
        c.font, c.fill, c.alignment = HEAD, HEADFILL, WRAP
    for r in rows:
        ws.append([text(v) for v in r] + [None, None])
    for i, w in enumerate(widths + [46, 18], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    n = len(rows)
    for row in ws.iter_rows(min_row=2, max_row=n + 1):
        for c in row:
            c.font, c.alignment = BODY, WRAP
            if isinstance(c.value, str) and c.value.startswith("="):
                c.data_type = "s"          # text, never a formula
        for c in row[-2:]:
            c.fill = INPUT
    if codes and n:
        col = get_column_letter(len(head))
        dv = DataValidation(type="list", formula1='"' + ",".join(codes) + '"',
                            allow_blank=True, showErrorMessage=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{n + 1}")
    ws.freeze_panes = "B2"
    return ws


def add_about(wb, stage_line, citation, tabs):
    ws = wb.create_sheet()
    ws.title = "about this workbook"
    lines = [stage_line, f"Paper: {citation}", "", "Tabs to code:"]
    for title, n, codes, ids in tabs:
        lines.append(f"  {title}: {n} row{'s' if n != 1 else ''}; codes: {codes}; "
                     f"ids: {ids}")
    lines += ["",
              "Your columns are shaded: YOUR REASON and YOUR CODE. Fill every shaded cell; "
              "never leave one blank.",
              "The first column is each row's id. Keep it as it is: it is how your codes "
              "are matched back to the rows.",
              "If you cannot attach the workbook when you have finished, give your codes as "
              f"plain text instead, one row per line, in the order given: {TEXT_FORMAT}. The ids "
              "are the ids printed in that row: fragment, claim, gap or pair, as the list above "
              "says (- for section 6). Write the reason before the code."]
    for l in lines:
        ws.append([text(l)])
    ws.column_dimensions["A"].width = 120
    for row in ws.iter_rows():
        for c in row:
            c.font, c.alignment = BODY, Alignment(wrap_text=True)
    return lines


def row_ids(title, n):
    return [f"{ROW_PREFIX[title]}.{i}" for i in range(1, n + 1)]


def ids_note(title):
    cols = ID_COLUMNS[title]
    return " and ".join(f"the {c} id" for c in cols) if cols else "-"


# ---------------------------------------------------------------- the build
def build(corpus_path, log_path, outdir, periods_override=None, pairs_override=None,
          allow_unfrozen=False, v_override=None, sidecar=None):
    from schema import load, validate
    name = os.path.basename(corpus_path)
    test = not re.match(r"^[a-z0-9_]+_v\d+\.json$", name)
    if test and not allow_unfrozen:
        raise Refused(f"{name} is not a frozen corpus (<paper_id>_v<N>.json). "
                      "Build from the frozen file; --allow-unfrozen makes a test pack only")
    if os.path.exists(outdir) and (not os.path.isdir(outdir) or os.listdir(outdir)):
        raise Refused(f"{outdir} exists and is not an empty folder: a pack is never rebuilt "
                      "over one already built")
    c_obj = load(corpus_path)
    errs, _ = validate(c_obj)
    if errs:
        raise Refused(f"{name} has {len(errs)} validation error(s); fix them first")
    corpus = json.load(open(corpus_path, encoding="utf-8"))
    # the extraction's sidecar, frozen with the corpus: named, checked and recorded
    if not sidecar:
        raise Refused("no sidecar given: pass the extraction's voice file, frozen with the "
                      "corpus, as --sidecar=<voice_gloss.json>. The key records its hash")
    _, _, side_warnings = read_sidecar(sidecar, corpus, c_obj)
    parent = os.path.dirname(os.path.abspath(outdir)) or "."
    os.makedirs(parent, exist_ok=True)
    work = tempfile.mkdtemp(prefix=".pack-", dir=parent)
    try:
        key, sends = _build(corpus, corpus_path, log_path, work, periods_override,
                            pairs_override, test, v_override, sidecar, side_warnings)
        if os.path.exists(outdir):
            os.rmdir(outdir)
        os.replace(work, outdir)
    except BaseException:
        shutil.rmtree(work, ignore_errors=True)
        raise
    kp = os.path.join(outdir, "key", f"{key['paper_id']}_pack_key.json")
    return key, kp, [os.path.join(outdir, os.path.basename(s)) for s in sends]


def _build(corpus, corpus_path, log_path, out, periods_override, pairs_override, test,
           v_override=None, sidecar=None, side_warnings=()):
    meta, pid = corpus["meta"], corpus["meta"]["paper_id"]
    cit = meta.get("citation") or pid
    frags, claims = corpus["fragments"], corpus["claims"]
    fr = {f["fragment_id"]: f for f in frags}
    cl = {c["claim_id"]: c for c in claims}
    fidx = {f["fragment_id"]: i for i, f in enumerate(frags)}
    cidx = {c["claim_id"]: i for i, c in enumerate(claims)}
    attach = os.path.join(out, "attach")
    os.makedirs(attach)
    os.makedirs(os.path.join(out, "key"))
    key = {"paper_id": pid, "test_pack": test, "builder": BUILDER_VERSION,
           "built_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
           "corpus": os.path.basename(corpus_path), "corpus_sha256": sha(corpus_path),
           "sidecar": os.path.basename(sidecar), "sidecar_sha256": sha(sidecar),
           "log": os.path.basename(log_path), "text_format": TEXT_FORMAT,
           "workbooks": {}, "warnings": list(side_warnings)}
    authored = []

    def save(wb, stage, fname, about):
        path = os.path.join(attach, fname)
        wb.save(path)
        authored.extend(about)
        tabs = {}
        for ws in wb.worksheets:
            if ws.title == "about this workbook":
                continue
            rows = list(ws.iter_rows(values_only=True))
            authored.extend(str(h) for h in rows[0])
            tabs[ws.title] = [r[0] for r in rows[1:]]
        key["workbooks"][stage] = {"file": fname, "sha256": sha(path), "rows": tabs}

    # Stage A
    wb = Workbook()
    ids = row_ids("0 outcome", len(claims))
    add_tab(wb, "0 outcome", ["row", "claim", "what the claim asserts", "location"],
            [7, 8, 80, 22],
            [[i, c["claim_id"], c["statement"], c.get("location") or ""]
             for i, c in zip(ids, claims)], first=True)
    about = add_about(wb, "Stage A of 4: the outcome of every claim (section 0 of the "
                          "instructions). Code this stage only.", cit,
                      [("0 outcome", len(claims), FREE_TEXT["0 outcome"], ids_note("0 outcome"))])
    save(wb, "A", f"{pid}_STAGE_A_outcome.xlsx", about)

    # Stage A2: only claims that name an outcome; the corpus's outcome is printed
    named = [c for c in claims if has_outcome(c)]
    wb = Workbook()
    ids = row_ids("0b asserted_direction", len(named))
    add_tab(wb, "0b asserted_direction",
            ["row", "claim", "what the claim asserts", "the outcome recorded for it"],
            [7, 8, 72, 40],
            [[i, c["claim_id"], c["statement"], c["outcome"].strip()] for i, c in zip(ids, named)],
            first=True)
    about = add_about(wb, "Stage A2 of 4: the asserted direction of each claim that names an "
                          "outcome (section 0b). Code this stage only.", cit,
                      [("0b asserted_direction", len(named), FREE_TEXT["0b asserted_direction"],
                        ids_note("0b asserted_direction"))])
    save(wb, "A2", f"{pid}_STAGE_A2_asserted_direction.xlsx", about)

    # Stage V: the location, and the stub with its voice cue removed; nothing else
    removed, vrows = [], []
    vov = json.load(open(v_override, encoding="utf-8")) if v_override else {}
    unknown = sorted(set(vov) - set(fr))
    if unknown:
        raise Refused(f"--v-descriptions names fragments not in the corpus: {unknown}")
    if vov:
        key["V_descriptions_override"] = {"source": os.path.basename(v_override),
                                          "fragments": sorted(vov)}
    ids = row_ids("V voice", len(frags))
    for i, f in zip(ids, frags):
        loc = (f.get("location") or "").strip()
        if not loc:
            raise Refused(f"{f['fragment_id']} (row {i}) has no location; V needs a page for "
                          "every fragment")
        if page_of(loc) is None:
            key["warnings"].append(f"{f['fragment_id']}: location has no page number: "
                                   f"{shown(loc)!r}")
        try:
            desc, rm, warn = v_description(f["stub"], vov.get(f["fragment_id"]))
        except Refused as e:
            raise Refused(f"{f['fragment_id']} (row {i}): {e}")
        if rm:
            removed.append({"row": i, "fragment_id": f["fragment_id"], "removed": rm})
        if warn:
            key["warnings"].append(f"{f['fragment_id']} (row {i}), V description {warn}")
        vrows.append([i, f["fragment_id"], loc, desc])
    wb = Workbook()
    add_tab(wb, "V voice", ["row", "fragment", "location",
                            "a short description of the fragment, to help you find it"],
            [7, 9, 30, 80], vrows, first=True, codes=CODES["V voice"])
    about = add_about(wb, "Stage V of 4: the voice of every fragment, coded from the paper "
                          "(section V). Code this stage only.", cit,
                      [("V voice", len(frags), ", ".join(CODES["V voice"]), ids_note("V voice"))])
    save(wb, "V", f"{pid}_STAGE_V_voice.xlsx", about)
    key["V_opening_removed"] = removed

    # Stage B
    wb = Workbook()
    links = sorted(corpus["claim_links"], key=lambda l: (fidx[l["fragment_id"]],
                                                         cidx[l["claim_id"]]))
    ids = row_ids("1 direction", len(links))

    def odir(c):
        if not has_outcome(c):
            return None, None
        return c["outcome"].strip(), (c.get("asserted_direction") or "").strip() or None

    add_tab(wb, "1 direction",
            ["row", "fragment", "what the fragment says (stub)", "claim",
             "what the claim asserts", "the claim's outcome", "which way the claim says it goes"],
            [7, 9, 56, 8, 52, 30, 40],
            [[i, l["fragment_id"], fr[l["fragment_id"]]["stub"], l["claim_id"],
              cl[l["claim_id"]]["statement"], *odir(cl[l["claim_id"]])]
             for i, l in zip(ids, links)],
            first=True, codes=CODES["1 direction"])
    ids = row_ids("2 core", len(claims))
    add_tab(wb, "2 core", ["row", "claim", "what the claim asserts"], [7, 8, 90],
            [[i, c["claim_id"], c["statement"]] for i, c in zip(ids, claims)],
            codes=CODES["2 core"])
    ids = row_ids("3 stated_as", len(claims))
    add_tab(wb, "3 stated_as", ["row", "claim", "what the claim asserts", "location"],
            [7, 8, 80, 22],
            [[i, c["claim_id"], c["statement"], c.get("location") or ""]
             for i, c in zip(ids, claims)], codes=CODES["3 stated_as"])
    gaps = corpus["declared_gaps"]
    ids = row_ids("4 gap kind", len(gaps))
    add_tab(wb, "4 gap kind", ["row", "gap", "what the paper says about itself", "location"],
            [7, 8, 80, 22],
            [[i, g["gap_id"], g["statement"], g.get("location") or ""]
             for i, g in zip(ids, gaps)], codes=CODES["4 gap kind"])
    pairs = corpus["contrast_pairs"]
    ptext, psrc = pair_texts(pairs, pairs_override)
    key["pair_texts"] = {"source": psrc,
                         "texts": {p: [shown(a), shown(b)] for p, (a, b) in ptext.items()}}

    def pages(ids_):
        locs = [fr[x].get("location") or "" for x in ids_ if x in fr]
        return "; ".join(dict.fromkeys(l for l in locs if l)) or "(no location)"

    ids = row_ids("5 pair drawn_by", len(pairs))
    add_tab(wb, "5 pair drawn_by",
            ["row", "pair", "episode A", "where episode A is", "episode B", "where episode B is"],
            [7, 8, 44, 24, 44, 24],
            [[i, p["pair_id"], ptext[p["pair_id"]][0], pages(p.get("fragments_a") or []),
              ptext[p["pair_id"]][1], pages(p.get("fragments_b") or [])]
             for i, p in zip(ids, pairs)], codes=CODES["5 pair drawn_by"])
    periods = meta.get("periods") or []
    descs, src = periods_text(log_path, periods, periods_override)
    key["period_descriptions"] = {"source": src,
                                  "descriptions": {p: shown(x) for p, x in descs.items()}}
    add_tab(wb, "6 resolution_period",
            ["row", "question", "the periods, in order"], [7, 44, 90],
            [["6.1", "In which period does the paper's story close?",
              "\n".join(f"{p}: {descs[p]}" for p in periods)]],
            codes=periods + ["none"])
    order = page_order(frags)
    pos = {k: n for n, k in enumerate(order)}
    glossed = [k for k in order if (frags[k].get("author_gloss") or "").strip()]
    ids = row_ids("7 gloss_relation", len(glossed))
    grows, neigh = [], {}
    for i, k in zip(ids, glossed):
        n = pos[k]
        before = frags[order[n - 1]] if n > 0 else None
        after = frags[order[n + 1]] if n + 1 < len(order) else None
        cell = lambda x: f"{x['fragment_id']}: {x['stub']}" if x else "(none)"
        f = frags[k]
        grows.append([i, f["fragment_id"], f["stub"], f["author_gloss"], cell(before), cell(after)])
        neigh[f["fragment_id"]] = [x["fragment_id"] if x else None for x in (before, after)]
    add_tab(wb, "7 gloss_relation",
            ["row", "fragment", "what the fragment says (stub)",
             "what the author says about it (gloss)", "the fragment before it on the page",
             "the fragment after it on the page"],
            [7, 9, 48, 48, 40, 40], grows, codes=CODES["7 gloss_relation"])
    key["gloss_neighbours"] = neigh
    tabs = [(t, n, c, ids_note(t)) for t, n, c in [
        ("1 direction", len(links), ", ".join(CODES["1 direction"])),
        ("2 core", len(claims), ", ".join(CODES["2 core"])),
        ("3 stated_as", len(claims), ", ".join(CODES["3 stated_as"])),
        ("4 gap kind", len(gaps), ", ".join(CODES["4 gap kind"])),
        ("5 pair drawn_by", len(pairs), ", ".join(CODES["5 pair drawn_by"])),
        ("6 resolution_period", 1, ", ".join(periods + ["none"])),
        ("7 gloss_relation", len(glossed), ", ".join(CODES["7 gloss_relation"]))]]
    about = add_about(wb, "Stage B of 4: sections 1 to 7 of the instructions, one tab each. "
                          "Code this stage only.", cit, tabs)
    save(wb, "B", f"{pid}_STAGE_B.xlsx", about)

    # the opening message for each stage's chat
    sends = []
    for n, (stage, label) in enumerate([("A", "Stage A"), ("A2", "Stage A2"),
                                        ("V", "Stage V"), ("B", "Stage B")], start=1):
        fname = key["workbooks"][stage]["file"]
        msg = (f"I've attached three files: CODING_INSTRUCTIONS.md, a PDF of a published "
               f"research paper, and a workbook for {label}.\n\n"
               f"Read the instructions in full first, then the whole paper, then the workbook. "
               f"Then code {label}, and only {label}, as the instructions describe: fill every "
               f"shaded cell.\n\n"
               f"When you have finished, give me the workbook back with your columns filled. "
               f"If you cannot attach a file, give me your codes as plain text instead, one row "
               f"per line, in the workbook's order: {TEXT_FORMAT}. The ids are the ids printed in "
               f"that row (fragment, claim, gap or pair; - for section 6), and the reason comes "
               f"before the code. If you cannot finish in one reply, stop at the end of a "
               f"row and wait for me to type \"continue\".\n\n"
               f"Work only from the attached PDF and workbook. Do not search the web, and set "
               f"aside anything you know about this paper or its author from elsewhere.")
        authored.append(msg)
        head = ("# TEST PACK: DO NOT SEND. Built from an unfrozen corpus.\n\n" if test else "")
        body = (f"{head}# {pid}: coder 2, {label}\n\n"
                f"Send {n} of 4. Send it only after the previous stage's workbook is back.\n\n"
                f"**A fresh incognito claude.ai chat**, outside any project; Opus 5.5 at the same "
                f"effort setting for every stage and paper (record it in the hours log); web "
                f"search, research and every connector off.\n\n"
                f"**Attach three files and nothing else:** `CODING_INSTRUCTIONS.md`, the paper's "
                f"PDF, and `attach/{fname}`. Never attach anything from `key/`.\n\n"
                f"**Paste this message as it is:**\n\n```\n{msg}\n```\n\n"
                f"**Then** type `continue` and nothing else, whenever it stops. If it asks a "
                f"question, reply: \"Please follow the instructions: make your best judgement "
                f"and say what made it hard in the reason column.\" Never answer what a code is "
                f"for.\n\n**Bring back** the returned workbook, or all of its replies' text in "
                f"one file, with the chat's start and end times. `intake_coder2.py` checks "
                f"either against the workbook as sent.\n")
        sp = os.path.join(out, f"SEND_{n}_STAGE_{stage}.md")
        open(sp, "w", encoding="utf-8").write(body)
        sends.append(sp)

    # checks on what was written
    mach = machinery_patterns()
    bad = machinery_named(authored, mach)
    if bad:
        raise Refused(f"builder-authored text names the machinery: {bad}")
    hits = []
    for stage, info in key["workbooks"].items():
        wbx = load_workbook(os.path.join(attach, info["file"]), read_only=True)
        for ws in wbx.worksheets:
            for r in ws.iter_rows(min_row=2, values_only=True):
                for v in r:
                    for p in mach:
                        if v and p.search(str(v)):
                            hits.append((stage, ws.title, r[0], p.pattern))
    if hits:
        key["warnings"].append(f"paper content matches a machinery word: {hits}")
    kp = os.path.join(out, "key", f"{pid}_pack_key.json")
    json.dump(key, open(kp, "w", encoding="utf-8"), indent=2, ensure_ascii=False)
    return key, sends


def report(key, kp, sends):
    w = key["workbooks"]
    if key["test_pack"]:
        print("TEST PACK: built from an unfrozen corpus. DO NOT SEND.")
    print(f"{key['paper_id']}: built from {key['corpus']} ({key['corpus_sha256'][:12]}), with "
          f"the sidecar {key['sidecar']} ({key['sidecar_sha256'][:12]}) recorded")
    for stage in ("A", "A2", "V", "B"):
        tabs = ", ".join(f"{t} {len(r)}" for t, r in w[stage]["rows"].items())
        print(f"  Stage {stage:<3} attach/{w[stage]['file']}: {tabs}")
    print(f"  V: opening voice cue removed from {len(key['V_opening_removed'])} stub(s)")
    print(f"  periods from {key['period_descriptions']['source']}; pairs from "
          f"{key['pair_texts']['source']}")
    print(f"  wrote {os.path.relpath(kp, os.path.dirname(os.path.dirname(kp)))} and "
          + ", ".join(os.path.basename(s) for s in sends))
    for x in key["warnings"]:
        print(f"  WARNING: {x}")
    print("  Send in order A, A2, V, B, each after the last is back.")


# ---------------------------------------------------------------- self-test
def selftest():
    """Left out of the repository's copy: the self-test ran on the Carlile pilot's files, which are not in the
    repository (see the README)."""
    print("selftest: left out of the repository's copy; it runs on the Carlile pilot's files, which are not "
          "here (see the README)")
    return 2


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    if "--selftest" in sys.argv:
        raise SystemExit(selftest())
    if len(args) != 3:
        raise SystemExit(__doc__)
    opt = {a.split("=", 1)[0]: a.split("=", 1)[1] for a in sys.argv[1:]
           if a.startswith("--") and "=" in a}
    try:
        report(*build(args[0], args[1], args[2], opt.get("--periods"), opt.get("--pairs"),
                      "--allow-unfrozen" in sys.argv, opt.get("--v-descriptions"),
                      opt.get("--sidecar")))
    except Refused as e:
        print(f"REFUSED: {e}")
        raise SystemExit(1)
