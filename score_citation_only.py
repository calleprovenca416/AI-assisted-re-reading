"""
Score the citation-only run (T36), by CITATION_ONLY_PLAN.md. Written on 16 Oct 2026,
before any reply of the run was opened.

    python3 -B score_citation_only.py selftest
    python3 -B score_citation_only.py score --key=citation_only_key.json
        --blank=citation_matching_sheet.xlsx --sheet=citation_matching_sheet_filled.xlsx
        --out=DIR REPLY_58 REPLY_59 REPLY_60 REPLY_61

score: the plan's measures, written to DIR/CITATION_ONLY_RESULTS.md and
DIR/citation_only_results.json. First the checks, each of which stops the run and
says why if it fails: the blank sheet is the one the key names (its sha256); each
reply is the one the key names, in the order of chats 58 to 61, and counts by
build_citation_sheet.py's check; no wager and no what_you_know, the strings
written out here, holds a quoted span of more than eight words (as
build_citation_sheet.py defines a span); matching_results.json and
stage2/markers/KEY.json are the files the key names, and the key's groups and
arms are theirs; the PI's sheet has the blank's sheets, and every cell but the
yellow ones (Match and Note on the sheets Bechky and Pratt, in the readings'
rows) is as on the blank; and every Match cell reads. On Bechky a match reads if,
its ends trimmed and capitals aside, it is one of the key's ten labels or none.
On Pratt it reads if it is none, capitals aside, or a list of readings of P01 to
P15, separated by commas, semicolons, white space or "and", none named twice.

The measures, for each paper, for each model (Fable 5.1 at max, GPT 6 Astra at
max) and for both together:

- M1: the readings made without the paper; how many the PI places with readings
  made with it (Bechky: in one of the PI's groups; Pratt: the same point as at least
  one of P01 to P15); how many the PI marks none, a new point.
- M2, Bechky: the PI's groups reached, each named by at least one reading made
  without the paper, by each model and by either; and of the groups found twice
  with the paper (those holding readings of two or more of chats 23, 26, 27 and
  46 by matching_results.json, which must be the plan's five: accreditation,
  capacity, different demands, firearms already changed and staff sorting), how
  many and which are reached.
- M3, Pratt: the readings of P01 to P15 named by at least one reading made
  without the paper, by arm (system, middle, plain, as KEY.json gives them), by
  each model and by either.
- M4: each chat's what_you_know, as it wrote it.

Also each reading made without the paper, with its paper, chat, model and wager,
and the PI's match and note. Nothing is tested: the plan fixes no bar.
"""
import hashlib
import json
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import build_citation_sheet as bcs   # noqa: E402

PAPERS = bcs.PAPERS
MODELS = ("Fable 5.1 at max", "GPT 6 Astra at max")
SHORT = {"Fable 5.1 at max": "Fable 5.1", "GPT 6 Astra at max": "GPT 6 Astra"}
EARLIER = ("23", "26", "27", "46")
FIVE = ("accreditation", "capacity", "different demands", "firearms already changed", "staff sorting")
ARMS = ("system", "middle", "plain")
P_IDS = tuple(f"P{i:02d}" for i in range(1, 16))
ANSWER_SHEETS = {"Bechky": "bechky", "Pratt": "pratt"}
MATCHING_RESULTS = bcs.MATCHING_RESULTS
KEY_11 = bcs.KEY_11
RESULTS_MD = "CITATION_ONLY_RESULTS.md"
RESULTS_JSON = "citation_only_results.json"


def sha(path):
    return hashlib.sha256(open(path, "rb").read()).hexdigest()


# ---------------------------------------------------------------- reading the matches

def read_bechky(v, labels):
    """(the group's label, or None for none; fault or None)."""
    s = "" if v is None else str(v).strip()
    if not s:
        return None, "no match"
    if s.casefold() == "none":
        return None, None
    if s.casefold() in labels:
        return labels[s.casefold()], None
    return None, f"{s!r} is not one of the ten labels or none"


def read_pratt(v):
    """(a list of P ids, empty for none; fault or None)."""
    s = "" if v is None else str(v).strip()
    if not s:
        return [], "no match"
    if s.casefold() == "none":
        return [], None
    toks = [t for t in re.split(r"[,;\s]+", s) if t and t.casefold() != "and"]
    out = []
    for t in toks:
        if t.upper() not in P_IDS:
            return [], f"{t!r} is not one of P01 to P15"
        if t.upper() in out:
            return [], f"{t.upper()} is named twice"
        out.append(t.upper())
    if not out:
        return [], f"{s!r} names no reading"
    return out, None


def read_sheet(filled, blank, key):
    """(faults, {x: {"match": ..., "note": ...}})."""
    import openpyxl
    wf, wb = openpyxl.load_workbook(filled), openpyxl.load_workbook(blank)
    if wf.sheetnames != wb.sheetnames:
        return [f"sheets {wf.sheetnames} are not {wb.sheetnames}"], None
    xs = {s: [e["x"] for e in key["readings"] if e["paper"] == p] for s, p in ANSWER_SHEETS.items()}
    faults = []
    for name in wb.sheetnames:
        a, b = wf[name], wb[name]
        last = 1 + len(xs.get(name, []))
        for r in range(1, max(a.max_row, b.max_row) + 1):
            for c in range(1, max(a.max_column, b.max_column) + 1):
                if name in xs and 2 <= r <= last and c in (3, 4):
                    continue
                if a.cell(r, c).value != b.cell(r, c).value:
                    faults.append(f"{name}!{a.cell(r, c).coordinate} differs from the blank sheet")
    labels = {g["label"].casefold(): g["label"] for g in key["groups"]}
    answers = {}
    for name, want in xs.items():
        ws = wf[name]
        if [ws.cell(r, 1).value for r in range(2, 2 + len(want))] != want:
            faults.append(f"{name}: the readings are not the key's, in order")
            continue
        for r, x in enumerate(want, start=2):
            v = ws.cell(r, 3).value
            m, f = read_bechky(v, labels) if name == "Bechky" else read_pratt(v)
            if f:
                faults.append(f"{name}!C{r} ({x}): {f}")
            note = ws.cell(r, 4).value
            answers[x] = {"match": m, "note": None if note is None else str(note)}
    return faults, (None if faults else answers)


# ---------------------------------------------------------------- the measures

def found_twice(grouping, chat_of):
    groups = {}
    for i, g in grouping.items():
        groups.setdefault(g, []).append(i)
    return sorted(g for g, m in groups.items() if len({chat_of[i] for i in m if chat_of[i] in EARLIER}) >= 2)


def measures(key, answers, arms, twice):
    """M1 to M3 from the key's readings, the PI's answers, P01 to P15's arms and the groups found twice."""
    rows = [dict(e, match=answers[e["x"]]["match"], note=answers[e["x"]]["note"]) for e in key["readings"]]
    order = [g["label"] for g in key["groups"]]
    of_cf = {g.casefold(): g for g in order}
    m1 = {}
    for paper in PAPERS:
        m1[paper] = {}
        for m in MODELS + ("both",):
            sel = [r for r in rows if r["paper"] == paper and (m == "both" or r["model"] == m)]
            placed = sum(1 for r in sel if r["match"])
            m1[paper][m] = {"readings": len(sel), "placed": placed, "new": len(sel) - placed}
    rb = [r for r in rows if r["paper"] == "bechky"]
    reached = {m: [g for g in order if any(r["match"] == g for r in rb if r["model"] == m)] for m in MODELS}
    reached["either"] = [g for g in order if any(g in reached[m] for m in MODELS)]
    five = [of_cf[g] for g in twice]
    m2 = {"groups": len(order),
          "reached": reached,
          "readings_of": {g: [r["x"] for r in rb if r["match"] == g] for g in order},
          "found_twice": five,
          "found_twice_reached": {m: [g for g in five if g in reached[m]] for m in reached}}
    rp = [r for r in rows if r["paper"] == "pratt"]
    named = {p: [r["x"] for r in rp if p in r["match"]] for p in P_IDS}
    counter = {m: [p for p in P_IDS if any(p in r["match"] for r in rp if r["model"] == m)] for m in MODELS}
    counter["either"] = [p for p in P_IDS if named[p]]
    by_arm = {}
    for arm in ARMS + ("all",):
        ps = [p for p in P_IDS if arm == "all" or arms[p] == arm]
        by_arm[arm] = {"readings": len(ps), **{m: sum(1 for p in counter[m] if p in ps) for m in counter}}
    m3 = {"named_by": named, "counterpart": counter, "by_arm": by_arm}
    return rows, m1, m2, m3


# ---------------------------------------------------------------- the run

def score(opts, replies):
    key = json.load(open(opts["key"], encoding="utf-8"))
    stop = []
    if sha(opts["blank"]) != key["sheet_sha256"]:
        stop.append("the blank sheet is not the one the key names")
    chats = sorted(bcs.CHATS)
    if len(replies) != 4:
        stop.append("four replies are needed, in the order of chats 58 to 61")
    wyk = {}
    for chat, path in zip(chats, replies):
        if sha(path) != key["sources"][chat]["sha256"]:
            stop.append(f"chat {chat}: {os.path.basename(path)} is not the reply the key names")
            continue
        rep, faults = bcs.load(path)
        faults = faults or bcs.faults_of(rep)
        if faults:
            stop.append(f"chat {chat}: the reply does not count: {'; '.join(faults)}")
            continue
        w = rep["what_you_know"]
        for s in [w if isinstance(w, str) else json.dumps(w, ensure_ascii=False)] + [r["wager"]
                                                                                      for r in rep["readings"]]:
            if any(n > bcs.MAX_QUOTE for n, _ in bcs.quoted_spans(s)):
                stop.append(f"chat {chat}: a wager or what_you_know holds a quoted span of more than eight words")
                break
        wyk[chat] = w
    for f in (MATCHING_RESULTS, KEY_11):
        if sha(os.path.join(HERE, f)) != key["inputs"][f]:
            stop.append(f"{f} is not the file the key names")
    if stop:
        raise SystemExit("not scored:\n  " + "\n  ".join(stop))
    res = json.load(open(os.path.join(HERE, MATCHING_RESULTS), encoding="utf-8"))
    k11 = json.load(open(os.path.join(HERE, KEY_11), encoding="utf-8"))["papers"]["pratt"]
    arms = {e["reading"]: e["arm"] for e in k11}
    if arms != {e["reading"]: e["arm"] for e in key["pratt_readings"]} or sorted(arms) != list(P_IDS):
        raise SystemExit("not scored: the key's arms are not KEY.json's")
    groups_cf = {i: g["label"].casefold() for g in key["groups"] for i in g["readings"]}
    if groups_cf != res["groupings"]["P"]:
        raise SystemExit("not scored: the key's groups are not those of matching_results.json")
    twice = found_twice(res["groupings"]["P"], res["chat_of"])
    if tuple(twice) != FIVE:
        raise SystemExit(f"not scored: the groups found twice are {twice}, not the plan's five")
    faults, answers = read_sheet(opts["sheet"], opts["blank"], key)
    if faults:
        raise SystemExit("not scored: the PI's sheet:\n  " + "\n  ".join(faults))
    rows, m1, m2, m3 = measures(key, answers, arms, twice)
    out = {"what": "the measures of the citation-only run (T36), by CITATION_ONLY_PLAN.md",
           "scorer_sha256": sha(os.path.abspath(__file__)),
           "files": {"key": [os.path.basename(opts["key"]), sha(opts["key"])],
                     "blank": [os.path.basename(opts["blank"]), sha(opts["blank"])],
                     "sheet": [os.path.basename(opts["sheet"]), sha(opts["sheet"])],
                     "replies": {c: [os.path.basename(p), sha(p)] for c, p in zip(chats, replies)},
                     MATCHING_RESULTS: key["inputs"][MATCHING_RESULTS], KEY_11: key["inputs"][KEY_11]},
           "checks": {"blank": "the key's", "replies": "the key's, and all four count",
                      "quoted_spans": "none of more than eight words in a wager or what_you_know",
                      "sheet": "no cell outside the yellow ones changed; every match reads"},
           "arms": arms, "rows": rows, "m1": m1, "m2": m2, "m3": m3,
           "m4": {c: {"paper": key["sources"][c]["paper"], "model": key["sources"][c]["model"],
                      "what_you_know": wyk[c]} for c in chats}}
    os.makedirs(opts["out"], exist_ok=True)
    open(os.path.join(opts["out"], RESULTS_JSON), "w", encoding="utf-8").write(
        json.dumps(out, indent=1, ensure_ascii=False) + "\n")
    open(os.path.join(opts["out"], RESULTS_MD), "w", encoding="utf-8").write(write_md(out, key))
    print(f"{RESULTS_MD}: sha256 {sha(os.path.join(opts['out'], RESULTS_MD))}")
    print(f"{RESULTS_JSON}: sha256 {sha(os.path.join(opts['out'], RESULTS_JSON))}")


def cell(s):
    return "" if s is None else str(s).replace("|", "\\|").replace("\n", " ")


def write_md(o, key):
    chat_of = {(key["sources"][c]["paper"], key["sources"][c]["model"]): c for c in key["sources"]}
    f = o["files"]
    L = ["# The citation-only run (T36): results", "",
         f"Written by `score_citation_only.py score` (sha256 opens {o['scorer_sha256'][:8]}) by "
         f"`CITATION_ONLY_PLAN.md`, from the key `{f['key'][0]}` (opens {f['key'][1][:8]}), the PI's sheet "
         f"`{f['sheet'][0]}` (opens {f['sheet'][1][:8]}), checked against the blank `{f['blank'][0]}` (opens "
         f"{f['blank'][1][:8]}), and the four replies: "
         + ", ".join(f"`{n}` (opens {h[:8]})" for n, h in f["replies"].values()) + ".", "",
         "Described, not tested: the plan fixes no bar. The PI ran the chats and is not blind to the source of what "
         "the PI matched; the sheet did not say which chat wrote which reading. One run per model and paper. A point "
         "reached without the paper is available without the paper's text, from memory of the paper or of writing "
         "about it, or by inference from the citation and the field; the run cannot say which. Pratt's readings "
         "are matched to single readings, Bechky's to the PI's groups, so the two papers' figures are not alike in "
         "kind.", "",
         "**Checks:** the blank sheet and the four replies are the key's, and all four replies count; no wager or "
         "what_you_know holds a quoted span of more than eight words; on the PI's sheet no cell outside the yellow "
         "ones changed, and every match reads.", "",
         "## M1. The readings made without the paper", "",
         "| paper | model | readings | placed with readings made with the paper | marked none, a new point |",
         "|---|---|---|---|---|"]
    for paper in PAPERS:
        for m in MODELS + ("both",):
            v = o["m1"][paper][m]
            who = "both" if m == "both" else f"{SHORT[m]} (chat {chat_of[(paper, m)]})"
            L.append(f"| {paper.capitalize()} | {who} | {v['readings']} | {v['placed']} | {v['new']} |")
    m2 = o["m2"]
    L += ["", "## M2. Bechky: the PI's groups of 14 Oct reached without the paper", "",
          "| the PI's group | readings made with the paper | found twice with the paper | reached by "
          + " | reached by ".join(SHORT[m] for m in MODELS) + " |",
          "|---|---|---|---|---|"]
    rows_b = [r for r in o["rows"] if r["paper"] == "bechky"]
    for g in key["groups"]:
        lab = g["label"]
        by = [", ".join(r["x"] for r in rows_b if r["match"] == lab and r["model"] == m) or "—" for m in MODELS]
        L.append(f"| {cell(lab)} | {len(g['readings'])} | {'yes' if lab in m2['found_twice'] else ''} | "
                 + " | ".join(by) + " |")
    L += [""]
    for m in MODELS + ("either",):
        who = "either model" if m == "either" else SHORT[m]
        rr, ft = m2["reached"][m], m2["found_twice_reached"][m]
        L.append(f"- **Reached by {who}:** {len(rr)} of {m2['groups']} groups"
                 + (f" ({', '.join(rr)})" if rr else "") + f"; of the {len(m2['found_twice'])} found twice with the "
                 f"paper, {len(ft)}" + (f" ({', '.join(ft)})" if ft else "") + ".")
    m3 = o["m3"]
    L += ["", "## M3. Pratt: the readings made with the paper that have a counterpart without it", "",
          "| arm | readings made with the paper | with a counterpart from "
          + " | from ".join(SHORT[m] for m in MODELS) + " | from either |",
          "|---|---|---|---|---|"]
    for arm in ARMS + ("all",):
        v = m3["by_arm"][arm]
        L.append(f"| {arm} | {v['readings']} | " + " | ".join(str(v[m]) for m in MODELS) + f" | {v['either']} |")
    L += [""]
    rows_p = {r["x"]: r for r in o["rows"] if r["paper"] == "pratt"}
    named = [(p, xs) for p, xs in m3["named_by"].items() if xs]
    if named:
        L.append("Named: " + "; ".join(
            f"{p} ({o['arms'][p]}) by " + ", ".join(f"{x} ({SHORT[rows_p[x]['model']]})" for x in xs)
            for p, xs in named) + ".")
    else:
        L.append("Named: none.")
    L += ["", "## M4. What each model says it knows of each paper", ""]
    for c, v in o["m4"].items():
        w = v["what_you_know"]
        w = w if isinstance(w, str) else json.dumps(w, ensure_ascii=False)
        L += [f"**Chat {c}, {v['paper'].capitalize()}, {SHORT[v['model']]}:**", ""]
        L += [("> " + line) if line.strip() else ">" for line in (w.strip() or "(empty)").split("\n")]
        L.append("")
    L += ["## The readings made without the paper, and the PI's matches", ""]
    for r in o["rows"]:
        match = (r["match"] or "none") if r["paper"] == "bechky" else (", ".join(r["match"]) or "none")
        L += [f"**{r['x']}** ({r['paper'].capitalize()}, chat {r['chat']}, {SHORT[r['model']]}, its reading "
              f"{r['place']}, `reading_id` {cell(r['reading_id'])}). **Match:** {match}."
              + (f" **Note:** {cell(r['note'])}" if r["note"] else ""), "", cell(r["wager"]), ""]
    return "\n".join(L)


# ---------------------------------------------------------------- self-test

def selftest():
    bad = []

    def eq(name, got, want):
        if got != want:
            bad.append(f"{name}: got {got!r}, want {want!r}")

    labels = {"capacity": "Capacity", "accreditation": "Accreditation", "status": "Status"}
    eq("b1", read_bechky(" capacity ", labels), ("Capacity", None))
    eq("b2", read_bechky("NONE", labels), (None, None))
    eq("b3", read_bechky(None, labels), (None, "no match"))
    eq("b4", read_bechky("Capacity, Status", labels)[1], "'Capacity, Status' is not one of the ten labels or none")
    eq("b5", read_bechky(3, labels)[1], "'3' is not one of the ten labels or none")
    eq("b6", read_bechky("ACCREDITATION", labels), ("Accreditation", None))
    eq("p1", read_pratt("p03, P11"), (["P03", "P11"], None))
    eq("p2", read_pratt("P03 and P11;P15"), (["P03", "P11", "P15"], None))
    eq("p3", read_pratt(" None "), ([], None))
    eq("p4", read_pratt("P3")[1], "'P3' is not one of P01 to P15")
    eq("p5", read_pratt("P16")[1], "'P16' is not one of P01 to P15")
    eq("p6", read_pratt("P03, p03")[1], "P03 is named twice")
    eq("p7", read_pratt("none, P03")[1], "'none' is not one of P01 to P15")
    eq("p8", read_pratt("")[1], "no match")
    eq("p9", read_pratt(" and ")[1], "'and' names no reading")
    eq("t1", found_twice({"a": "x", "b": "x", "c": "y", "d": "y", "e": "z", "f": "z"},
                         {"a": "23", "b": "26", "c": "46", "d": "46", "e": "49", "f": "27"}), ["x"])

    F, A = MODELS
    key = {"groups": [{"label": "Capacity", "readings": ["M01"]}, {"label": "Accreditation", "readings": ["M02"]},
                      {"label": "Status", "readings": ["M03"]}],
           "readings": [{"x": "X01", "paper": "bechky", "model": F}, {"x": "X02", "paper": "bechky", "model": A},
                        {"x": "X03", "paper": "bechky", "model": F}, {"x": "X04", "paper": "bechky", "model": A},
                        {"x": "X05", "paper": "bechky", "model": F}, {"x": "X06", "paper": "pratt", "model": F},
                        {"x": "X07", "paper": "pratt", "model": A}, {"x": "X08", "paper": "pratt", "model": A},
                        {"x": "X09", "paper": "pratt", "model": F}]}
    answers = {"X01": "Capacity", "X02": None, "X03": "Accreditation", "X04": "Capacity", "X05": "Status",
               "X06": ["P01", "P04"], "X07": [], "X08": ["P04", "P02"], "X09": []}
    answers = {x: {"match": m, "note": None} for x, m in answers.items()}
    arms = {p: "plain" for p in P_IDS}
    arms.update({"P02": "middle", "P04": "system"})
    rows, m1, m2, m3 = measures(key, answers, arms, ["accreditation", "capacity"])
    eq("m1b", m1["bechky"], {F: {"readings": 3, "placed": 3, "new": 0}, A: {"readings": 2, "placed": 1, "new": 1},
                             "both": {"readings": 5, "placed": 4, "new": 1}})
    eq("m1p", m1["pratt"], {F: {"readings": 2, "placed": 1, "new": 1}, A: {"readings": 2, "placed": 1, "new": 1},
                            "both": {"readings": 4, "placed": 2, "new": 2}})
    eq("m2r", m2["reached"], {F: ["Capacity", "Accreditation", "Status"], A: ["Capacity"],
                              "either": ["Capacity", "Accreditation", "Status"]})
    eq("m2f", m2["found_twice"], ["Accreditation", "Capacity"])
    eq("m2fr", m2["found_twice_reached"], {F: ["Accreditation", "Capacity"], A: ["Capacity"],
                                           "either": ["Accreditation", "Capacity"]})
    eq("m2o", m2["readings_of"], {"Capacity": ["X01", "X04"], "Accreditation": ["X03"], "Status": ["X05"]})
    eq("m3c", m3["counterpart"], {F: ["P01", "P04"], A: ["P02", "P04"], "either": ["P01", "P02", "P04"]})
    eq("m3a", m3["by_arm"], {"system": {"readings": 1, F: 1, A: 1, "either": 1},
                             "middle": {"readings": 1, F: 0, A: 1, "either": 1},
                             "plain": {"readings": 13, F: 1, A: 0, "either": 1},
                             "all": {"readings": 15, F: 2, A: 2, "either": 3}})
    eq("m3n", {p: x for p, x in m3["named_by"].items() if x}, {"P01": ["X06"], "P02": ["X08"],
                                                              "P04": ["X06", "X08"]})
    eq("rows", [r["match"] for r in rows][:2], ["Capacity", None])
    if bad:
        print("SELFTEST FAILED:\n  " + "\n  ".join(bad))
        return 1
    print("selftest: all checks pass")
    return 0


def main(argv):
    opts, files = {}, []
    for a in argv[1:]:
        if a.startswith("--"):
            k, _, v = a[2:].partition("=")
            opts[k] = v
        else:
            files.append(a)
    mode = argv[0] if argv else None
    if mode == "selftest" and not opts and not files:
        sys.exit(selftest())
    elif mode == "score" and {"key", "blank", "sheet", "out"} <= set(opts) and len(files) == 4:
        score(opts, files)
    else:
        raise SystemExit(__doc__)


if __name__ == "__main__":
    main(sys.argv[1:])
