"""
Build a voice sidecar -- the `<paper_id>_voice_gloss.json` format that
EXTRACTION_BRIEF.md defines -- from a returned coding workbook.

WHY
---
audit.py applies the voice rule (27 Sep 2026) only when it is given
a sidecar. An extraction session writes its own. Carlile v4.4 has none: its
`voice` was coded once from the page, by the second coder in Stage B (25 Sep),
and that coding lives in the returned workbook. This turns the workbook's voice
and gloss-relation tabs into a sidecar, so v4.4 can be audited under the rule.
The same script turns a cold paper's returned stage V workbook into a sidecar
for comparison with the extraction's own.

On a cold paper the two fields come back in two workbooks: voice in Stage V's,
gloss relation in Stage B's. So the Stage B workbook can be given as a fourth
argument (added 7 Oct 2026), and the sidecar then carries coder 2's codes for
both.

USAGE
    python3 make_voice_sidecar.py <workbook.xlsx> <paper_id> <out.json> [<stage_B.xlsx>]

Reads the tab whose title ends in "voice" and, where present, the one ending in
"gloss_relation"; takes the YOUR CODE and YOUR REASON columns. Refuses if a voice
row is blank or carries a code outside the two values.

With a fourth argument, gloss relation is read from that workbook's tab. The
first workbook must then have no gloss tab of its own, the second must have
one, and every row of it must carry a code: the script refuses otherwise.

Carlile v4.4:
    python3 make_voice_sidecar.py returned_carlile_STAGE_B_coder2.xlsx \
        carlile_2002 carlile_2002_v4_4_voice_gloss.json

A cold paper, coder 2:
    python3 make_voice_sidecar.py returned_<paper_id>_STAGE_V_coder2.xlsx \
        <paper_id> <out.json> returned_<paper_id>_STAGE_B_coder2.xlsx

Two refusals added on 7 Oct 2026:
- It never writes over a file. It once could, and in a scratch copy it
  replaced an extraction's own sidecar without a word. To make a sidecar
  again, remove or rename the old one first.
- With a fourth argument, both workbooks must carry the record that
  intake_coder2.py writes, for this paper, as Stage V and as Stage B, with no
  cell changed since. So a Stage V of one paper cannot be joined to a Stage B
  of another.
And one more after that review's second reading: with one workbook, a cold
paper's Stage V (its voice tab is titled `V voice`) must carry that record
too. Carlile's workbook, whose voice tab is `8 voice`, never passed that
intake and is read as before.
"""
import json
import os
import sys

from openpyxl import load_workbook

VOICE = ("reported", "author_insight")
GLOSS = ("consistent", "goes_beyond", "diverges", "none")


def read_tab(ws):
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h or "") for h in rows[0]]
    code = next(i for i, h in enumerate(header) if h.startswith("YOUR CODE"))
    reason = next(i for i, h in enumerate(header) if h.startswith("YOUR REASON"))
    frag = header.index("fragment")
    out = []
    for r in rows[1:]:
        if r is None or r[frag] is None:
            continue
        out.append({"fragment_id": str(r[frag]).strip(),
                    "code": ("" if r[code] is None else str(r[code])).strip().strip("`").lower(),
                    "reason": ("" if r[reason] is None else str(r[reason])).strip()})
    return out


def main(path, paper_id, out_path, gloss_path=None):
    if os.path.exists(out_path):
        raise SystemExit(f"REFUSED: {out_path} exists, and this script never writes over a file")
    if gloss_path is not None:
        sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
        from intake_coder2 import stamp_problem
        for p, stage in ((path, "V"), (gloss_path, "B")):
            why = stamp_problem(p, paper_id, stage)
            if why:
                raise SystemExit(f"REFUSED: {why}")
    wb = load_workbook(path, read_only=True, data_only=True)
    tabs = {ws.title: ws for ws in wb.worksheets}
    vt = [t for t in tabs if t.lower().endswith("voice")]
    gt = [t for t in tabs if t.lower().endswith("gloss_relation")]
    if len(vt) != 1:
        raise SystemExit(f"expected one voice tab, found {vt}")
    if gloss_path is None and vt[0] == "V voice":       # a cold paper's Stage V, given alone
        sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
        from intake_coder2 import stamp_problem
        why = stamp_problem(path, paper_id, "V")
        if why:
            raise SystemExit(f"REFUSED: {why}")
    voice = read_tab(tabs[vt[0]])
    bad = [v["fragment_id"] for v in voice if v["code"] not in VOICE]
    if bad:
        raise SystemExit(f"REFUSED: blank or unknown voice code on {', '.join(bad)}")
    if gloss_path is None:
        gloss = [g for g in read_tab(tabs[gt[0]]) if g["code"]] if gt else []
        source = (f"{path}: tab '{vt[0]}'" + (f" and tab '{gt[0]}'" if gt else "")
                  + ", YOUR CODE and YOUR REASON columns, as returned")
    else:
        if gt:
            raise SystemExit(f"REFUSED: {path} has a gloss tab of its own ({gt}); give one "
                             "workbook, or a second one for gloss relation, not both")
        gtabs = {ws.title: ws for ws in
                 load_workbook(gloss_path, read_only=True, data_only=True).worksheets}
        gt = [t for t in gtabs if t.lower().endswith("gloss_relation")]
        if len(gt) != 1:
            raise SystemExit(f"expected one gloss_relation tab in {gloss_path}, found {gt}")
        gloss = read_tab(gtabs[gt[0]])
        blank = [g["fragment_id"] for g in gloss if not g["code"]]
        if blank:
            raise SystemExit(f"REFUSED: blank gloss_relation code on {', '.join(blank)}")
        source = (f"{path}: tab '{vt[0]}'; {gloss_path}: tab '{gt[0]}'; "
                  "YOUR CODE and YOUR REASON columns, as returned")
    bad = [g["fragment_id"] for g in gloss if g["code"] not in GLOSS]
    if bad:
        raise SystemExit(f"REFUSED: unknown gloss_relation code on {', '.join(bad)}")
    side = {"paper_id": paper_id, "source": source, "voice": voice, "gloss_relation": gloss}
    with open(out_path, "w") as fh:
        json.dump(side, fh, indent=2, ensure_ascii=False)
        fh.write("\n")
    n_ai = sum(v["code"] == "author_insight" for v in voice)
    print(f"wrote {out_path}: voice for {len(voice)} fragments "
          f"({n_ai} author_insight), gloss_relation for {len(gloss)}")


if __name__ == "__main__":
    if len(sys.argv) not in (4, 5):
        raise SystemExit(__doc__)
    main(*sys.argv[1:])
