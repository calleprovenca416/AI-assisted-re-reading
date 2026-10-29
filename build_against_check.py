"""
The `against` check (AGAINST_CHECK.md): build the
one table the check session receives, and read its answers back.
Written 28 Sep 2026.

WHEN: once, after both codings of all four papers are in, before any
disagreement is discussed.

THE TABLE (as AGAINST_CHECK.md has it):
- every link either coder coded `against`, in any paper, each fragment-claim
  pair once;
- each row shows the claim's statement, its outcome and its asserted direction,
  and the fragment's stub, all from the frozen corpus;
- never the paper, the coder, the link's basis or the gloss;
- rows shuffled with a recorded seed, ids random (R01, R02, ...);
- an `against` link on a claim with no outcome breaks the rule both coders
  work under; it is listed in the key, separately, and never sent.

Coder 1's codes come from the frozen corpus; coder 2's from the returned
Stage B workbook's "1 direction" tab.

USAGE
    python3 build_against_check.py build <outdir> --seed=<int> \\
        <paper_v1.json>=<returned_paper_STAGE_B_coder2.xlsx> [more pairs ...]
    python3 build_against_check.py read <outdir>/against_check_key.json <answers.txt>
    python3 build_against_check.py --selftest

build writes against_check_table.md (the table, sent with AGAINST_CHECK.md and
nothing else) and against_check_key.json (for no chat), into a folder
that holds neither yet: a table already sent is never rebuilt over. read checks
that the table is the one the key records, then the answers, one per line as
AGAINST_CHECK.md asks: row id | reason | answer, with "| unclear" optionally at
the end; answers are 1, 2, 3, 4a, 4b, 4c, 4x. It writes
against_check_results.json beside the key, listing per paper and per coder the
links flagged 4a-4x: in the second version of every result that rests on
`against`, those links count as `illustrates`. The links held back because
their claim has no outcome are carried into the results as well.

Revised 28 Sep: the reason now comes before the
answer, as it does for the coders; an existing key is never overwritten; the
table's hash is checked when answers are read.

Revised 7 Oct 2026: a returned Stage B workbook must carry the record
that intake_coder2.py writes, for its paper, as Stage B, with no cell changed
since. Before this the table could be built from a workbook that agreement.py
refuses, and the check session would have been spent before anyone knew.
"""

import hashlib
import json
import os
import random
import re
import sys
import tempfile

from openpyxl import load_workbook

ANSWERS = ("1", "2", "3", "4a", "4b", "4c", "4x")
ROW_AT = re.compile(r"(?<![\w.])(R\d+)[`*\s]{0,4}\|")
HEADER_LINE = re.compile(r"^[|\s`*]*(row\b|id\b|#\s*[|`*]|[-:|\s`]+$)", re.I)


def sha(path):
    return hashlib.sha256(open(path, "rb").read()).hexdigest()


def coder2_against(stage_b):
    wb = load_workbook(stage_b, read_only=True, data_only=True)
    ws = wb["1 direction"]
    rows = list(ws.iter_rows(values_only=True))
    head = [str(h or "") for h in rows[0]]
    code = head.index("YOUR CODE")
    frag, claim = head.index("fragment"), head.index("claim")
    return {(str(r[frag]), str(r[claim])) for r in rows[1:]
            if r and str(r[code] or "").strip().strip("`").lower() == "against"}


def cell(s):
    return " ".join(str(s or "").split()).replace("|", "\\|")


def build(outdir, seed, pairs):
    os.makedirs(outdir, exist_ok=True)
    for f in ("against_check_table.md", "against_check_key.json"):
        if os.path.exists(os.path.join(outdir, f)):
            raise SystemExit(f"REFUSED: {outdir} already holds {f}; a check table is never rebuilt "
                             "over one that may have been sent")
    rows, broken, sources = [], [], []
    for corpus_path, stage_b in pairs:
        c = json.load(open(corpus_path, encoding="utf-8"))
        pid = c["meta"]["paper_id"]
        from build_cold_pack import has_outcome
        from intake_coder2 import stamp_problem
        why = stamp_problem(stage_b, pid, "B")
        if why:
            raise SystemExit(f"REFUSED: {why}")
        fr = {f["fragment_id"]: f for f in c["fragments"]}
        cl = {x["claim_id"]: x for x in c["claims"]}
        one = {(l["fragment_id"], l["claim_id"]) for l in c["claim_links"]
               if (l.get("direction") or "").lower() == "against"}
        two = coder2_against(stage_b)
        unknown = [k for k in two if k[0] not in fr or k[1] not in cl]
        if unknown:
            raise SystemExit(f"{stage_b}: links not in {corpus_path}: {unknown}")
        sources.append({"paper_id": pid, "corpus": os.path.basename(corpus_path),
                        "corpus_sha256": sha(corpus_path), "stage_b": os.path.basename(stage_b),
                        "stage_b_sha256": sha(stage_b), "coder1_against": len(one),
                        "coder2_against": len(two)})
        for f_id, c_id in sorted(one | two):
            who = [n for n, s in (("coder1", one), ("coder2", two)) if (f_id, c_id) in s]
            claim = cl[c_id]
            item = {"paper_id": pid, "fragment_id": f_id, "claim_id": c_id, "coded_by": who}
            if not has_outcome(claim):
                broken.append(item)
                continue
            item["_show"] = [claim["statement"], claim["outcome"],
                             claim.get("asserted_direction") or "", fr[f_id]["stub"]]
            rows.append(item)
    rng = random.Random(seed)
    rng.shuffle(rows)
    width = max(2, len(str(len(rows))))
    lines = ["| row | the claim | its outcome | which way the claim says it goes | the stub |",
             "|---|---|---|---|---|"]
    for n, item in enumerate(rows, 1):
        item["row"] = f"R{n:0{width}d}"
        lines.append(f"| {item['row']} | " + " | ".join(cell(x) for x in item.pop("_show")) + " |")
    table = os.path.join(outdir, "against_check_table.md")
    open(table, "w", encoding="utf-8").write("\n".join(lines) + "\n")
    key = {"seed": seed, "sources": sources, "rows": rows,
           "not_sent_claim_has_no_outcome": broken, "table_sha256": sha(table)}
    kp = os.path.join(outdir, "against_check_key.json")
    json.dump(key, open(kp, "w", encoding="utf-8"), indent=2, ensure_ascii=False)
    print(f"wrote {table}: {len(rows)} row(s), seed {seed}")
    print(f"wrote {kp} (for no chat)")
    for s in sources:
        print(f"  {s['paper_id']}: coder 1 {s['coder1_against']}, coder 2 {s['coder2_against']} "
              f"`against` link(s)")
    if broken:
        print(f"  not sent, claim has no outcome: {len(broken)} "
              + ", ".join(f"{b['paper_id']} {b['fragment_id']}->{b['claim_id']}" for b in broken))
    return key, table, kp


def read(key_path, answers_path):
    key = json.load(open(key_path, encoding="utf-8"))
    table = os.path.join(os.path.dirname(key_path), "against_check_table.md")
    if not os.path.exists(table) or sha(table) != key["table_sha256"]:
        print("REFUSED: the table beside this key is not the one the key records")
        return 1
    ids = {r["row"]: r for r in key["rows"]}
    got, errors, superseded = {}, [], []
    for n, raw in enumerate(open(answers_path, encoding="utf-8").read().splitlines(), 1):
        line = re.sub(r"^[\s|]+", "", raw)
        m = ROW_AT.search(line)
        if not (m and "|" not in line[:m.start()]):
            if "|" in line and not HEADER_LINE.match(line):
                errors.append(f"line {n} not read as row id | reason | answer: {raw.strip()[:80]}")
            continue
        parts = [p.strip().strip("`*").strip() for p in line[m.start():].strip().strip("|").split("|")]
        unclear = re.sub(r"[^a-z]", "", parts[-1].lower()) == "unclear"
        if unclear:
            parts = parts[:-1]
        if len(parts) < 3:
            errors.append(f"line {n} not read as row id | reason | answer: {raw.strip()[:80]}")
            continue
        rid, ans = m.group(1), parts[-1].lower().strip(" .`*")
        reason = " | ".join(parts[1:-1]).strip()
        if rid not in ids:
            errors.append(f"unknown row {rid} (line {n})")
        elif ans not in ANSWERS:
            errors.append(f"{rid}: answer `{parts[-1]}` is not one of {', '.join(ANSWERS)} (line {n})")
        else:
            if rid in got and got[rid]["answer"] != ans:
                superseded.append({"row": rid, **got[rid]})
            got[rid] = {"answer": ans, "reason": reason, "unclear": unclear, "line": n}
    missing = sorted(set(ids) - set(got))
    if missing:
        errors.append(f"unanswered: {missing}")
    if errors:
        for e in errors:
            print("  ", e)
        print("REFUSED: nothing written")
        return 1
    results = []
    for rid, r in ids.items():
        results.append({**r, **got[rid], "flagged": got[rid]["answer"].startswith("4")})
    flagged = {}
    for r in results:
        if r["flagged"]:
            for who in r["coded_by"]:
                flagged.setdefault(r["paper_id"], {}).setdefault(who, []).append(
                    [r["fragment_id"], r["claim_id"]])
    out = os.path.join(os.path.dirname(key_path), "against_check_results.json")
    if superseded:
        print(f"  CHECK BY HAND: {len(superseded)} row(s) answered twice with different answers; "
              f"the last is kept, the earlier are in the results file: "
              f"{[x['row'] for x in superseded]}")
    json.dump({"key": os.path.basename(key_path), "answers": os.path.basename(answers_path),
               "results": results, "flagged_as_illustrates": flagged,
               "superseded_answers": superseded,
               "not_sent_claim_has_no_outcome": key["not_sent_claim_has_no_outcome"],
               "rule": "in the second version of every result resting on `against`, P2 included, "
                       "a link answered 4a-4x counts as `illustrates`"},
              open(out, "w", encoding="utf-8"), indent=2, ensure_ascii=False)
    counts = {a: sum(r["answer"] == a for r in results) for a in ANSWERS}
    print(f"wrote {out}: " + ", ".join(f"{a} {n}" for a, n in counts.items() if n)
          + f"; unclear {sum(r['unclear'] for r in results)}")
    return 0


def selftest():
    """Left out of the repository's copy: the self-test ran on the Carlile pilot's files, which are not in the
    repository (see the README)."""
    print("selftest: left out of the repository's copy; it runs on the Carlile pilot's files, which are not "
          "here (see the README)")
    return 2


if __name__ == "__main__":
    a = sys.argv[1:]
    if a[:1] == ["--selftest"]:
        raise SystemExit(selftest())
    if a[:1] == ["read"] and len(a) == 3:
        raise SystemExit(read(a[1], a[2]))
    if a[:1] == ["build"] and len(a) >= 3:
        seed = next((int(x.split("=", 1)[1]) for x in a if x.startswith("--seed=")), None)
        pairs = [tuple(x.split("=", 1)) for x in a[2:] if not x.startswith("--")]
        if seed is None or not pairs or any(len(p) != 2 for p in pairs):
            raise SystemExit(__doc__)
        build(a[1], seed, pairs)
        raise SystemExit(0)
    raise SystemExit(__doc__)
