"""
Build the files of the matching test (T33), by MATCHING_TEST_PLAN.md. Written on
14 Oct 2026, on the wager test's
`build_marker_packs.py` and `build_worth_sheet.py` and T32's `build_places_packs.py`.

    python3 -B build_matching_pack.py reference --out=DIR
    python3 -B build_matching_pack.py pack --seed=20261014 --out=DIR REPLY_49 REPLY_50

reference writes REFERENCE_bechky.md, for arm B's coders: Bechky's 23 claims,
numbered, each with its statement and the pages the record gives it, and its 172
passages, numbered, each with its page, the unit it concerns, its voice and its
summary (the stub), from bechky_2020_v1.json and its voice sidecar,
bechky_2020_v1_voice_gloss.json (coder 1's codes, as frozen). Nothing else of the
record goes in: no links, pairs, gaps, absences, `core`, `stated_as`, glosses,
speakers or other fields. Voice is printed as "reported" or "author's account"
(the record's author_insight).

pack writes, from the two replies of chats 49 and 50 and the earlier readings:

- READINGS_matching.md, for the coders of both arms: every reading of the test
  as one paragraph, shuffled with the seed and numbered M01 onward.
- matching_key.json: which reading came from which chat; for nobody.
- grouping_sheet.xlsx: the PI's grouping sheet: the rule, the made-up example
  and the readings in the same order and words, with a yellow cell for a group
  label and one for a note.
- worth_sheet_matching.xlsx: the PI's worth sheet for the readings of chats 49
  and 50 only, on the wager test's scale and layout, for after the PI's grouping.

The readings, in this order before the shuffle:

- B01 to B14 of stage2/markers/KEY.json (Bechky), their sentences joined with a
  space: chat 23 for the system's cards, 26 for reply_26, 27 for reply_27.
- W01, W03, W05 and W06 of places_packs_key.json, the readings of chat 46,
  their sentences joined with a space (the edits of places_pack_edits.json are
  already in them).
- The readings of reply 49, then of reply 50, in their order: each put in the
  plain arm's form of 11 Oct (build_places_packs.reading_text: the wager, then
  each piece of evidence's job with its page in parentheses at the end of the
  job's first sentence, pages made alike), then split into sentences with the
  wager test's splitter and joined again with a space.

Chat 45's readings are left out, as the plan says. A reply counts if it is one
JSON object (from its first "{" to its last "}") with a list `readings` of one
to five items, each with a `reading_id`, a non-empty `wager` and a list
`evidence` of items with a `page` and a `job`; otherwise the build stops and
says why. Both workbooks carry fixed dates inside, so two builds give the same
bytes.
"""
import datetime
import hashlib
import json
import os
import random
import re
import sys
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import build_marker_packs as bmp     # noqa: E402
import build_places_packs as bpp     # noqa: E402
import build_worth_sheet as bws      # noqa: E402
from openpyxl import Workbook         # noqa: E402

RECORD = "bechky_2020_v1.json"
VOICE = "bechky_2020_v1_voice_gloss.json"
KEY_11 = "stage2/markers/KEY.json"
KEY_46 = "places_packs_key.json"
CHAT_OF_SOURCE = {
    "stage2/baseline_replies/reply_26_bechky_reading.json": "26",
    "stage2/baseline_replies/reply_27_bechky_plain.json": "27",
}
W_OF_46 = ["W01", "W03", "W05", "W06"]
VOICE_WORDS = {"reported": "reported", "author_insight": "author's account"}

RULE = ("Two readings make the same point when they contest the same claim of the paper with the same "
        "alternative account, whatever their wording, emphasis or evidence. Group each reading by its main "
        "point, the one its wager states first. A reading whose point no other reading makes is a group of its "
        "own. Every reading is in exactly one group.")
EXAMPLE = [
    ("The paper", "A made-up paper on restaurant kitchens claims that cooks under time pressure skip the tasting "
                  "step because they value speed over care (claim 1), and that head chefs who taste every dish "
                  "themselves keep the step alive (claim 2)."),
    ("R1", "The step is skipped because the recipe cards go missing in the rush, not because cooks value speed."),
    ("R2", "What the paper calls valuing speed is the absence of the cards: without them, no one knows what a dish "
           "should taste like."),
    ("R3", "Cooks skip the step in the rush because the head chef forbids tasting at the pass."),
    ("R4", "The head chefs who taste everything are those who arrive before the rush begins, so the step survives "
           "only where the evening starts calm."),
    ("Groups", "R1 and R2 contest claim 1 with the same alternative, the missing cards: one label for both. R3 "
               "contests claim 1 with another alternative, the head chef's rule: a label of its own. R4 contests "
               "claim 2: a label of its own. So: R1 a, R2 a, R3 b, R4 c."),
]
GROUP_RULES = [
    "Group every reading by its main point, the one its wager states first.",
    "Give readings that make the same point the same label, and a reading whose point no other reading makes a "
    "label of its own. Any short label will do: a, b, c, or a few words. Labels that differ only in capitals "
    "count as the same.",
    "Group by the point, not by worth, and each reading on its own words, whoever wrote it.",
    "The readings are shuffled, and nothing on the sheet shows where each came from. Please do not open "
    "matching_key.json, and do not run or read any of chats 51 to 54, until your grouping is sent.",
    "You may consult the paper. A page in a reading is a pointer, not a proof.",
    "Save the workbook as grouping_sheet_filled.xlsx, with no other change, and send it in our conversation or "
    "put it in replies/.",
]
WORTH_RULES = [
    "Rate worth only: whether the reading, as an idea about the paper's theory, opens something worth a "
    "scholar's time.",
    "Whether its facts hold is not checked in this test. If you know a premise to be false, rate the reading as "
    "written and say so in the note.",
    "Rate each reading on its own, not against the others, and rate every one.",
    "These are the readings of chats 49 and 50, which you ran. Please rate them only after your grouping sheet "
    "is sent.",
    "You may consult the paper. A page in a reading is a pointer, not a proof.",
    "Save the workbook as worth_sheet_matching_filled.xlsx, with no other change, and send it in our "
    "conversation or put it in replies/.",
]
STAMP = (2026, 10, 14, 0, 0, 0)


def here(name):
    return os.path.join(HERE, name)


def sha(path):
    return hashlib.sha256(open(path, "rb").read()).hexdigest()


def steady(path):
    stamp = b"2026-10-14T00:00:00Z"
    with zipfile.ZipFile(path) as z:
        items = [(i, z.read(i.filename)) for i in z.infolist()]
    with zipfile.ZipFile(path, "w") as z:
        for i, data in items:
            if i.filename == "docProps/core.xml":
                data, n = re.subn(rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)",
                                  rb"\g<1>" + stamp + rb"\g<2>", data)
                if n != 1:
                    raise SystemExit("docProps/core.xml: no single modified date")
            j = zipfile.ZipInfo(i.filename, date_time=STAMP)
            j.compress_type = i.compress_type
            j.external_attr = i.external_attr
            z.writestr(j, data)


def fixed_dates(wb, creator):
    fixed = datetime.datetime(*STAMP)
    wb.properties.created = fixed
    wb.properties.modified = fixed
    wb.properties.creator = creator


def build_reference(out):
    rec = json.load(open(here(RECORD), encoding="utf-8"))
    side = json.load(open(here(VOICE), encoding="utf-8"))
    voice = {v["fragment_id"]: v["code"] for v in side["voice"]}
    claims, frags = rec["claims"], rec["fragments"]
    if len(claims) != 23 or len(frags) != 172:
        raise SystemExit(f"expected 23 claims and 172 passages, found {len(claims)} and {len(frags)}")
    missing = [f["fragment_id"] for f in frags if voice.get(f["fragment_id"]) not in VOICE_WORDS]
    if missing:
        raise SystemExit(f"passages without a voice code: {missing}")
    lines = [f"# A reference for {rec['meta']['citation']}", "",
             "Made from the paper before any of the readings was written. It lists the paper's claims, numbered,",
             "each with its statement and the pages where the paper makes it, and the passages of the paper's",
             "evidence, numbered, each with its page, the group or unit it concerns, its voice and a one-sentence",
             "summary. Voice: `reported`, the passage reports what happened, what was done or what someone said;",
             "`author's account`, the passage is the author's own account, across events. A summary is a pointer,",
             "not the passage: check the page.", "",
             f"## The paper's claims ({len(claims)})", ""]
    for c in claims:
        lines += [f"**{c['claim_id']}** ({c['location']}). {c['statement'].strip()}", ""]
    lines += [f"## The passages of the paper's evidence ({len(frags)})", ""]
    for f in frags:
        lines += [f"**{f['fragment_id']}** · {f['location']} · {f['unit']} · {VOICE_WORDS[voice[f['fragment_id']]]}. "
                  f"{f['stub'].strip()}", ""]
    text = "\n".join(lines)
    os.makedirs(out, exist_ok=True)
    open(os.path.join(out, "REFERENCE_bechky.md"), "w", encoding="utf-8").write(text)
    print(f"REFERENCE_bechky.md: {len(claims)} claims, {len(frags)} passages, "
          f"sha256 {hashlib.sha256(text.encode()).hexdigest()[:16]}")


def check_reply(path, chat):
    rep = bpp.load_reply(path)
    rs = rep.get("readings") if isinstance(rep, dict) else None
    if not isinstance(rs, list) or not 1 <= len(rs) <= 5:
        raise SystemExit(f"chat {chat}: `readings` is not a list of one to five items")
    for r in rs:
        if not isinstance(r, dict) or not str(r.get("reading_id") or "").strip():
            raise SystemExit(f"chat {chat}: a reading has no reading_id")
        if not isinstance(r.get("wager"), str) or not r["wager"].strip():
            raise SystemExit(f"chat {chat}, {r.get('reading_id')}: no wager")
        ev = r.get("evidence")
        if not isinstance(ev, list) or any(not isinstance(e, dict) or not isinstance(e.get("page"), str)
                                           or not isinstance(e.get("job"), str) or not e["job"].strip()
                                           for e in ev):
            raise SystemExit(f"chat {chat}, {r.get('reading_id')}: evidence is not a list of page and job")
    ids = [r["reading_id"] for r in rs]
    if len(set(ids)) != len(ids):
        raise SystemExit(f"chat {chat}: a reading_id is repeated")
    return rep


def earlier_readings():
    out = []
    k11 = json.load(open(here(KEY_11), encoding="utf-8"))["papers"]["bechky"]
    if [e["reading"] for e in k11] != [f"B{i:02d}" for i in range(1, 15)]:
        raise SystemExit("KEY.json: Bechky's readings are not B01 to B14 in order")
    for e in k11:
        if e["arm"] == "system":
            if not e["source"].startswith("stage2/bechky_2020_cards/"):
                raise SystemExit(f"{e['reading']}: a system reading from {e['source']}")
            chat = "23"
        else:
            chat = CHAT_OF_SOURCE[e["source"]]
        out.append({"chat": chat, "model": "Fable 5.1 at max", "earlier_id": e["reading"],
                    "source": e["source"], "source_id": e["source_id"],
                    "text": " ".join(u["text"] for u in e["units"])})
    k46 = json.load(open(here(KEY_46), encoding="utf-8"))
    by_id = {e["reading"]: e for e in k46["readings"]}
    for w in W_OF_46:
        e = by_id[w]
        if e["chat"] != "46":
            raise SystemExit(f"{w} is not chat 46's")
        out.append({"chat": "46", "model": "Fable 5.1 at max", "earlier_id": w,
                    "source": "reply_46_bechky_open.json", "source_id": e["source_id"],
                    "text": " ".join(u["text"] for u in e["units"])})
    if sorted(e["chat"] for e in k46["readings"] if e["chat"] == "46") != ["46"] * 4:
        raise SystemExit("places_packs_key.json: chat 46 does not have four readings")
    return out


def readings_sheet(wb, title, entries, third, fourth, widths=(10, 100, 16, 45), validate=None):
    from openpyxl.styles import Alignment
    ws = wb.create_sheet(title)
    for col, w in zip("ABCD", widths):
        ws.column_dimensions[col].width = w
    for i, h in enumerate(["Reading", "The reading", third, fourth]):
        bws.cell(ws, f"{'ABCD'[i]}1", h, bold=True, fill=bws.GREY).border = bws.BOX
    if validate is not None:
        ws.add_data_validation(validate)
    for r, e in enumerate(entries, start=2):
        bws.cell(ws, f"A{r}", e["reading"], bold=True).border = bws.BOX
        bws.cell(ws, f"B{r}", e["text"]).border = bws.BOX
        c = bws.cell(ws, f"C{r}", None, fill=bws.YELLOW)
        c.border = bws.BOX
        c.alignment = Alignment(horizontal="center", vertical="top")
        bws.cell(ws, f"D{r}", None, fill=bws.YELLOW).border = bws.BOX
        if validate is not None:
            validate.add(f"C{r}")
        ws.row_dimensions[r].height = bws.height(e["text"], 112)
    ws.freeze_panes = "B2"
    bws.printable(ws)


def instructions(wb, rows):
    ins = wb.active
    ins.title = "Instructions"
    ins.column_dimensions["A"].width = 16
    ins.column_dimensions["B"].width = 110
    r = 1
    for label, t, bold, size in rows:
        bws.cell(ins, f"A{r}", label, bold=bold, size=size, wrap=r != 1)
        if t is not None:
            bws.cell(ins, f"B{r}", t, bold=False, size=size)
            ins.row_dimensions[r].height = bws.height(t, 120)
        r += 1
    bws.printable(ins)


def build_pack(out, seed, reply_49, reply_50):
    rec = json.load(open(here(RECORD), encoding="utf-8"))
    rs = earlier_readings()
    sources = {}
    for chat, path in (("49", reply_49), ("50", reply_50)):
        rep = check_reply(path, chat)
        sources[chat] = {"file": os.path.basename(path), "sha256": sha(path), "paper": rep.get("paper")}
        for r in rep["readings"]:
            text = " ".join(bmp.sentences(bpp.reading_text(r)))
            rs.append({"chat": chat, "model": "GPT 6 Astra at max", "earlier_id": None,
                       "source": os.path.basename(path), "source_id": r["reading_id"], "text": text})
    rng = random.Random(seed)
    rng.shuffle(rs)
    for i, r in enumerate(rs, 1):
        r["reading"] = f"M{i:02d}"
    n = len(rs)
    lines = [f"# Readings of {rec['meta']['citation']}", "",
             f"{n} readings of this paper, in no particular order, numbered M01 to M{n:02d}. A page in a reading",
             "is a pointer, not a proof.", ""]
    for r in rs:
        lines += [f"## Reading {r['reading']}", "", r["text"], ""]
    text = "\n".join(lines)
    os.makedirs(out, exist_ok=True)
    open(os.path.join(out, "READINGS_matching.md"), "w", encoding="utf-8").write(text)
    key = {"what": "the key to READINGS_matching.md, grouping_sheet.xlsx and worth_sheet_matching.xlsx: for nobody",
           "seed": seed, "sources": sources,
           "earlier": {"KEY.json": sha(here(KEY_11)), "places_packs_key.json": sha(here(KEY_46))},
           "readings_sha256": hashlib.sha256(text.encode()).hexdigest(),
           "readings": [{"reading": r["reading"], "chat": r["chat"], "model": r["model"],
                         "earlier_id": r["earlier_id"], "source": r["source"], "source_id": r["source_id"],
                         "text": r["text"]} for r in rs]}
    open(os.path.join(out, "matching_key.json"), "w", encoding="utf-8").write(
        json.dumps(key, indent=1, ensure_ascii=False) + "\n")

    wb = Workbook()
    rows = [("Grouping the readings by the point they make", None, True, 13),
            ("What", f"Each reading proposes another way of reading the paper's own account: it contests something "
                     f"the paper claims and puts something else in its place. Group the {n} readings, on the sheet "
                     f"Bechky, by the point each makes.", True, 10),
            ("Fill in", "Only the yellow cells: Group (a short label) and Note. Leave every other cell as it is.",
             True, 10),
            ("The rule", RULE, True, 10),
            ("Example", "From a made-up paper, not one of the readings:", True, 10)]
    rows += [(a, b, False, 10) for a, b in EXAMPLE]
    rows.append(("Rules", None, True, 10))
    rows += [(str(i), t, False, 10) for i, t in enumerate(GROUP_RULES, 1)]
    instructions(wb, rows)
    readings_sheet(wb, "Bechky", rs, "Group", "Note", widths=(10, 100, 18, 45))
    fixed_dates(wb, "build_matching_pack.py")
    p = os.path.join(out, "grouping_sheet.xlsx")
    wb.save(p)
    steady(p)

    new = [r for r in rs if r["chat"] in ("49", "50")]
    wb = Workbook()
    rows = [("Rating the worth of the new readings", None, True, 13),
            ("What", f"Each reading proposes another way of reading the paper's own account. Give each a worth from "
                     f"0 to 3, and a note where it helps. The {len(new)} readings of chats 49 and 50 are on the sheet "
                     f"Bechky, numbered as on your grouping sheet.", True, 10),
            ("Fill in", "Only the yellow cells: Worth (a whole number, 0 to 3; the cell takes nothing else) and Note. "
                        "Leave every other cell as it is.", True, 10),
            ("Scale", None, True, 10)]
    rows += [(str(v), d, False, 10) for v, d in bws.SCALE]
    rows.append(("Rules", None, True, 10))
    rows += [(str(i), t, False, 10) for i, t in enumerate(WORTH_RULES, 1)]
    instructions(wb, rows)
    from openpyxl.worksheet.datavalidation import DataValidation
    dv = DataValidation(type="whole", operator="between", formula1="0", formula2="3", allow_blank=True,
                        showErrorMessage=True, errorTitle="Worth", error="Enter a whole number from 0 to 3.",
                        showInputMessage=True, promptTitle="Worth", prompt="0, 1, 2 or 3: see Instructions.")
    readings_sheet(wb, "Bechky", new, "Worth (0 to 3)", "Note", widths=(10, 100, 14, 45), validate=dv)
    fixed_dates(wb, "build_matching_pack.py")
    p = os.path.join(out, "worth_sheet_matching.xlsx")
    wb.save(p)
    steady(p)
    by_chat = {}
    for r in rs:
        by_chat[r["chat"]] = by_chat.get(r["chat"], 0) + 1
    print(f"{n} readings {dict(sorted(by_chat.items()))}; READINGS_matching.md {key['readings_sha256'][:16]}; "
          f"{len(new)} on the worth sheet")


def main(argv):
    opts, files = {}, []
    for a in argv[1:]:
        if a.startswith("--"):
            k, _, v = a[2:].partition("=")
            opts[k] = v
        else:
            files.append(a)
    mode = argv[0] if argv else None
    if mode == "reference" and "out" in opts and not files:
        build_reference(opts["out"])
    elif mode == "pack" and {"out", "seed"} <= set(opts) and len(files) == 2:
        build_pack(opts["out"], int(opts["seed"]), files[0], files[1])
    else:
        raise SystemExit(__doc__)


if __name__ == "__main__":
    main(sys.argv[1:])
