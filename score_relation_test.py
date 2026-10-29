"""
Score the relation test (T37), by RELATION_TEST_PLAN.md. Written on 16 Oct 2026,
before any reply of the test was opened.

    python3 -B score_relation_test.py selftest
    python3 -B score_relation_test.py score --key=relation_key.json --sheet=relation_sheet_filled.xlsx
        --check=planted_check_filled.xlsx --out=DIR REPLY_64 REPLY_65 REPLY_66 REPLY_67

The replies in the order of the chats: 64, Bechky, Gemini 3.1 Pro; 65, Bechky, GPT 6 Astra; 66, Pratt, Gemini;
67, Pratt, Astra.

score: the plan's measures, written to DIR/RELATION_RESULTS.md and DIR/relation_results.json. First the checks,
each of which stops the run and says why if it fails: the key was made by this folder's build_relation_test.py
(its sha256) from the inputs as they are in this folder; the readings files and the two blank sheets beside the
key are the ones the key names; each reply counts by build_relation_test.py's check against its paper's readings
and claims, and holds no quoted span of more than eight words (as filed, the replies hold none); the PI's coding
sheet has the blank's sheets and no cell changed but Main claim, Main type and Note in the readings' rows, and
each of the PI's twelve codes reads as a coder's would (a claim of the list or none, one of the six types or none,
none in both or neither); the PI's check of the planted readings changes no cell but Of this type? and Note, and each
answer is yes or no (ends trimmed, capitals aside).

The measures, by the plan, on the real readings unless said:

- M1: for each paper and both together, raw agreement and Cohen's kappa (agreement.py's) between the two
  coders' main types, in the six types with none and in the four levels with none, each with its table; for both
  papers together a 95% interval by agreement.py's boot_interval, 2,000 resamples with the seed 20261023.
- M2: how many readings the coders give the same main claim, and of those how many the same main type.
- M3: for each planted reading, each coder's code: caught (its main claim and main type are the planted ones),
  caught aside (the planted claim among its other claims with the planted type), or missed; counted over the
  planted readings the PI confirmed, the others shown and left out of the count.
- M4: on the PI's twelve, each coder against the PI: the same main claim, the same main type, the same level, each
  counted on its own; and the two coders against each other on the same twelve.
- M5: each coder's counts of types and levels; where the two agree on the level, the PI's worth by level (count and
  mean), worth taken from the PI's sheets of 11 Oct (B01 to B14, P01 to P15), of T32 (chat 46's) and of T33 (chats 49
  and 50's) as matching_worth.py takes it; and each coder's levels by the chat that wrote the reading (Bechky) and
  by arm (Pratt).

Nothing is tested: the plan fixes no bar.
"""
import hashlib
import json
import os
import sys
from collections import Counter

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import agreement as A               # noqa: E402
import build_citation_sheet as bcs  # noqa: E402
import build_relation_test as brt   # noqa: E402

PAPERS = brt.PAPERS
CHATS = (("64", "bechky", "gemini"), ("65", "bechky", "astra"), ("66", "pratt", "gemini"), ("67", "pratt", "astra"))
CODERS = ("gemini", "astra")
CODER_NAME = {"gemini": "Gemini 3.1 Pro", "astra": "GPT 6 Astra at max"}
TYPES6 = brt.TYPES + ("none",)
LEVELS = ("inverts", "dismantles", "complements", "restates", "none")
SCOPES = (("bechky", "Bechky"), ("pratt", "Pratt"), ("both", "both papers"))
SEED_BOOT = 20261023
REPS = 2000
WORTH_T29 = "stage2/worth/worth_sheet_filled_2.xlsx"
WORTH_T32 = "worth_sheet_places_filled.xlsx"
WORTH_T33 = "worth_sheet_matching_filled.xlsx"
RESULTS_MD = "RELATION_RESULTS.md"
RESULTS_JSON = "relation_results.json"


def sha(path):
    return hashlib.sha256(open(path, "rb").read()).hexdigest()


def here(name):
    return os.path.join(HERE, name)


def level(t):
    return brt.LEVEL[t]


# ---------------------------------------------------------------- the sheets

def compare_cells(filled, blank, free):
    """(faults, the filled workbook): faults where the filled workbook differs from the blank outside `free`,
    {sheet: (rows, cols)}."""
    import openpyxl
    wf, wb = openpyxl.load_workbook(filled), openpyxl.load_workbook(blank)
    if wf.sheetnames != wb.sheetnames:
        return [f"sheets {wf.sheetnames} are not {wb.sheetnames}"], None
    faults = []
    for name in wb.sheetnames:
        a, b = wf[name], wb[name]
        rows, cols = free.get(name, ((), ()))
        for r in range(1, max(a.max_row, b.max_row) + 1):
            for c in range(1, max(a.max_column, b.max_column) + 1):
                if r in rows and c in cols:
                    continue
                if a.cell(r, c).value != b.cell(r, c).value:
                    faults.append(f"{name}!{a.cell(r, c).coordinate} differs from the blank sheet")
    return faults, wf


def read_pi_sheet(filled, blank, key):
    """(faults, remarks, {(paper, R id): {main_claim, main_type, note}})."""
    free = {brt.NAME[p]: (range(2, 2 + brt.SAMPLE), (3, 4, 5)) for p in PAPERS}
    faults, wf = compare_cells(filled, blank, free)
    if wf is None:
        return faults, [], None
    codes, remarks = {}, []
    for p in PAPERS:
        ws = wf[brt.NAME[p]]
        if [ws.cell(r, 1).value for r in range(2, 2 + brt.SAMPLE)] != key["sample"][p]:
            faults.append(f"{brt.NAME[p]}: the readings are not the key's sample, in order")
            continue
        ids = key["claims"][p]
        for r in range(2, 2 + brt.SAMPLE):
            rid = ws.cell(r, 1).value
            raw_c, raw_t = ws.cell(r, 3).value, ws.cell(r, 4).value
            mc, loose = brt.norm_claim(raw_c, ids)
            mt = brt.norm_type(raw_t, allow_none=True)
            if mc is None or mt is None or (mc == "none") != (mt == "none"):
                faults.append(f"{brt.NAME[p]} {rid}: main claim {raw_c!r} and main type {raw_t!r} do not read")
                continue
            if loose:
                remarks.append(f"{brt.NAME[p]} {rid}: main claim {raw_c!r} read as {mc}")
            if raw_t != mt:
                remarks.append(f"{brt.NAME[p]} {rid}: main type {raw_t!r} read as {mt}")
            note = ws.cell(r, 5).value
            codes[(p, rid)] = {"main_claim": mc, "main_type": mt, "note": None if note is None else str(note)}
    return faults, remarks, (None if faults else codes)


def read_check(filled, blank, key):
    """(faults, {(paper, R id): {confirmed, note}})."""
    n = len(key["planted"])
    faults, wf = compare_cells(filled, blank, {"Planted": (range(2, 2 + n), (5, 6))})
    if wf is None:
        return faults, None
    ws = wf["Planted"]
    out = {}
    for r in range(2, 2 + n):
        label = str(ws.cell(r, 1).value)
        paper, rid = label.split()[0].casefold(), label.split()[1]
        v = "" if ws.cell(r, 5).value is None else str(ws.cell(r, 5).value).strip().casefold()
        if v not in ("yes", "no"):
            faults.append(f"Planted!E{r} ({label}): {ws.cell(r, 5).value!r} is not yes or no")
            continue
        note = ws.cell(r, 6).value
        out[(paper, rid)] = {"confirmed": v == "yes", "note": None if note is None else str(note)}
    if not faults and sorted(out) != sorted((x["paper"], x["r"]) for x in key["planted"]):
        faults.append("the planted readings on the sheet are not the key's")
    return faults, (None if faults else out)


def worth_values(path, sheet):
    """{id: worth} from one of the PI's worth sheets: column A the id, column C the worth, a whole number 0 to 3
    (kept as text or as a number), as matching_worth.py reads it."""
    import openpyxl
    ws = openpyxl.load_workbook(here(path))[sheet]
    out = {}
    for r in range(2, ws.max_row + 1):
        rid = ws.cell(r, 1).value
        if rid is None:
            continue
        s = "" if ws.cell(r, 3).value is None else str(ws.cell(r, 3).value).strip()
        if s not in ("0", "1", "2", "3"):
            raise SystemExit(f"{path} {rid}: worth {ws.cell(r, 3).value!r}")
        out[rid] = int(s)
    return out


def worth_of_readings(key):
    """{(paper, R id): worth} for the real readings, by matching_key.json as matching_worth.py maps it: chats 49
    and 50's from the sheet of T33 by their M id; the others by their earlier id, B from the sheet of 11 Oct, W
    from the sheet of T32; Pratt's from the sheet of 11 Oct by their P id."""
    mk = {r["reading"]: r for r in json.load(open(here(bcs.MATCHING_KEY), encoding="utf-8"))["readings"]}
    t29b, t29p = worth_values(WORTH_T29, "Bechky"), worth_values(WORTH_T29, "Pratt")
    t32, t33 = worth_values(WORTH_T32, "Bechky"), worth_values(WORTH_T33, "Bechky")
    out = {}
    for e in key["readings"]:
        if e["source"] == "planted":
            continue
        if e["paper"] == "pratt":
            out[("pratt", e["r"])] = t29p[e["source"]]
            continue
        m = mk[e["source"]]
        if m["chat"] in ("49", "50"):
            out[("bechky", e["r"])] = t33[e["source"]]
        elif str(m["earlier_id"]).startswith("B"):
            out[("bechky", e["r"])] = t29b[m["earlier_id"]]
        else:
            out[("bechky", e["r"])] = t32[m["earlier_id"]]
    return out


# ---------------------------------------------------------------- the measures

def pair_stats(pairs):
    po, k = A.kappa(pairs)
    return {"n": len(pairs), "agree": sum(a == b for a, b in pairs), "raw": po, "kappa": k,
            "table": {f"{a}|{b}": n for (a, b), n in sorted(Counter(pairs).items())}}


def planted_verdict(code, claim, typ):
    if code["main_claim"] == claim and code["main_type"] == typ:
        return "caught"
    if (claim, typ) in [tuple(x) for x in code["other_claims"]]:
        return "caught aside"
    return "missed"


def against(x, y, ks):
    """How many of the readings ks two codings x and y give the same main claim, main type and level."""
    return {"n": len(ks),
            "claim": sum(1 for k in ks if x[k]["main_claim"] == y[k]["main_claim"]),
            "type": sum(1 for k in ks if x[k]["main_type"] == y[k]["main_type"]),
            "level": sum(1 for k in ks if level(x[k]["main_type"]) == level(y[k]["main_type"]))}


def measures(key, codes, pi, check, worth):
    """codes: {coder: {(paper, R id): code}}; pi: {(paper, R id): code}; check: {(paper, R id): {...}};
    worth: {(paper, R id): worth} for the real readings."""
    real = {p: [e["r"] for e in key["readings"] if e["paper"] == p and e["source"] != "planted"] for p in PAPERS}
    scopes = {"bechky": [("bechky", r) for r in real["bechky"]], "pratt": [("pratt", r) for r in real["pratt"]]}
    scopes["both"] = scopes["bechky"] + scopes["pratt"]
    g, a = codes["gemini"], codes["astra"]
    m1, m2 = {}, {}
    for s, ks in scopes.items():
        tp = [(g[k]["main_type"], a[k]["main_type"]) for k in ks]
        m1[s] = {"types": pair_stats(tp), "levels": pair_stats([(level(x), level(y)) for x, y in tp])}
        same = [k for k in ks if g[k]["main_claim"] == a[k]["main_claim"]]
        m2[s] = {"readings": len(ks), "same_claim": len(same),
                 "same_claim_same_type": sum(1 for k in same if g[k]["main_type"] == a[k]["main_type"])}
    for name, fn in (("types", lambda x: x), ("levels", level)):
        pairs = [(fn(g[k]["main_type"]), fn(a[k]["main_type"])) for k in scopes["both"]]
        ci = A.boot_interval(pairs, reps=REPS, seed=SEED_BOOT)
        m1["both"][name]["interval"] = None if ci is None else {"low": ci[0], "high": ci[1], "resamples": ci[2]}
    m3 = []
    for x in key["planted"]:
        k = (x["paper"], x["r"])
        row = {"paper": x["paper"], "r": x["r"], "claim": x["claim"], "type": x["type"],
               "confirmed": check[k]["confirmed"], "note": check[k]["note"]}
        for c in CODERS:
            code = codes[c][k]
            row[c] = {"verdict": planted_verdict(code, x["claim"], x["type"]), "main_claim": code["main_claim"],
                      "main_type": code["main_type"], "other_claims": [list(o) for o in code["other_claims"]]}
        m3.append(row)
    sample = {"bechky": [("bechky", r) for r in key["sample"]["bechky"]],
              "pratt": [("pratt", r) for r in key["sample"]["pratt"]]}
    sample["both"] = sample["bechky"] + sample["pratt"]
    m4 = {c: {s: against(codes[c], pi, ks) for s, ks in sample.items()} for c in CODERS}
    m4["coders"] = {s: against(g, a, ks) for s, ks in sample.items()}
    by = {(e["paper"], e["r"]): e for e in key["readings"]}
    chats = sorted({by[k]["chat"] for k in scopes["bechky"]})
    m5 = {"types": {}, "levels": {}, "worth_by_level": {}, "by_chat": {}, "by_arm": {}}
    for c in CODERS:
        lv = {k: level(codes[c][k]["main_type"]) for k in scopes["both"]}
        m5["types"][c] = {t: sum(1 for k in scopes["both"] if codes[c][k]["main_type"] == t) for t in TYPES6}
        m5["levels"][c] = {L: sum(1 for k in scopes["both"] if lv[k] == L) for L in LEVELS}
        m5["by_chat"][c] = {ch: {L: sum(1 for k in scopes["bechky"] if by[k]["chat"] == ch and lv[k] == L)
                                 for L in LEVELS} for ch in chats}
        m5["by_arm"][c] = {arm: {L: sum(1 for k in scopes["pratt"] if by[k]["arm"] == arm and lv[k] == L)
                                 for L in LEVELS} for arm in ("system", "middle", "plain")}
    for L in LEVELS:
        ws = sorted(worth[k] for k in scopes["both"]
                    if level(g[k]["main_type"]) == L and level(a[k]["main_type"]) == L)
        m5["worth_by_level"][L] = {"n": len(ws), "mean": (sum(ws) / len(ws)) if ws else None, "values": ws}
    return {"m1": m1, "m2": m2, "m3": m3, "m4": m4, "m5": m5}


# ---------------------------------------------------------------- the run

def score(opts, replies):
    key = json.load(open(opts["key"], encoding="utf-8"))
    kdir = os.path.dirname(os.path.abspath(opts["key"]))
    stop = []
    if key.get("builder_sha256") != sha(os.path.abspath(brt.__file__)):
        stop.append("the key was not made by this folder's build_relation_test.py")
    for f, h in key.get("inputs", {}).items():
        if sha(here(f)) != h:
            stop.append(f"{f} in this folder is not the input the key names")
    for f, h in key.get("files", {}).items():
        if sha(os.path.join(kdir, f)) != h:
            stop.append(f"{f} beside the key is not the file the key names")
    if len(replies) != 4:
        stop.append("four replies are needed, in the order of chats 64 to 67")
    codes = {c: {} for c in CODERS}
    files = {}
    for (chat, paper, coder), path in zip(CHATS, replies):
        rids = [r["r"] for r in key["readings"] if r["paper"] == paper]
        rep, faults = bcs.load(path)
        remarks, got = [], None
        if not faults:
            faults, remarks, got = brt.coder_codes(rep, rids, key["claims"][paper], paper)
        if faults:
            stop.append(f"chat {chat}: the reply does not count: {'; '.join(faults)}")
            continue
        if bcs.long_spans(rep, key["citations"][paper]):
            stop.append(f"chat {chat}: the reply holds a quoted span of more than eight words")
            continue
        for rid, c in got.items():
            codes[coder][(paper, rid)] = c
        files[chat] = {"file": os.path.basename(path), "sha256": sha(path), "remarks": remarks}
    if stop:
        raise SystemExit("not scored:\n  " + "\n  ".join(stop))
    faults, pi_remarks, pi = read_pi_sheet(opts["sheet"], os.path.join(kdir, brt.SHEET), key)
    f2, check = read_check(opts["check"], os.path.join(kdir, brt.CHECK), key)
    if faults or f2:
        raise SystemExit("not scored: the PI's sheets:\n  " + "\n  ".join(faults + f2))
    worth = worth_of_readings(key)
    res = measures(key, codes, pi, check, worth)
    out = {"what": "the measures of the relation test (T37), by RELATION_TEST_PLAN.md",
           "scorer_sha256": sha(os.path.abspath(__file__)),
           "files": {"key": [os.path.basename(opts["key"]), sha(opts["key"])],
                     "sheet": [os.path.basename(opts["sheet"]), sha(opts["sheet"])],
                     "check": [os.path.basename(opts["check"]), sha(opts["check"])],
                     "replies": files,
                     "worth": {f: sha(here(f)) for f in (WORTH_T29, WORTH_T32, WORTH_T33)}},
           "pi_remarks": pi_remarks,
           **res,
           "codes": {c: {f"{p}|{r}": v for (p, r), v in sorted(codes[c].items())} for c in CODERS},
           "pi": {f"{p}|{r}": v for (p, r), v in sorted(pi.items())},
           "worth": {f"{p}|{r}": v for (p, r), v in sorted(worth.items())}}
    out = json.loads(json.dumps(out))
    os.makedirs(opts["out"], exist_ok=True)
    open(os.path.join(opts["out"], RESULTS_JSON), "w", encoding="utf-8").write(
        json.dumps(out, indent=1, ensure_ascii=False) + "\n")
    open(os.path.join(opts["out"], RESULTS_MD), "w", encoding="utf-8").write(write_md(out, key))
    print(f"{RESULTS_MD}: sha256 {sha(os.path.join(opts['out'], RESULTS_MD))}")
    print(f"{RESULTS_JSON}: sha256 {sha(os.path.join(opts['out'], RESULTS_JSON))}")


# ---------------------------------------------------------------- the report

def f3(x):
    return "—" if x is None else f"{x:.3f}"


def f2(x):
    return "—" if x is None else f"{x:.2f}"


def table_md(stats, codes, rows_name):
    used = [c for c in codes if any(c in k.split("|") for k in stats["table"])]
    L = [f"| {rows_name} | " + " | ".join(f"`{c}`" for c in used) + " |", "|---|" + "---|" * len(used)]
    for a in used:
        L.append(f"| `{a}` | " + " | ".join(str(stats["table"].get(f"{a}|{b}", 0)) for b in used) + " |")
    return L


def write_md(o, key):
    f = o["files"]
    n_real = o["m2"]["both"]["readings"]
    L = ["# The relation test (T37): results", "",
         f"Written by `score_relation_test.py score` (sha256 opens {o['scorer_sha256'][:8]}) by "
         f"`RELATION_TEST_PLAN.md`, from the key `{f['key'][0]}` (opens {f['key'][1][:8]}), the PI's coding sheet "
         f"`{f['sheet'][0]}` (opens {f['sheet'][1][:8]}), the PI's check of the planted readings `{f['check'][0]}` (opens "
         f"{f['check'][1][:8]}), and the four coders' replies: "
         + ", ".join(f"chat {c} `{v['file']}` (opens {v['sha256'][:8]})" for c, v in f["replies"].items()) + ".", "",
         "Described, not tested: the plan fixes no bar. Agreement between the coders shows that the coding can be "
         "repeated across two vendors, not that it is right; the PI's sample and the planted readings speak to "
         "that, and both are small. A type says what a reading does to a claim of the index, not whether it holds "
         "or is worth pursuing. The readings are coded as written, not as standard statements. GPT 6 Astra codes "
         "ten readings its own model wrote; Gemini none.", "",
         "**Checks:** the key is the builder's, from the inputs as they are; the readings files and the blank "
         "sheets are the key's; all four replies count and quote no span of over eight words; on both of the PI's "
         "sheets no cell outside the yellow ones changed, and every answer reads.", ""]
    remarks = [f"chat {c}: {r}" for c, v in f["replies"].items() for r in v["remarks"]]
    remarks += [f"the PI's sheet: {r}" for r in o["pi_remarks"]]
    if remarks:
        L += ["**The checks' remarks** (none of them a fault): " + "; ".join(remarks) + ".", ""]
    L += [f"## M1. The two coders on the main type ({n_real} real readings)", "",
          "| readings | types: agree | types: kappa | levels: agree | levels: kappa |", "|---|---|---|---|---|"]
    for s, name in SCOPES:
        t, lv = o["m1"][s]["types"], o["m1"][s]["levels"]
        L.append(f"| {name} ({t['n']}) | {t['agree']} | {f3(t['kappa'])} | {lv['agree']} | {f3(lv['kappa'])} |")
    L.append("")
    for name in ("types", "levels"):
        ci = o["m1"]["both"][name]["interval"]
        L.append(f"Both papers, {name}: 95% interval "
                 + ("not given" if ci is None else f"{f3(ci['low'])} to {f3(ci['high'])} ({ci['resamples']} of "
                                                   f"{REPS} resamples gave a kappa)") + ".")
        L.append("")
    if any(o["m1"][s][n]["kappa"] is None for s, _ in SCOPES for n in ("types", "levels")):
        L += ["A kappa shown as — is not given: in its table both coders used one and the same code throughout.", ""]
    for s, name in SCOPES:
        L += [f"**{name.capitalize()}, types:**", ""]
        L += table_md(o["m1"][s]["types"], TYPES6, "Gemini in rows, Astra in columns") + [""]
        L += [f"**{name.capitalize()}, levels:**", ""]
        L += table_md(o["m1"][s]["levels"], LEVELS, "Gemini in rows, Astra in columns") + [""]
    L += ["## M2. The main claim", "", "| readings | the same main claim | of those, the same main type |",
          "|---|---|---|"]
    for s, name in SCOPES:
        v = o["m2"][s]
        L.append(f"| {name} ({v['readings']}) | {v['same_claim']} | {v['same_claim_same_type']} |")
    L += ["", "## M3. The planted readings", "",
          "| reading | its claim | written to be | the PI | Gemini | Astra |", "|---|---|---|---|---|---|"]

    def show(v):
        aside = "; " + ", ".join(f"{c} {t}" for c, t in v["other_claims"]) if v["other_claims"] else ""
        return f"{v['verdict']} ({v['main_claim']} {v['main_type']}{aside})"
    for x in o["m3"]:
        L.append(f"| {x['paper'].capitalize()} {x['r']} | {x['claim']} | {x['type']} | "
                 f"{'confirmed' if x['confirmed'] else 'not confirmed'} | {show(x['gemini'])} | {show(x['astra'])} |")
    conf = [x for x in o["m3"] if x["confirmed"]]
    L.append("")
    for c in CODERS:
        n = Counter(x[c]["verdict"] for x in conf)
        L.append(f"{CODER_NAME[c]}: of the {len(conf)} the PI confirmed, caught {n['caught']}, caught aside "
                 f"{n['caught aside']}, missed {n['missed']}.")
        L.append("")
    if len(conf) < len(o["m3"]):
        L += ["The readings the PI did not confirm are shown and left out of these counts, as the plan says.", ""]
    L += ["## M4. Against the PI, on the PI's twelve", "",
          "| | readings | the same main claim | the same main type | the same level |", "|---|---|---|---|---|"]
    for c, name in (("gemini", "Gemini against the PI"), ("astra", "Astra against the PI"),
                    ("coders", "Gemini against Astra")):
        for s, sname in SCOPES:
            v = o["m4"][c][s]
            L.append(f"| {name}, {sname} | {v['n']} | {v['claim']} | {v['type']} | {v['level']} |")
    L += ["", "Each count on its own: the same main type is counted whether or not the main claim is the same.", "",
          "## M5. Described", "", f"**Each coder's types and levels, on the {n_real} real readings:**", "",
          "| | " + " | ".join(f"`{t}`" for t in TYPES6) + " |", "|---|" + "---|" * len(TYPES6)]
    for c in CODERS:
        L.append(f"| {CODER_NAME[c]} | " + " | ".join(str(o["m5"]["types"][c][t]) for t in TYPES6) + " |")
    L += ["", "| | " + " | ".join(LEVELS) + " |", "|---|" + "---|" * len(LEVELS)]
    for c in CODERS:
        L.append(f"| {CODER_NAME[c]} | " + " | ".join(str(o["m5"]["levels"][c][lv]) for lv in LEVELS) + " |")
    L += ["", "**The PI's worth by level, where the two coders agree on the level:**", "",
          "| level | readings | mean worth | worths |", "|---|---|---|---|"]
    for lv in LEVELS:
        v = o["m5"]["worth_by_level"][lv]
        L.append(f"| {lv} | {v['n']} | {f2(v['mean'])} | {', '.join(map(str, v['values'])) or '—'} |")
    L += ["", "The PI's worth was given on three occasions knowing different things (the sheets of 11 Oct, of T32 and "
              "of T33), so these means are described, not compared.", "",
          "**Levels by the chat that wrote the reading (Bechky):**", "",
          "| coder | chat | " + " | ".join(LEVELS) + " |", "|---|---|" + "---|" * len(LEVELS)]
    for c in CODERS:
        for ch, v in o["m5"]["by_chat"][c].items():
            L.append(f"| {CODER_NAME[c]} | {ch} | " + " | ".join(str(v[lv]) for lv in LEVELS) + " |")
    L += ["", "**Levels by arm (Pratt):**", "", "| coder | arm | " + " | ".join(LEVELS) + " |",
          "|---|---|" + "---|" * len(LEVELS)]
    for c in CODERS:
        for arm, v in o["m5"]["by_arm"][c].items():
            L.append(f"| {CODER_NAME[c]} | {arm} | " + " | ".join(str(v[lv]) for lv in LEVELS) + " |")
    L += ["", "## Each reading", "",
          "| reading | source | Gemini | Astra | the PI | worth |", "|---|---|---|---|---|---|"]
    for e in key["readings"]:
        k = f"{e['paper']}|{e['r']}"
        g, a = o["codes"]["gemini"][k], o["codes"]["astra"][k]
        pi = o["pi"].get(k)
        src = e["source"] if e["source"] != "planted" else f"planted (chat {e['chat']})"
        L.append(f"| {e['paper'].capitalize()} {e['r']} | {src} | {g['main_claim']} {g['main_type']} | "
                 f"{a['main_claim']} {a['main_type']} | {'' if pi is None else pi['main_claim'] + ' ' + pi['main_type']} "
                 f"| {o['worth'].get(k, '')} |")
    return "\n".join(L) + "\n"


# ---------------------------------------------------------------- self-test

def selftest():
    bad = []

    def eq(name, got, want):
        if got != want:
            bad.append(f"{name}: got {got!r}, want {want!r}")

    eq("level", [level(t) for t in TYPES6], ["inverts", "dismantles", "dismantles", "dismantles", "complements",
                                             "restates", "none"])
    eq("pv1", planted_verdict({"main_claim": "C05", "main_type": "inverts", "other_claims": []}, "C05", "inverts"),
       "caught")
    eq("pv2", planted_verdict({"main_claim": "C01", "main_type": "complements", "other_claims": [("C05", "inverts")]},
                              "C05", "inverts"), "caught aside")
    eq("pv3", planted_verdict({"main_claim": "C05", "main_type": "complements", "other_claims": []}, "C05",
                              "inverts"), "missed")
    eq("pv4", planted_verdict({"main_claim": "C01", "main_type": "inverts", "other_claims": [("C05", "complements")]},
                              "C05", "inverts"), "missed")
    eq("pv5", planted_verdict({"main_claim": "C01", "main_type": "inverts", "other_claims": [["C05", "inverts"]]},
                              "C05", "inverts"), "caught aside")
    s = pair_stats([("a", "a"), ("a", "b"), ("b", "b"), ("b", "b")])
    eq("ps", (s["n"], s["agree"], s["raw"], s["kappa"]), (4, 3, 0.75, 0.5))
    eq("ps one code", pair_stats([("a", "a"), ("a", "a")])["kappa"], None)
    eq("against", against({1: {"main_claim": "C01", "main_type": "rereads_evidence"}},
                          {1: {"main_claim": "C02", "main_type": "challenges_premise"}}, [1]),
       {"n": 1, "claim": 0, "type": 0, "level": 1})
    eq("nc1", brt.norm_claim(" c3 ", ["C01", "C03"]), ("C03", True))
    eq("nc2", brt.norm_claim("C03", ["C01", "C03"]), ("C03", False))
    eq("nc3", brt.norm_claim("None", ["C01"]), ("none", True))
    eq("nc4", brt.norm_claim("C3a", ["C03"]), (None, False))
    eq("nc5", brt.norm_claim(None, ["C03"]), (None, False))
    eq("nt", (brt.norm_type(" Inverts ", False), brt.norm_type("none", False), brt.norm_type("none", True)),
       ("inverts", None, "none"))

    def code(mc, mt, oc=()):
        return {"main_claim": mc, "main_type": mt, "other_claims": list(oc), "reason": "r", "page": "1"}
    key = {"readings": [{"paper": "bechky", "r": "R01", "source": "M01", "chat": "23", "arm": None},
                        {"paper": "bechky", "r": "R02", "source": "planted", "chat": "62", "arm": None},
                        {"paper": "bechky", "r": "R03", "source": "M02", "chat": "49", "arm": None},
                        {"paper": "bechky", "r": "R04", "source": "M03", "chat": "23", "arm": None},
                        {"paper": "pratt", "r": "R01", "source": "P01", "chat": None, "arm": "plain"},
                        {"paper": "pratt", "r": "R02", "source": "P02", "chat": None, "arm": "system"},
                        {"paper": "pratt", "r": "R03", "source": "planted", "chat": "63", "arm": None}],
           "planted": [{"paper": "bechky", "r": "R02", "claim": "C05", "type": "inverts"},
                       {"paper": "pratt", "r": "R03", "claim": "C02", "type": "restates"}],
           "sample": {"bechky": ["R01", "R03"], "pratt": ["R01"]}}
    g = {("bechky", "R01"): code("C01", "inverts"), ("bechky", "R02"): code("C05", "inverts"),
         ("bechky", "R03"): code("C02", "complements"), ("bechky", "R04"): code("C03", "other_mechanism"),
         ("pratt", "R01"): code("C04", "restates"), ("pratt", "R02"): code("none", "none"),
         ("pratt", "R03"): code("C01", "complements", [("C02", "restates")])}
    a = {("bechky", "R01"): code("C01", "other_mechanism"), ("bechky", "R02"): code("C04", "inverts"),
         ("bechky", "R03"): code("C02", "complements"), ("bechky", "R04"): code("C07", "other_mechanism"),
         ("pratt", "R01"): code("C04", "restates"), ("pratt", "R02"): code("C09", "complements"),
         ("pratt", "R03"): code("C02", "restates")}
    pi = {("bechky", "R01"): {"main_claim": "C01", "main_type": "other_mechanism", "note": None},
          ("bechky", "R03"): {"main_claim": "C02", "main_type": "challenges_premise", "note": None},
          ("pratt", "R01"): {"main_claim": "C04", "main_type": "restates", "note": None}}
    check = {("bechky", "R02"): {"confirmed": True, "note": None}, ("pratt", "R03"): {"confirmed": False, "note": "x"}}
    worth = {("bechky", "R01"): 2, ("bechky", "R03"): 1, ("bechky", "R04"): 3, ("pratt", "R01"): 0,
             ("pratt", "R02"): 2}
    res = measures(key, {"gemini": g, "astra": a}, pi, check, worth)
    eq("m1 bechky types", (res["m1"]["bechky"]["types"]["n"], res["m1"]["bechky"]["types"]["agree"]), (3, 2))
    eq("m1 bechky levels", res["m1"]["bechky"]["levels"]["agree"], 2)
    eq("m1 both types agree", res["m1"]["both"]["types"]["agree"], 3)
    eq("m1 both levels agree", res["m1"]["both"]["levels"]["agree"], 3)
    eq("m1 bechky table", res["m1"]["bechky"]["types"]["table"],
       {"complements|complements": 1, "inverts|other_mechanism": 1, "other_mechanism|other_mechanism": 1})
    eq("m1 interval given", sorted(res["m1"]["both"]["types"]["interval"]), ["high", "low", "resamples"])
    eq("m1 no interval by paper", "interval" in res["m1"]["bechky"]["types"], False)
    eq("m2 both", res["m2"]["both"], {"readings": 5, "same_claim": 3, "same_claim_same_type": 2})
    eq("m3", [(x["r"], x["gemini"]["verdict"], x["astra"]["verdict"], x["confirmed"]) for x in res["m3"]],
       [("R02", "caught", "missed", True), ("R03", "caught aside", "caught", False)])
    eq("m4 gemini both", res["m4"]["gemini"]["both"], {"n": 3, "claim": 3, "type": 1, "level": 1})
    eq("m4 astra both", res["m4"]["astra"]["both"], {"n": 3, "claim": 3, "type": 2, "level": 2})
    eq("m4 astra pratt", res["m4"]["astra"]["pratt"], {"n": 1, "claim": 1, "type": 1, "level": 1})
    eq("m4 coders both", res["m4"]["coders"]["both"], {"n": 3, "claim": 3, "type": 2, "level": 2})
    eq("m5 types astra", res["m5"]["types"]["astra"], {"inverts": 0, "other_mechanism": 2, "rereads_evidence": 0,
                                                       "challenges_premise": 0, "complements": 2, "restates": 1,
                                                       "none": 0})
    eq("m5 levels gemini", res["m5"]["levels"]["gemini"], {"inverts": 1, "dismantles": 1, "complements": 1,
                                                           "restates": 1, "none": 1})
    eq("m5 worth", {lv: (v["n"], v["values"]) for lv, v in res["m5"]["worth_by_level"].items()},
       {"inverts": (0, []), "dismantles": (1, [3]), "complements": (1, [1]), "restates": (1, [0]), "none": (0, [])})
    eq("m5 by chat", res["m5"]["by_chat"]["astra"], {"23": {"inverts": 0, "dismantles": 2, "complements": 0,
                                                            "restates": 0, "none": 0},
                                                     "49": {"inverts": 0, "dismantles": 0, "complements": 1,
                                                            "restates": 0, "none": 0}})
    eq("m5 by arm", res["m5"]["by_arm"]["gemini"]["system"], {"inverts": 0, "dismantles": 0, "complements": 0,
                                                              "restates": 0, "none": 1})
    if bad:
        print("SELFTEST FAILED:\n  " + "\n  ".join(bad))
        return 1
    print("selftest: all checks pass")
    return 0


def main(argv):
    opts, files = {}, []
    for a in argv[1:]:
        if a.startswith("--"):
            k, _, v = a[2:].partition("=")
            opts[k] = v
        else:
            files.append(a)
    mode = argv[0] if argv else None
    if mode == "selftest" and not opts and not files:
        sys.exit(selftest())
    elif mode == "score" and {"key", "sheet", "check", "out"} <= set(opts) and len(files) == 4:
        score(opts, files)
    else:
        raise SystemExit(__doc__)


if __name__ == "__main__":
    main(sys.argv[1:])
