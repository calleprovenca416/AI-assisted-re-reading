"""
Build the files of the relation test (T37), by RELATION_TEST_PLAN.md. Written on 16 Oct 2026, before any
chat of the test opened, on T36's `build_citation_sheet.py`.

    python3 -B build_relation_test.py specs --out=DIR
    python3 -B build_relation_test.py check-writer --paper=bechky|pratt REPLY
    python3 -B build_relation_test.py pack --out=DIR WRITER_62 WRITER_63
    python3 -B build_relation_test.py check-coder --key=relation_key.json --paper=bechky|pratt REPLY

specs writes, for each paper, its claims file (CLAIMS_bechky.md, CLAIMS_pratt.md: the claims of the frozen
index, each with its number, pages and statement, and nothing else) and the specification of its two planted
readings (PLANTED_SPEC_bechky.md, PLANTED_SPEC_pratt.md). The planted claims are drawn with
random.Random(20261021), Bechky's two first, then Pratt's, each draw rng.sample of two from the paper's claims
that the index marks `core` and gives an `outcome` and an `asserted_direction`, in the index's order; the first
drawn is for the first type. The types: Bechky, inverts then complements; Pratt, other_mechanism then
restates.

check-writer: whether a planted-reading chat's reply counts: one JSON object (from its first "{" to its last
"}") with a list `readings` of exactly two items, the first about the spec's first claim and the second about
its second (`claim`), each with a non-empty string `wager` and a list `evidence` of one or more items, each a
`page` (a string) and a non-empty `job`. It names every quoted span of more than eight words in any string of
the reply, by where it is and its length, as `build_citation_sheet.py` defines a span. Exits 1 if the reply does
not count, 2 if it counts but holds such a span, 0 otherwise.

pack: from the two writers' replies as filed (each must count and hold no long span), writes to DIR:
- RELATION_READINGS_bechky.md and RELATION_READINGS_pratt.md, for the coders: each paper's real readings
  (Bechky: M01 to M28 as READINGS_matching.md gives them; Pratt: P01 to P15 as stage2/markers/READINGS_pratt.md
  gives them, sentences joined) and its two planted ones (each put in the plain arm's form, as T33 put chats 49
  and 50's: the wager, then each piece of evidence's job with its page, split into sentences and joined), in
  that order, shuffled and numbered R01 onward.
- relation_sheet.xlsx, the PI's coding sheet: six real readings of each paper, with yellow cells for the main
  claim and the main type (each offering its choices) and a note, the types, the rule of order, the example,
  and both papers' claims.
- planted_check.xlsx, the PI's check of the four planted readings: each with its claim and the type it was
  written to be, and a yellow cell for yes or no, and one for a note.
- relation_key.json, for nobody: which reading is which, the planted claims and types, the PI's sample, the seeds,
  the hashes of every file read and written, and this script's.
One generator, random.Random(20261022), in this order: Bechky's shuffle, Pratt's, then the PI's six of Bechky's real
readings and six of Pratt's (each rng.sample from the real readings in R order, then sorted). Workbooks carry
fixed dates, so two builds give the same bytes.

check-coder: whether a coder chat's reply counts, by the plan: one JSON object with a list `codes` naming every
reading of its paper's file exactly once; each `main_claim` one of the paper's claims or `none`; each
`main_type` one of the six types, or `none` exactly when the main claim is `none`; `other_claims` a list of at
most two items, each a claim of the paper other than the main one, named once, with one of the six types; a
non-empty `reason`; and a `page`. A claim number written without its zero or in lower case (c3, C3) is read as
the paper's (C03), and a type with its ends trimmed and capitals aside; the check says so of each. Long quoted
spans are named as in check-writer. Exits as check-writer does.

Both checks remark, without counting it a fault, when the reply's `paper` does not name the paper's author:
the chat may have had another PDF.
"""
import datetime
import hashlib
import json
import os
import random
import re
import sys
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import build_citation_sheet as bcs   # noqa: E402
import build_marker_packs as bmp     # noqa: E402
import build_places_packs as bpp     # noqa: E402
import build_worth_sheet as bws      # noqa: E402
from openpyxl import Workbook        # noqa: E402
from openpyxl.styles import Alignment   # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation   # noqa: E402

PAPERS = ("bechky", "pratt")
CORPUS = {"bechky": "bechky_2020_v1.json", "pratt": "pratt_2000_v1.json"}
TYPES = ("inverts", "other_mechanism", "rereads_evidence", "challenges_premise", "complements", "restates")
LEVEL = {"inverts": "inverts", "other_mechanism": "dismantles", "rereads_evidence": "dismantles",
         "challenges_premise": "dismantles", "complements": "complements", "restates": "restates", "none": "none"}
PLANTED_TYPES = {"bechky": ("inverts", "complements"), "pratt": ("other_mechanism", "restates")}
WRITER_CHAT = {"bechky": "62", "pratt": "63"}
SEED_CLAIMS = 20261021
SEED_PACK = 20261022
SAMPLE = 6
STAMP = (2026, 10, 16, 0, 0, 0)
READINGS_FILE = {"bechky": "RELATION_READINGS_bechky.md", "pratt": "RELATION_READINGS_pratt.md"}
CLAIMS_FILE = {"bechky": "CLAIMS_bechky.md", "pratt": "CLAIMS_pratt.md"}
SPEC_FILE = {"bechky": "PLANTED_SPEC_bechky.md", "pratt": "PLANTED_SPEC_pratt.md"}
SHEET = "relation_sheet.xlsx"
CHECK = "planted_check.xlsx"
KEY = "relation_key.json"
NAME = {"bechky": "Bechky", "pratt": "Pratt"}

DEFINE = {
    "inverts": "the reading holds that the claim's outcome goes the other way under the condition the claim "
               "names, or that a process the claim states runs in the reverse order.",
    "other_mechanism": "the reading keeps the claim's outcome but puts a different mechanism or cause in place "
                       "of the paper's.",
    "rereads_evidence": "the reading takes passages the paper itself offers for the claim and reads them so that "
                        "they no longer support it as stated.",
    "challenges_premise": "the reading questions something the paper's account of the claim takes for granted "
                          "(a premise, a category, a boundary), rather than its outcome or its mechanism.",
    "complements": "the reading adds a mechanism, a condition or a level to the claim, or changes its scope, "
                   "without denying it.",
    "restates": "the reading says what the claim already says.",
}
ASK = {
    "inverts": "**invert** it: hold that the claim's outcome goes the other way under the condition the claim "
               "names, or that a process the claim states runs in the reverse order",
    "complements": "**complement** it: add a mechanism, a condition or a level to the claim, or change its "
                   "scope, without denying it",
    "other_mechanism": "put **another mechanism** in place of the paper's: keep the claim's outcome, but explain "
                       "it by a different mechanism or cause than the paper's",
    "restates": "**restate** it: say what the claim already says, in other words, while presenting itself as "
                "another reading of the paper",
}
EXAMPLE = [
    ("Claims", "From a made-up paper: C1 (p. 4), cooks under time pressure skip the tasting step because they "
               "value speed over care; C2 (p. 7), head chefs who taste every dish themselves keep the step alive."),
    ("R1", "\"Cooks under time pressure taste more, not less, since a mistake costs more when the pass is full.\" "
           "C1, inverts."),
    ("R2", "\"The step is skipped because the recipe cards go missing in the rush, not because cooks value "
           "speed.\" C1, other_mechanism."),
    ("R3", "\"The paper's own episodes of skipping all fall on nights the head chef was away; read again, they "
           "show missing supervision, not speed.\" C1, rereads_evidence."),
    ("R4", "\"The paper treats tasting as one step; in these kitchens it is three different checks, and only one "
           "is ever skipped.\" C1, challenges_premise."),
    ("R5", "\"Head chefs keep the step alive only in kitchens with a separate pass station.\" C2, complements."),
    ("R6", "\"Where the head chef tastes every dish, the kitchen keeps tasting.\" C2, restates."),
]


def here(name):
    return os.path.join(HERE, name)


def sha(path):
    return hashlib.sha256(open(path, "rb").read()).hexdigest()


def corpus(paper):
    return json.load(open(here(CORPUS[paper]), encoding="utf-8"))


def claims(paper):
    return [(c["claim_id"], c["location"].strip(), c["statement"].strip()) for c in corpus(paper)["claims"]]


def citation(paper):
    return corpus(paper)["meta"]["citation"]


def planted_claims():
    """{paper: [(claim, type), (claim, type)]}, drawn as the docstring says."""
    rng = random.Random(SEED_CLAIMS)
    out = {}
    for paper in PAPERS:
        cands = [c["claim_id"] for c in corpus(paper)["claims"]
                 if c.get("core") and c.get("outcome") and c.get("asserted_direction")]
        picks = rng.sample(cands, 2)
        out[paper] = list(zip(picks, PLANTED_TYPES[paper]))
    return out


# ---------------------------------------------------------------- specs

def claims_text(paper):
    cs = claims(paper)
    lines = [f"# The claims of {citation(paper)}", "",
             f"The paper's {len(cs)} claims as its index records them, numbered, each with the pages where the paper",
             "makes it and its statement. A statement is a summary, not the paper's words: check the page.", ""]
    for cid, loc, st in cs:
        lines += [f"**{cid}** ({loc}). {st}", ""]
    return "\n".join(lines)


def spec_text(paper, picks):
    by_id = {cid: (loc, st) for cid, loc, st in claims(paper)}
    lines = ["# Two readings, written to a specification", "",
             "The attached PDF is a published qualitative paper. Write two alternative readings of it, each about "
             "one of the paper's claims below, and each doing one particular thing to its claim. A claim's statement "
             "is a summary from an index of the paper, not the paper's words.", ""]
    for i, (cid, typ) in enumerate(picks, 1):
        loc, st = by_id[cid]
        lines += [f"{i}. **About claim {cid}** ({loc}): {st} The reading must {ASK[typ]}.", ""]
    lines += ["Write each as an alternative reading is written: a **wager** in plain prose, of 80 to 150 words, "
              "saying what the paper's evidence would support if read this way; and its **evidence**, the passages "
              "of the paper it rests on, each cited by the journal's page number, with what each passage does in "
              "the reading. Make it as plausible as the paper's evidence allows. Write it as a reading, not as an "
              "exercise: do not name the claim's number, and do not say what the reading does to the claim in the "
              "words above.", "",
              "Do not quote more than eight words of the paper in a row. Paraphrase, and cite the page. Work only "
              "from the paper. Do not search the web.", "",
              "Return a single JSON object and nothing else, as a file if you can, otherwise as one code block:", "",
              "```json",
              "{",
              '  "paper": "author and year, as the paper prints them",',
              '  "readings": [',
              '    {',
              f'      "reading_id": "S1",',
              f'      "claim": "{picks[0][0]}",',
              '      "wager": "...",',
              '      "evidence": [{"page": "p. ..", "job": "..."}]',
              '    },',
              '    {',
              f'      "reading_id": "S2",',
              f'      "claim": "{picks[1][0]}",',
              '      "wager": "...",',
              '      "evidence": [{"page": "p. ..", "job": "..."}]',
              '    }',
              '  ]',
              "}",
              "```", ""]
    return "\n".join(lines)


def specs(out):
    os.makedirs(out, exist_ok=True)
    picks = planted_claims()
    for paper in PAPERS:
        for name, text in ((CLAIMS_FILE[paper], claims_text(paper)), (SPEC_FILE[paper], spec_text(paper, picks[paper]))):
            open(os.path.join(out, name), "w", encoding="utf-8").write(text)
            print(f"{name}: sha256 {hashlib.sha256(text.encode()).hexdigest()}")
        print(f"  {paper}: planted {picks[paper]}")


# ---------------------------------------------------------------- the replies

def spans_report(rep, cite):
    return bcs.long_spans(rep, cite)


def writer_faults(rep, paper):
    if not isinstance(rep, dict):
        return ["not a JSON object"]
    rs = rep.get("readings")
    if not isinstance(rs, list) or len(rs) != 2:
        return ["`readings` is not a list of exactly two items"]
    faults = []
    picks = planted_claims()[paper]
    for i, (r, (cid, _)) in enumerate(zip(rs, picks), 1):
        if not isinstance(r, dict):
            faults.append(f"reading {i}: not an object")
            continue
        if str(r.get("claim", "")).strip() != cid:
            faults.append(f"reading {i}: claim {r.get('claim')!r}, not {cid}")
        if not isinstance(r.get("wager"), str) or not r["wager"].strip():
            faults.append(f"reading {i}: no non-empty wager")
        ev = r.get("evidence")
        if not isinstance(ev, list) or not ev or any(
                not isinstance(e, dict) or not isinstance(e.get("page"), str)
                or not isinstance(e.get("job"), str) or not e["job"].strip() for e in ev):
            faults.append(f"reading {i}: evidence is not a list of one or more items, each a page and a job")
    return faults


def check_writer(path, paper):
    rep, faults = bcs.load(path)
    faults = faults or writer_faults(rep, paper)
    print(f"chat {WRITER_CHAT[paper]} ({paper}, the planted readings): {os.path.basename(path)}, sha256 {sha(path)}")
    if faults:
        for f in faults:
            print(f"  FAULT: {f}")
        print("  DOES NOT COUNT")
        return 1
    for r in paper_remark(rep, paper):
        print(f"  remark: {r}")
    for i, r in enumerate(rep["readings"], 1):
        n = len(r["evidence"])
        print(f"  reading {i}: claim {r['claim']}, wager {len(r['wager'].split())} words, "
              f"{n} {'piece' if n == 1 else 'pieces'} of evidence")
    spans = spans_report(rep, citation(paper))
    for where, words, unclosed, inside in spans:
        print(f"  QUOTED SPAN over {bcs.MAX_QUOTE} words: {where}: {words} words"
              f"{', unclosed' if unclosed else ''}{', within the citation' if inside else ''}")
    if spans:
        print("  COUNTS; to be shown to the PI before it is filed (the plan): the spans above are not filed")
        return 2
    print("  COUNTS; no quoted span of more than eight words")
    return 0


def norm_claim(v, ids):
    """(the paper's claim id or 'none', or None if unreadable; whether it was read loosely)."""
    s = "" if v is None else str(v).strip()
    if s in ids or s == "none":
        return s, False
    if s.casefold() == "none":
        return "none", True
    m = re.fullmatch(r"[Cc]\s*0*(\d+)", s)
    if m:
        c = f"C{int(m.group(1)):02d}"
        if c in ids:
            return c, True
    return None, False


def norm_type(v, allow_none):
    s = "" if v is None else str(v).strip().casefold()
    if s in TYPES or (allow_none and s == "none"):
        return s
    return None


def paper_remark(rep, paper):
    """A remark if the reply's `paper` does not name the paper's author: the chat may have had another PDF."""
    given = rep.get("paper") if isinstance(rep, dict) else None
    if NAME[paper].casefold() not in str(given).casefold():
        return [f"the reply gives the paper as {given!r}, which does not name {NAME[paper]}: check the chat's PDF"]
    return []


def coder_codes(rep, rids, ids, paper=None):
    """(faults, remarks, {rid: code}) for one coder's reply."""
    if not isinstance(rep, dict):
        return ["not a JSON object"], [], None
    codes = rep.get("codes")
    if not isinstance(codes, list):
        return ["no list `codes`"], [], None
    faults, remarks, out = [], (paper_remark(rep, paper) if paper else []), {}
    for k, c in enumerate(codes, 1):
        if not isinstance(c, dict):
            faults.append(f"code {k}: not an object")
            continue
        rid = str(c.get("reading", "")).strip()
        if rid not in rids:
            faults.append(f"code {k}: reading {c.get('reading')!r} is not in the file")
            continue
        if rid in out:
            faults.append(f"{rid}: coded twice")
            continue
        mc, loose = norm_claim(c.get("main_claim"), ids)
        if mc is None:
            faults.append(f"{rid}: main_claim {c.get('main_claim')!r} is not one of the paper's claims or none")
            continue
        if loose:
            remarks.append(f"{rid}: main_claim {c.get('main_claim')!r} read as {mc}")
        mt = norm_type(c.get("main_type"), allow_none=True)
        if mt is None or (mt == "none") != (mc == "none"):
            faults.append(f"{rid}: main_type {c.get('main_type')!r} with main_claim {mc}")
            continue
        if c.get("main_type") != mt:
            remarks.append(f"{rid}: main_type {c.get('main_type')!r} read as {mt}")
        others = c.get("other_claims", [])
        if others is None:
            others = []
        if not isinstance(others, list) or len(others) > 2:
            faults.append(f"{rid}: other_claims is not a list of at most two")
            continue
        oc, bad = [], False
        for o in others:
            if not isinstance(o, dict):
                bad = True
                break
            cid, loose = norm_claim(o.get("claim"), ids)
            typ = norm_type(o.get("type"), allow_none=False)
            if cid in (None, "none") or typ is None or cid == mc or cid in [x[0] for x in oc]:
                bad = True
                break
            if loose:
                remarks.append(f"{rid}: other claim {o.get('claim')!r} read as {cid}")
            if o.get("type") != typ:
                remarks.append(f"{rid}: other claim's type {o.get('type')!r} read as {typ}")
            oc.append((cid, typ))
        if bad:
            faults.append(f"{rid}: other_claims {others!r} do not read")
            continue
        reason = c.get("reason")
        if not isinstance(reason, str) or not reason.strip():
            faults.append(f"{rid}: no reason")
            continue
        page = c.get("page")
        if page is None or not isinstance(page, (str, int)):
            faults.append(f"{rid}: no page")
            continue
        out[rid] = {"main_claim": mc, "main_type": mt, "other_claims": oc, "reason": reason.strip(),
                    "page": str(page).strip()}
    missing = [r for r in rids if r not in out]
    if missing and not faults:
        faults.append(f"readings not coded: {missing}")
    order = [str(c.get("reading", "")).strip() for c in codes if isinstance(c, dict)]
    if not faults and order != list(rids):
        remarks.append("the codes are not in the file's order")
    return faults, remarks, (None if faults else out)


def check_coder(key_path, paper, path):
    key = json.load(open(key_path, encoding="utf-8"))
    rids = [r["r"] for r in key["readings"] if r["paper"] == paper]
    ids = key["claims"][paper]
    rep, faults = bcs.load(path)
    remarks, codes = [], None
    if not faults:
        faults, remarks, codes = coder_codes(rep, rids, ids, paper)
    print(f"{paper} coder reply: {os.path.basename(path)}, sha256 {sha(path)}")
    for r in remarks:
        print(f"  remark: {r}")
    if faults:
        for f in faults:
            print(f"  FAULT: {f}")
        print("  DOES NOT COUNT")
        return 1
    from collections import Counter
    print(f"  {len(codes)} readings coded; main types {dict(Counter(c['main_type'] for c in codes.values()))}")
    spans = spans_report(rep, key["citations"][paper])
    for where, words, unclosed, inside in spans:
        print(f"  QUOTED SPAN over {bcs.MAX_QUOTE} words: {where}: {words} words"
              f"{', unclosed' if unclosed else ''}{', within the citation' if inside else ''}")
    if spans:
        print("  COUNTS; to be shown to the PI before it is filed (the plan): the spans above are not filed")
        return 2
    print("  COUNTS; no quoted span of more than eight words")
    return 0


# ---------------------------------------------------------------- the workbooks

def steady(path):
    stamp = b"2026-10-16T00:00:00Z"
    with zipfile.ZipFile(path) as z:
        items = [(i, z.read(i.filename)) for i in z.infolist()]
    with zipfile.ZipFile(path, "w") as z:
        for i, data in items:
            if i.filename == "docProps/core.xml":
                data, n = re.subn(rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)",
                                  rb"\g<1>" + stamp + rb"\g<2>", data)
                if n != 1:
                    raise SystemExit("docProps/core.xml: no single modified date")
            j = zipfile.ZipInfo(i.filename, date_time=STAMP)
            j.compress_type = i.compress_type
            j.external_attr = i.external_attr
            z.writestr(j, data)


def save(wb, path):
    wb.properties.created = wb.properties.modified = datetime.datetime(*STAMP)
    wb.properties.creator = "build_relation_test.py"
    wb.save(path)
    steady(path)


def instructions(wb, rows):
    ins = wb.active
    ins.title = "Instructions"
    ins.column_dimensions["A"].width = 18
    ins.column_dimensions["B"].width = 110
    for r, (label, t, bold, size) in enumerate(rows, 1):
        bws.cell(ins, f"A{r}", label, bold=bold, size=size, wrap=r != 1)
        if t is not None:
            bws.cell(ins, f"B{r}", t, bold=False, size=size)
            ins.row_dimensions[r].height = bws.height(t, 120)
    bws.printable(ins)


def table_sheet(wb, title, heads, widths, rows, text_col, yellow=(), validations=()):
    ws = wb.create_sheet(title)
    cols = "ABCDEFG"[:len(heads)]
    for col, w in zip(cols, widths):
        ws.column_dimensions[col].width = w
    for i, h in enumerate(heads):
        bws.cell(ws, f"{cols[i]}1", h, bold=True, fill=bws.GREY).border = bws.BOX
    for dv, _ in validations:
        ws.add_data_validation(dv)
    for r, row in enumerate(rows, start=2):
        for i, v in enumerate(row):
            c = bws.cell(ws, f"{cols[i]}{r}", v, bold=(i == 0), fill=(bws.YELLOW if i in yellow else None))
            c.border = bws.BOX
            if i in yellow:
                c.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
        for dv, col in validations:
            dv.add(f"{col}{r}")
        ws.row_dimensions[r].height = bws.height(row[text_col], max(20, widths[text_col] + 12))
    ws.freeze_panes = "B2"
    bws.printable(ws)
    return ws


def list_validation(options, title, prompt):
    formula = '"' + ",".join(options) + '"'
    if len(formula) > 255:
        raise SystemExit(f"{title}: the list of choices is too long for a cell")
    return DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=True,
                          errorTitle=title, error=f"Choose one of the list: {prompt}", showInputMessage=True,
                          promptTitle=title, prompt=prompt)


# ---------------------------------------------------------------- pack

def real_readings():
    """{paper: [{source, chat, arm, text}]} in the pre-shuffle order."""
    texts = bcs.m_readings()
    mk = {r["reading"]: r for r in json.load(open(here(bcs.MATCHING_KEY), encoding="utf-8"))["readings"]}
    bech = [{"source": m, "chat": mk[m]["chat"], "arm": None, "text": t} for m, t in texts.items()]
    pratt = []
    for e in bcs.p_readings():
        src = e["reading"]
        pratt.append({"source": src, "chat": None, "arm": e["arm"], "text": e["text"]})
    if [r["source"] for r in bech] != [f"M{i:02d}" for i in range(1, 29)] or len(pratt) != 15:
        raise SystemExit("the real readings are not M01 to M28 and P01 to P15")
    return {"bechky": bech, "pratt": pratt}


def pack(out, writer_paths):
    picks = planted_claims()
    reals = real_readings()
    writers, lists = {}, {}
    for paper, path in zip(PAPERS, writer_paths):
        rep, faults = bcs.load(path)
        faults = faults or writer_faults(rep, paper)
        if faults:
            raise SystemExit(f"chat {WRITER_CHAT[paper]}: the reply does not count: {'; '.join(faults)}")
        if spans_report(rep, citation(paper)):
            raise SystemExit(f"chat {WRITER_CHAT[paper]}: the reply holds a quoted span of more than eight words")
        writers[WRITER_CHAT[paper]] = {"paper": paper, "file": os.path.basename(path), "sha256": sha(path)}
        planted = []
        for r, (cid, typ) in zip(rep["readings"], picks[paper]):
            planted.append({"source": "planted", "chat": WRITER_CHAT[paper], "arm": None,
                            "reading_id": r.get("reading_id"), "claim": cid, "type": typ,
                            "text": " ".join(bmp.sentences(bpp.reading_text(r)))})
        lists[paper] = reals[paper] + planted
    rng = random.Random(SEED_PACK)
    for paper in PAPERS:
        rng.shuffle(lists[paper])
        for i, e in enumerate(lists[paper], 1):
            e["r"] = f"R{i:02d}"
    sample = {}
    for paper in PAPERS:
        real_rids = [e["r"] for e in lists[paper] if e["source"] != "planted"]
        sample[paper] = sorted(rng.sample(real_rids, SAMPLE))
    os.makedirs(out, exist_ok=True)
    files = {}
    for paper in PAPERS:
        es = lists[paper]
        lines = [f"# Readings of {citation(paper)}", "",
                 f"{len(es)} readings of this paper, in no particular order, numbered R01 to R{len(es):02d}. A page "
                 "in a reading", "is a pointer, not a proof.", ""]
        for e in es:
            lines += [f"## Reading {e['r']}", "", e["text"], ""]
        p = os.path.join(out, READINGS_FILE[paper])
        open(p, "w", encoding="utf-8").write("\n".join(lines))
        files[READINGS_FILE[paper]] = sha(p)

    # the PI's coding sheet
    wb = Workbook()
    rows = [("Coding how each reading relates to the paper's claims", None, True, 13),
            ("What", "Twelve readings, six of each paper, on the sheets Bechky and Pratt. For each, give the claim it "
                     "bears on most directly (its main claim) and what it does to that claim (its main type), as the "
                     "coders are asked to. Each paper's claims are on the sheets Claims Bechky and Claims Pratt.",
             True, 10),
            ("Fill in", "Only the yellow cells: Main claim, Main type (each cell offers its choices) and Note. Leave "
                        "every other cell as it is. If a reading bears on no claim in the list, give none in both.",
             True, 10),
            ("The six types", "For one reading and one claim, one type:", True, 10)]
    rows += [(t, DEFINE[t][0].upper() + DEFINE[t][1:], False, 10) for t in TYPES]
    rows += [("The order", "If more than one type fits, give the first that fits in the order above.", True, 10),
             ("Example", "From a made-up paper, not one of these:", True, 10)]
    rows += [(a, b, False, 10) for a, b in EXAMPLE]
    rows += [("Rules", None, True, 10),
             ("1", "Code each reading on its own words, whoever may have written it, and not by its worth.", False, 10),
             ("2", "Please fill this sheet before you open any of chats 64 to 67, and do not open relation_key.json "
                   "until the coders' replies are in.", False, 10),
             ("3", "You may consult the papers. A page in a reading is a pointer, not a proof.", False, 10),
             ("4", "Save the workbook as relation_sheet_filled.xlsx, with no other change, and send it in our "
                   "conversation or put it in replies/.", False, 10)]
    instructions(wb, rows)
    by_r = {p: {e["r"]: e for e in lists[p]} for p in PAPERS}
    for paper in PAPERS:
        ids = [c[0] for c in claims(paper)]
        dvc = list_validation(ids + ["none"], "Main claim", "a claim of the list, or none")
        dvt = list_validation(list(TYPES) + ["none"], "Main type", "one of the six types, or none")
        table_sheet(wb, NAME[paper], ["Reading", "The reading", "Main claim", "Main type", "Note"],
                    (10, 95, 14, 22, 40), [(r, by_r[paper][r]["text"], None, None, None) for r in sample[paper]], 1,
                    yellow=(2, 3, 4), validations=((dvc, "C"), (dvt, "D")))
    for paper in PAPERS:
        table_sheet(wb, f"Claims {NAME[paper]}", ["Claim", "Pages", "Statement"], (10, 30, 100),
                    [(cid, loc, st) for cid, loc, st in claims(paper)], 2)
    p = os.path.join(out, SHEET)
    save(wb, p)
    files[SHEET] = sha(p)

    # the PI's check of the planted readings
    wb = Workbook()
    rows = [("Checking the four planted readings", None, True, 13),
            ("What", "Four readings were written by a model to a specification: two of each paper, each to do one "
                     "thing to one claim. They are mixed, unmarked, among the readings the coders get. Say of each "
                     "whether it does that thing to that claim: yes or no.", True, 10),
            ("Fill in", "Only the yellow cells: Of this type? (yes or no; the cell offers them) and Note.", True, 10),
            ("Rules", None, True, 10),
            ("1", "A reading you answer no stays in the coders' files and is left out of the count of planted "
                  "readings caught, and reported.", False, 10),
            ("2", "Save the workbook as planted_check_filled.xlsx, with no other change, and send it in our "
                  "conversation or put it in replies/.", False, 10)]
    instructions(wb, rows)
    dvy = list_validation(["yes", "no"], "Of this type?", "yes or no")
    prow = []
    for paper in PAPERS:
        cl = {cid: (loc, st) for cid, loc, st in claims(paper)}
        for e in lists[paper]:
            if e["source"] == "planted":
                loc, st = cl[e["claim"]]
                prow.append((f"{NAME[paper]} {e['r']}", e["text"], f"{e['claim']} ({loc}): {st}",
                             f"{e['type']}: {DEFINE[e['type']]}", None, None))
    prow.sort(key=lambda x: (x[0].split()[0] != "Bechky", x[0]))
    table_sheet(wb, "Planted", ["Reading", "The reading", "Its claim", "Written to be", "Of this type?", "Note"],
                (14, 80, 45, 40, 14, 30), prow, 1, yellow=(4, 5), validations=((dvy, "E"),))
    p = os.path.join(out, CHECK)
    save(wb, p)
    files[CHECK] = sha(p)

    inputs = [bcs.READINGS_M, bcs.MATCHING_KEY, bcs.READINGS_P, bcs.KEY_11, CORPUS["bechky"], CORPUS["pratt"]]
    key = {"what": "the key to the relation test's readings files and the PI's two sheets: for nobody, and not for "
                   "the PI until the coders' replies are in",
           "seeds": {"claims": SEED_CLAIMS, "pack": SEED_PACK},
           "writers": writers,
           "inputs": {f: sha(here(f)) for f in inputs},
           "builder_sha256": sha(os.path.abspath(__file__)),
           "citations": {paper: citation(paper) for paper in PAPERS},
           "claims": {paper: [c[0] for c in claims(paper)] for paper in PAPERS},
           "planted": [{"paper": paper, "r": e["r"], "chat": e["chat"], "reading_id": e["reading_id"],
                        "claim": e["claim"], "type": e["type"]}
                       for paper in PAPERS for e in lists[paper] if e["source"] == "planted"],
           "readings": [{"paper": paper, "r": e["r"], "source": e["source"], "chat": e["chat"], "arm": e["arm"],
                         "text": e["text"]} for paper in PAPERS for e in lists[paper]],
           "sample": sample,
           "files": files}
    open(os.path.join(out, KEY), "w", encoding="utf-8").write(json.dumps(key, indent=1, ensure_ascii=False) + "\n")
    for f, h in files.items():
        print(f"{f}: sha256 {h}")
    print(f"{KEY}: sha256 {sha(os.path.join(out, KEY))}")
    print(f"planted: {[(x['paper'], x['r'], x['claim'], x['type']) for x in key['planted']]}; sample: {sample}")


def main(argv):
    opts, files = {}, []
    for a in argv[1:]:
        if a.startswith("--"):
            k, _, v = a[2:].partition("=")
            opts[k] = v
        else:
            files.append(a)
    mode = argv[0] if argv else None
    if mode == "specs" and "out" in opts and not files:
        specs(opts["out"])
    elif mode == "check-writer" and opts.get("paper") in PAPERS and len(files) == 1:
        sys.exit(check_writer(files[0], opts["paper"]))
    elif mode == "pack" and "out" in opts and len(files) == 2:
        pack(opts["out"], files)
    elif mode == "check-coder" and opts.get("paper") in PAPERS and "key" in opts and len(files) == 1:
        sys.exit(check_coder(opts["key"], opts["paper"], files[0]))
    else:
        raise SystemExit(__doc__)


if __name__ == "__main__":
    main(sys.argv[1:])
