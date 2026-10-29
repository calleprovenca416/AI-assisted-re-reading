"""
Score T34, the authors' own worth for the wager readings, by AUTHOR_RATINGS_PLAN.md
and its two dated changes. Written on 15 Oct 2026, before any reply came, and
tested on made-up answers.

    python3 -B score_author_ratings.py selftest
    python3 -B score_author_ratings.py sheet --paper=bechky|pratt --out=DIR
    python3 -B score_author_ratings.py score --paper=bechky|pratt --answers=FILE --out=DIR

sheet writes author_answers_<paper>.xlsx, for the PI alone: a sheet on which the PI
passes on the author's numbers and second answers, reading by reading, as the
plan's change of 14 Oct asks. It holds the readings' numbers as in the
author's file (1 to 14 for Bechky, 1 to 15 for Pratt), a cell for her worth (0
to 3) and one for her second answer, each chosen from a list, and nothing else:
no reading's text, id or arm, and no cell for her words. It carries a fixed date
inside, so two builds give the same bytes.

score reads her answers, from that sheet filled (.xlsx) or from a JSON file
{"answers": [{"number": 1, "worth": 2, "second": "new to me"}, ...]}, with the
paper's key (author_ratings_key.json, or author_ratings_key_pratt.json) and the
PI's worth of 11 Oct (stage2/worth/worth_sheet_filled_2.xlsx, the paper's tab),
and writes AUTHOR_RATINGS_<PAPER>.md and author_ratings_<paper>.json with what
the plan's "What will be reported" lists, and nothing else:

1. her worth by arm (system, middle, plain), as means and as counts of 0, 1, 2
   and 3, beside the PI's worth for the same readings;
2. reading by reading, on how many she and the PI give the same worth, and on
   how many they are within one point;
3. how many she rates 2 or more, and how many of those she marks "new to me";
4. her second answers, counted by arm.

All of it for the readings she rates (those with a worth), with their count, and
no other cut or total. The readings the PI rated late, after the PI had been told
their arm (B06; P12 and P14), are marked. A worth must
be a whole number from 0 to 3, and a second answer one of "new to me", "I had
considered it" and "it is in the paper", read without regard to capitals,
spaces or the kind of apostrophe; anything else stops the score and names the
reading's number, for the PI to say how it is to be read. A second answer given
to a reading without a worth is counted apart. No test of significance; each
paper on its own, never pooled.
"""
import datetime
import hashlib
import json
import os
import sys
import unicodedata

HERE = os.path.dirname(os.path.abspath(__file__))
PAPERS = {"bechky": {"key": "author_ratings_key.json", "tab": "Bechky", "n": 14, "late": ("B06",),
                     "name": "Bechky (2020)"},
          "pratt": {"key": "author_ratings_key_pratt.json", "tab": "Pratt", "n": 15, "late": ("P12", "P14"),
                    "name": "Pratt (2000)"}}
PI_SHEET = "stage2/worth/worth_sheet_filled_2.xlsx"
ARMS = ("system", "middle", "plain")
SECONDS = ("new to me", "I had considered it", "it is in the paper")
STAMP = (2026, 10, 15, 0, 0, 0)


class Refused(Exception):
    pass


def sha(path):
    return hashlib.sha256(open(path, "rb").read()).hexdigest()


def norm(s):
    s = unicodedata.normalize("NFKC", str(s)).replace("’", "'").replace("‘", "'")
    return " ".join(s.split()).casefold()


SECOND_OF = {norm(x): x for x in SECONDS}


def read_worth(v, where):
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, bool):
        raise Refused(f"{where}: worth {v!r} is not 0 to 3")
    if isinstance(v, (int, float)) and float(v).is_integer() and 0 <= v <= 3:
        return int(v)
    if isinstance(v, str) and v.strip() in ("0", "1", "2", "3"):
        return int(v.strip())
    raise Refused(f"{where}: worth {v!r} is not a whole number from 0 to 3")


def read_second(v, where):
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    s = SECOND_OF.get(norm(v))
    if s is None:
        raise Refused(f"{where}: second answer {v!r} is not one of {SECONDS}")
    return s


def load_answers(path, n):
    """{number: (worth or None, second or None)} for every number 1..n."""
    if path.endswith(".xlsx"):
        from openpyxl import load_workbook
        ws = load_workbook(path)["Answers"]
        if ws.cell(1, 1).value != "Reading":
            raise Refused("the sheet's first column is not headed Reading")
        raw = {}
        for r in range(2, ws.max_row + 1):
            num = ws.cell(r, 1).value
            if num is None:
                continue
            raw[num] = (ws.cell(r, 2).value, ws.cell(r, 3).value)
    else:
        obj = json.load(open(path, encoding="utf-8"))
        raw = {}
        for a in obj.get("answers", []):
            if a.get("number") in raw:
                raise Refused(f"reading {a.get('number')} is given twice")
            raw[a.get("number")] = (a.get("worth"), a.get("second"))
    if set(raw) - set(range(1, n + 1)):
        raise Refused(f"numbers outside 1 to {n}: {sorted(map(str, set(raw) - set(range(1, n + 1))))}")
    out = {}
    for i in range(1, n + 1):
        w, s = raw.get(i, (None, None))
        out[i] = (read_worth(w, f"reading {i}"), read_second(s, f"reading {i}"))
    return out


def load_pi(paper):
    from openpyxl import load_workbook
    ws = load_workbook(os.path.join(HERE, PI_SHEET))[PAPERS[paper]["tab"]]
    return {ws.cell(r, 1).value: read_worth(ws.cell(r, 3).value, f"the PI's {ws.cell(r, 1).value}")
            for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value}


def mean(xs):
    return sum(xs) / len(xs) if xs else None


def counts(xs):
    return {str(k): sum(1 for x in xs if x == k) for k in range(4)}


def figures(answers, readings, pi, late):
    """answers {number: (worth, second)}; readings [{number, reading, arm}]; pi {reading: worth}."""
    rated = [r for r in readings if answers[r["number"]][0] is not None]
    for r in rated:
        if pi.get(r["reading"]) is None:
            raise Refused(f"the PI's sheet has no worth for {r['reading']}")
    her = {r["reading"]: answers[r["number"]][0] for r in rated}
    sec = {r["reading"]: answers[r["number"]][1] for r in rated}
    out = {"rated": len(rated), "of": len(readings),
           "numbers_rated": [r["number"] for r in rated],
           "late_among_rated": [r["reading"] for r in rated if r["reading"] in late], "by_arm": {}}
    for arm in ARMS:
        rs = [r["reading"] for r in rated if r["arm"] == arm]
        out["by_arm"][arm] = {
            "rated": len(rs), "of": sum(1 for r in readings if r["arm"] == arm),
            "her_mean": mean([her[x] for x in rs]), "her_counts": counts([her[x] for x in rs]),
            "pi_mean": mean([pi[x] for x in rs]), "pi_counts": counts([pi[x] for x in rs]),
            "second": {s: sum(1 for x in rs if sec[x] == s) for s in SECONDS} | {
                "none given": sum(1 for x in rs if sec[x] is None)}}
    allr = [r["reading"] for r in rated]

    def agree(rs):
        return {"n": len(rs), "same": sum(1 for x in rs if her[x] == pi[x]),
                "within_one": sum(1 for x in rs if abs(her[x] - pi[x]) <= 1)}
    out["reading_by_reading"] = agree(allr)
    two = [x for x in allr if her[x] >= 2]
    out["two_or_more"] = {"n": len(two), "new_to_me": sum(1 for x in two if sec[x] == "new to me")}
    out["second_without_worth"] = sum(1 for r in readings
                                      if answers[r["number"]][0] is None and answers[r["number"]][1] is not None)
    return out


def fmt(x):
    return "—" if x is None else f"{x:.2f}"


def write_md(paper, res, path):
    P = PAPERS[paper]
    L = [f"# The author's own worth for the wager readings of {P['name']}: T34", "",
         "Written by `score_author_ratings.py score`, as `AUTHOR_RATINGS_PLAN.md` and its two dated changes say: "
         "what will be reported, and nothing else. Described, not tested: one rater. The author's answers came "
         "through the PI, the numbers and second answers only, reading by reading. The PI's worth is the PI's of 11 Oct, "
         "which was "
         "not blind" + (f"; {' and '.join(P['late'])} the PI rated late, after the PI had been told the arm." if P["late"]
                        else "."), "",
         f"The author rated {res['rated']} of the {res['of']} readings"
         + (f" (by their numbers in the author's file: {', '.join(map(str, res['numbers_rated']))})." if res["rated"]
            else "."),
         "", "## 1. The author's worth by arm, beside the PI's", "",
         "| arm | rated | the author's mean | the author's 0 / 1 / 2 / 3 | the PI's mean | the PI's 0 / 1 / 2 / 3 |",
         "|---|---|---|---|---|---|"]
    for arm in ARMS:
        a = res["by_arm"][arm]
        hc, pc = a["her_counts"], a["pi_counts"]
        L.append(f"| {arm} | {a['rated']} of {a['of']} | {fmt(a['her_mean'])} | {hc['0']} / {hc['1']} / {hc['2']} / "
                 f"{hc['3']} | {fmt(a['pi_mean'])} | {pc['0']} / {pc['1']} / {pc['2']} / {pc['3']} |")
    rb = res["reading_by_reading"]
    L += ["", "## 2. Reading by reading", "",
          f"The author and the PI give the same worth on {rb['same']} of {rb['n']}, and are within one point on "
          f"{rb['within_one']}. "
          + (f"Among them, {' and '.join(res['late_among_rated'])}, which the PI rated late." if res["late_among_rated"]
             else "None of the readings the PI rated late is among them."), "",
          "## 3. Rated 2 or more", "",
          f"The author rates {res['two_or_more']['n']} readings 2 or more, and marks "
          f"{res['two_or_more']['new_to_me']} of those \"new to me\".", "",
          "## 4. The author's second answers, by arm", "",
          "| arm | " + " | ".join(SECONDS) + " | none given |", "|---|---|---|---|---|"]
    for arm in ARMS:
        s = res["by_arm"][arm]["second"]
        L.append(f"| {arm} | " + " | ".join(str(s[x]) for x in SECONDS) + f" | {s['none given']} |")
    if res["second_without_worth"]:
        L += ["", f"A second answer without a worth: {res['second_without_worth']} reading(s), counted apart and not "
                  "in the figures above."]
    L.append("")
    with open(path, "w", encoding="utf-8") as fh:
        fh.write("\n".join(L))


def build_sheet(paper, out):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    n = PAPERS[paper]["n"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Answers"
    ws.append(["Reading", "The author's worth (0 to 3)", "The author's second answer"])
    for c in range(1, 4):
        ws.cell(1, c).font = Font(bold=True)
        ws.cell(1, c).alignment = Alignment(wrap_text=True, vertical="top")
    fill = PatternFill("solid", fgColor="FFF2CC")
    for i in range(1, n + 1):
        ws.append([i, None, None])
        ws.cell(i + 1, 2).fill = fill
        ws.cell(i + 1, 3).fill = fill
    dv1 = DataValidation(type="list", formula1='"0,1,2,3"', allow_blank=True)
    dv2 = DataValidation(type="list", formula1='"' + ",".join(SECONDS) + '"', allow_blank=True)
    ws.add_data_validation(dv1)
    ws.add_data_validation(dv2)
    dv1.add(f"B2:B{n + 1}")
    dv2.add(f"C2:C{n + 1}")
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 24
    ws2 = wb.create_sheet("About")
    for line in (f"The author's answers for {PAPERS[paper]['name']}, as you pass them on (T34).",
                 "For each reading the author answered, by its number in the author's file: the worth, 0 to 3, and the "
                 "second answer, each chosen from the cell's list.",
                 "Leave a reading blank if the author did not answer it. The author's words stay with you: there is no cell "
                 "for them.",
                 "If an answer fits neither list, leave it blank and tell me in our conversation, by reading number, "
                 "without the author's words."):
        ws2.append([line])
    ws2.column_dimensions["A"].width = 110
    fixed = datetime.datetime(*STAMP)
    wb.properties.created = fixed
    wb.properties.modified = fixed
    wb.properties.creator = "score_author_ratings.py"
    os.makedirs(out, exist_ok=True)
    path = os.path.join(out, f"author_answers_{paper}.xlsx")
    wb.save(path)
    import zipfile
    import re
    with zipfile.ZipFile(path) as z:
        items = [(i, z.read(i.filename)) for i in z.infolist()]
    with zipfile.ZipFile(path, "w") as z:
        for i, data in items:
            if i.filename == "docProps/core.xml":
                data = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)",
                              rb"\g<1>2026-10-15T00:00:00Z\g<2>", data)
            j = zipfile.ZipInfo(i.filename, date_time=STAMP)
            j.compress_type = i.compress_type
            j.external_attr = i.external_attr
            z.writestr(j, data)
    return path


def score(paper, answers_path, out):
    P = PAPERS[paper]
    key = json.load(open(os.path.join(HERE, P["key"]), encoding="utf-8"))
    readings = key["readings"]
    if [r["number"] for r in readings] != list(range(1, P["n"] + 1)):
        raise Refused("the key does not number the readings 1 to n")
    answers = load_answers(answers_path, P["n"])
    pi = load_pi(paper)
    res = figures(answers, readings, pi, P["late"])
    res.update({"paper": paper, "key_sha256": sha(os.path.join(HERE, P["key"])),
                "pi_sheet_sha256": sha(os.path.join(HERE, PI_SHEET)), "answers_sha256": sha(answers_path),
                "plan_sha256": sha(os.path.join(HERE, "AUTHOR_RATINGS_PLAN.md"))})
    os.makedirs(out, exist_ok=True)
    write_md(paper, res, os.path.join(out, f"AUTHOR_RATINGS_{paper.upper()}.md"))
    with open(os.path.join(out, f"author_ratings_{paper}.json"), "w", encoding="utf-8") as fh:
        json.dump(res, fh, ensure_ascii=False, indent=1)
        fh.write("\n")
    return res


# ---------------------------------------------------------------- self-test

def selftest():
    import tempfile
    fails = []

    def check(cond, msg):
        if not cond:
            fails.append(msg)
    # a tiny made-up case, figures by hand
    readings = [{"number": 1, "reading": "X1", "arm": "system"}, {"number": 2, "reading": "X2", "arm": "middle"},
                {"number": 3, "reading": "X3", "arm": "middle"}, {"number": 4, "reading": "X4", "arm": "plain"},
                {"number": 5, "reading": "X5", "arm": "plain"}]
    pi = {"X1": 2, "X2": 0, "X3": 1, "X4": 3, "X5": 2}
    ans = {1: (3, "new to me"), 2: (0, "it is in the paper"), 3: (3, "I had considered it"), 4: (None, "new to me"),
           5: (2, None)}
    f = figures(ans, readings, pi, ("X5",))
    check(f["rated"] == 4 and f["numbers_rated"] == [1, 2, 3, 5], "rated")
    check(f["by_arm"]["middle"]["her_mean"] == 1.5 and f["by_arm"]["middle"]["pi_mean"] == 0.5, "means by arm")
    check(f["by_arm"]["middle"]["her_counts"] == {"0": 1, "1": 0, "2": 0, "3": 1}, "counts by arm")
    check(f["by_arm"]["plain"]["rated"] == 1 and f["by_arm"]["plain"]["of"] == 2, "a part rated")
    # X1 3 against 2, X2 0 against 0, X3 3 against 1, X5 2 against 2 (late)
    check(f["reading_by_reading"] == {"n": 4, "same": 2, "within_one": 3},
          f"same and within one: {f['reading_by_reading']}")
    check(f["late_among_rated"] == ["X5"], "the late one marked")
    check(f["two_or_more"] == {"n": 3, "new_to_me": 1}, "two or more")
    check(f["by_arm"]["plain"]["second"]["none given"] == 1 and f["second_without_worth"] == 1,
          "second answers: none given, and one without a worth")
    check("all" not in f and "reading_by_reading_without_late" not in f, "no figure the plan does not list")
    # reading answers
    for bad in (2.5, "2-3", "high", True, 4, -1):
        try:
            read_worth(bad, "x")
            check(False, f"a worth {bad!r} passed")
        except Refused:
            pass
    check(read_worth("2", "x") == 2 and read_worth(3.0, "x") == 3 and read_worth("", "x") is None, "worth read")
    check(read_second("  New to ME ", "x") == "new to me" and read_second("I had considered it", "x") ==
          "I had considered it" and read_second("It is in the paper", "x") == "it is in the paper", "second read")
    try:
        read_second("maybe", "x")
        check(False, "a second answer outside the three passed")
    except Refused:
        pass
    with tempfile.TemporaryDirectory() as d:
        # the sheet: two builds alike; filled, it scores as the JSON does
        p1 = build_sheet("bechky", os.path.join(d, "a"))
        p2 = build_sheet("bechky", os.path.join(d, "b"))
        check(open(p1, "rb").read() == open(p2, "rb").read(), "two sheets differ")
        from openpyxl import load_workbook
        wb = load_workbook(p1)
        ws = wb["Answers"]
        check([ws.cell(r, 1).value for r in range(2, 16)] == list(range(1, 15)), "the sheet's numbers")
        made = {1: (2, "new to me"), 2: (3, "I had considered it"), 3: (1, "it is in the paper"), 4: (0, None),
                5: (2, "new to me")}
        for i, (w, s) in made.items():
            ws.cell(i + 1, 2).value = str(w)
            ws.cell(i + 1, 3).value = s
        filled = os.path.join(d, "filled.xlsx")
        wb.save(filled)
        jpath = os.path.join(d, "answers.json")
        json.dump({"answers": [{"number": i, "worth": w, "second": s} for i, (w, s) in made.items()]},
                  open(jpath, "w"))
        r1 = score("bechky", filled, os.path.join(d, "o1"))
        r2 = score("bechky", jpath, os.path.join(d, "o2"))
        for k in ("by_arm", "reading_by_reading", "two_or_more", "rated"):
            check(r1[k] == r2[k], f"sheet and JSON differ on {k}")
        # by the key: 1 B08 middle, 2 B10 system, 3 B14 middle, 4 B06 plain, 5 B01 middle;
        # the PI's sheet: B08 0, B10 2, B14 2, B06 2 (late), B01 0
        check(r1["by_arm"]["middle"]["rated"] == 3 and r1["by_arm"]["middle"]["her_mean"] == 5 / 3 and
              r1["by_arm"]["middle"]["pi_mean"] == 2 / 3, f"Bechky middle: {r1['by_arm']['middle']}")
        check(r1["by_arm"]["system"]["her_mean"] == 3 and r1["by_arm"]["system"]["pi_mean"] == 2, "Bechky system")
        # her 2, 3, 1, 0, 2 against the PI's 0, 2, 2, 2, 0: none the same, two within one point
        check(r1["reading_by_reading"] == {"n": 5, "same": 0, "within_one": 2}, f"Bechky: {r1['reading_by_reading']}")
        check(r1["late_among_rated"] == ["B06"], "late marked")
        check(r1["two_or_more"] == {"n": 3, "new_to_me": 2}, "Bechky two or more")
        md = open(os.path.join(d, "o1", "AUTHOR_RATINGS_BECHKY.md"), encoding="utf-8").read()
        check("The author rated 5 of the 14 readings" in md and "B06 the PI rated late" in md, "the report's words")
        bad = os.path.join(d, "bad.json")
        json.dump({"answers": [{"number": 16, "worth": 2}]}, open(bad, "w"))
        try:
            score("bechky", bad, os.path.join(d, "o3"))
            check(False, "a number outside 1 to 14 passed")
        except Refused:
            pass
        rp = score("pratt", jpath, os.path.join(d, "o4"))
        # Pratt: 1 P09 system, 2 P14 plain (late), 3 P12 plain (late), 4 P07 middle, 5 P04 system
        check(rp["by_arm"]["plain"]["rated"] == 2 and rp["late_among_rated"] == ["P14", "P12"], "Pratt")
    if fails:
        print("SELF-TEST FAILED:")
        for x in fails:
            print("  ", x)
        return 1
    print("self-test passed: figures by hand on a made-up case, answers read and refused, the sheet built alike "
          "twice and scored as the JSON, Bechky's and Pratt's keys with the PI's sheet, the report")
    return 0


def main(argv):
    if not argv or argv[0] not in ("selftest", "sheet", "score"):
        print(__doc__)
        return 2
    if argv[0] == "selftest":
        return selftest()
    opts = dict(a[2:].split("=", 1) for a in argv[1:] if a.startswith("--") and "=" in a)
    paper = opts.get("paper")
    if paper not in PAPERS:
        print("--paper=bechky or --paper=pratt")
        return 2
    try:
        if argv[0] == "sheet":
            p = build_sheet(paper, opts["out"])
            print(f"{os.path.basename(p)}: sha256 {sha(p)}")
            return 0
        res = score(paper, opts["answers"], opts["out"])
    except Refused as e:
        print(f"STOPPED: {e}")
        return 1
    print(f"{paper}: {res['rated']} of {res['of']} rated; same worth as the PI on "
          f"{res['reading_by_reading']['same']}, within one on {res['reading_by_reading']['within_one']}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
