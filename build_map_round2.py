#!/usr/bin/env python3
"""
The map test (T39), round 2: the files the placement chats get, the check of their replies, the openings and their
check, and the PI's comparison sheet.

WHY
---
`MAP_TEST_PLAN.md`, part E: Gemini 4 Argon places each of T37's 47 readings on its paper's settled map (chats 77 and
78), given T37's readings file and both T37 coders' codes, and writes each reading's two-sentence opening; GPT 6 checks
each opening against its reading (chats 79 and 80); and the PI compares, on the twelve readings of the PI's T37 sample, T37's
account of each reading with the map's. The plan's order asks that this script be fixed by its sha256 before any chat
of round 2. It judges nothing.

USAGE
-----
    python3 -B build_map_round2.py files <out dir>
        SETTLED_MAP_bechky.md, SETTLED_MAP_pratt.md, T37_CODES_bechky.md, T37_CODES_pratt.md, map_round2_key.json
    python3 -B build_map_round2.py check-placement <paper> <map_round2_key.json> <reply>
        exit 0: counts; 1: does not; 2: counts, with a quoted span to show the PI first
    python3 -B build_map_round2.py openings <map_round2_key.json> <reply 77> <reply 78> <out dir>
        OPENINGS_bechky.md, OPENINGS_pratt.md and map_openings_key.json
    python3 -B build_map_round2.py check-openings <paper> <map_openings_key.json> <reply>
        exit codes as for check-placement
    python3 -B build_map_round2.py comparison --seed=20261026 [--take=3] <map_round2_key.json> <reply 77> <reply 78>
            <out dir>
        MAP_COMPARISON.xlsx (the PI's sheet) and map_comparison_key.json
    python3 -B build_map_round2.py read-comparison <map_comparison_key.json> <filled sheet> <out dir>
        map_comparison_answers.json (the blank sheet is read from beside the key)
    python3 -B build_map_round2.py --selftest

Every source is checked against its sha256, and `openings`, `comparison` and `read-comparison` refuse a key whose
recorded inputs are not the sources as they are. A reply is used by `openings` and `comparison` only if it counts by
`check-placement` with no quoted span of more than eight words (a reply with such a span is filed without it, and the
filed copy is the one given here). Nothing is written over, and the same inputs always give the same bytes.

THE SOURCES
-----------
The two frozen indexes; the settled maps (`map_settled_<paper>.json`) and their closures (`map_closure_<paper>.json`);
T37's readings files (`RELATION_READINGS_<paper>.md`), which the placement chats get as they are; T37's coding task
(`RELATION_CODING.md`), whose section on the six types the codes file gives as it stands; T37's key
(`relation_key.json`: the readings' texts, which must equal the files', and the PI's sample) and its results
(`relation_results.json`: both coders' codes).

THE MAP'S ELEMENTS, WITH THEIR IDS
----------------------------------
The claims, as the index numbers them; the edges, E01 onward, the settled edges sorted by the number in the claim's id
and then in the base's; the constructs, K01 onward, one for each canonical construct of the settled registry and one
for each construct name in no group, sorted by name (case aside), each with its names and the claims that use them;
the factors, by the settled map's ids (X1 onward); the units of the comparative table, U1 onward, in the settled
map's order; the contrast pairs, declared gaps and absences, as the index numbers them. The passages (F..) are the
evidence, not elements.

THE FILES OF THE PLACEMENT CHATS
--------------------------------
`SETTLED_MAP_<paper>.md`: the citation; the claims (pages, core, stated as, statement, outcome, constructs each with
its construct's id, passages by direction, and the edges into and out of the claim); the edges ("C03 rests on C02 ·
mechanistic, necessary"); the constructs; the factors with the claim stating each; where the map has them, the units
with their outcome and each factor's value; the passages, as the round 1 sheets give them; the contrast pairs, each
with each factor's control; the declared gaps; the absences. No map coder, no choice of the PI's and no model is
named. `T37_CODES_<paper>.md`: the six types as T37's task gives them, then for each reading in the file's order
coder A's code (main claim and type, other claims and types, page, reason), then coder B's: coder A is T37's coder
`gemini` (chats 64 and 66), coder B its coder `astra` (65 and 67); the file does not say so.

THE CHECK OF A PLACEMENT REPLY
------------------------------
The text from its first "{" to its last "}". Read leniently first, each a remark: an id in other case or with spaces
(a reading, an element, a passage, a tie); a tie "off map", "offmap" or "none", in any case, for "off-map". Faults: not
one JSON object; a key given twice in one object; no `paper` or `placements`; `placements` not a list of objects; a
reading not the paper's, placed twice, or not placed; `bears_on` not a list of objects, or an element in it that is
neither an element of the map nor a passage, or with no line on `how`; `evidence` not a list of strings; `off_map` not a
non-empty string; an empty `bears_on` with `off_map` "none"; `opening` not two objects, or a sentence empty, or a tie
neither an element, a passage nor "off-map". Remarks: `paper` not naming the author; a passage given as an element or
a tie (kept); an evidence id not a passage of the paper; one element given twice for a reading. In a reply that has no
fault, as in round 1's check, every quoted span of more than eight words in any string is named (where, how many
words), never printed.

THE OPENINGS AND THEIR CHECK
----------------------------
`OPENINGS_<paper>.md`: the citation and a fixed line on what the file holds, then for each reading in order, under
its heading, its text as T37's readings file gives it and its opening, the two sentences of the placement reply joined
by a space, each under a label; nothing else (no map, no tie, no code, no paper). The check of a reply (chats 79 and
80): one JSON object, no key twice; `paper` (a remark if it does not name the author) and `checks`, a list of objects
with every reading once and no other, each with `faithful` true or false ("yes", "no", "true" or "false", in any case,
read as such, a remark) and a non-empty `reason`; in a reply with no fault, every long quoted span named.

THE COMPARISON SHEET
--------------------
The readings of the PI's T37 sample, Bechky's and then Pratt's, each in the sample's order, one tab each; with
`--take=3`, only the first three of each paper's sample in that order, should the PI cut the PI's comparison to six readings
before any chat of round 2. Which account
is account 1 is drawn for each tab in that order with one random.Random(seed): `rng.choice(("T37", "map"))`. T37's
account: for the first coder (T37's `gemini`) and then the second (`astra`), the main claim and each other claim, with
its type and the claim's statement, then the page and the reason. The map's account: each element the reading bears
on, with what it is and the line on how; the labels the closure gives core claims under the first coder's codes, then
the second's; the passages its evidence rests on, with their pages; what is off the map; the opening, each sentence
with its tie. The accounts are not labelled. Each tab has yellow cells for the PI's answer to the two questions (account 1,
account 2, the same, or neither) and a note; the first tab has two for the PI's start and end. `read-comparison` compares
the filled sheet with the blank outside the yellow cells and refuses any other change, and an answer not in the list
(read with case and spacing aside); an empty answer is recorded as unanswered. Each answer is mapped to T37, map, the
same or neither by the key.
"""
import collections
import copy
import datetime
import hashlib
import json
import os
import random
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import build_map_test as bmt  # noqa: E402
import build_map_settle as bms  # noqa: E402
import build_worth_sheet as bws  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402
from openpyxl.styles import Alignment  # noqa: E402

bcs = bmt.bcs
PAPERS = bmt.PAPERS
AUTHOR = bmt.AUTHOR
SHA = {
    "map_settled_bechky.json": "01b7fc1b3045c294af9593573cf9b5ba0f059da901bfd62b1b25998db605380b",
    "map_settled_pratt.json": "323c704a457c2b86357f4f7c370abc961e277a044fb6a43efe7cd6f861736532",
    "map_closure_bechky.json": "7ff976788af4392a40fbe8bd40b0b7b8151690da18b484aa41f19b24680d339f",
    "map_closure_pratt.json": "dde21e59c1423aa3877bc6dcd6254c7fec6fe013c4274196e3e5723ab91d6ea7",
    "relation_key.json": "b8598261afbe915c4d5c58832c7d4516f7500b7f87493c531cf1682490937ea5",
    "relation_results.json": "5dff925eb5d23484ae44642628f29fd12e28366ef4db1e5fc907c4ad97a1fc47",
    "RELATION_READINGS_bechky.md": "20cadee86efc3d6a29dcc6d4ac29c1006fe6527fb21612a561e7e4d6a378cbf1",
    "RELATION_READINGS_pratt.md": "c6de0290d9fedd02ddc4a92c465940f5eb8c4985ae92e4d864bede0cd994f580",
    "RELATION_CODING.md": "be24fa430a3eb9b10ae1c83348935250b66f3e038876e6e32452263ed72fba30",
}
CODERS = (("A", "gemini"), ("B", "astra"))
ORDINAL = {"gemini": "the first coder", "astra": "the second coder"}
SEED = 20261026
ACCOUNTS = ("T37", "map")
ANSWERS = ("account 1", "account 2", "the same", "neither")
QUESTIONS = ("Which account tells you more about what the reading does to the paper's argument?",
             "Which account could another reader check more readily against the paper?")
OFF = "off-map"
OFF_WORDS = ("off-map", "off map", "offmap", "none")
YES_NO = {"yes": True, "true": True, "no": False, "false": False}
KIND_ORDER = ("claim", "edge", "construct", "factor", "unit", "contrast pair", "declared gap", "absence")
ANSWER_ROWS = (6, 8, 9)
TIME_ROWS = (7, 8)
num = bmt.num
sha_bytes = bmt.sha_bytes
sha = bmt.sha
nonempty = bmt.nonempty
st = bmt.st
plain = bmt.plain
dumps = bms.dumps


# ---------------------------------------------------------------- the sources

def read_src(name, folder=None):
    data = open(os.path.join(folder or HERE, name), "rb").read()
    if folder is None and sha_bytes(data) != SHA[name]:
        raise SystemExit(f"{name}: sha256 {sha_bytes(data)}, not {SHA[name]}")
    return data


def readings_of(text):
    """{reading: text} in the file's order."""
    parts = re.split(r"^## Reading (R\d\d)\n", text, flags=re.M)
    out = collections.OrderedDict()
    for i in range(1, len(parts), 2):
        if parts[i] in out:
            raise SystemExit(f"reading {parts[i]} twice in a readings file")
        out[parts[i]] = parts[i + 1].strip()
    return out


def load_sources(folder=None):
    s = {"index": {}, "settled": {}, "closure": {}, "readings": {}, "inputs": {}}
    for paper in PAPERS:
        s["index"][paper], s["inputs"][bmt.INDEX[paper]] = bmt.load_index(paper, folder)
        for kind, name in (("settled", f"map_settled_{paper}.json"), ("closure", f"map_closure_{paper}.json")):
            data = read_src(name, folder)
            s[kind][paper] = json.loads(data)
            s["inputs"][name] = sha_bytes(data)
        name = f"RELATION_READINGS_{paper}.md"
        data = read_src(name, folder)
        s["inputs"][name] = sha_bytes(data)
        s["readings"][paper] = readings_of(data.decode("utf-8"))
    for name in ("relation_key.json", "relation_results.json", "RELATION_CODING.md"):
        data = read_src(name, folder)
        s["inputs"][name] = sha_bytes(data)
        s[name] = data.decode("utf-8") if name.endswith(".md") else json.loads(data)
    check_sources(s)
    return s


def check_sources(s):
    rkey, codes = s["relation_key.json"], s["relation_results.json"]["codes"]
    for paper in PAPERS:
        claims = bmt.structure(s["index"][paper])[0]
        order = list(s["readings"][paper])
        if s["settled"][paper].get("paper") != paper or s["closure"][paper].get("paper") != paper:
            raise SystemExit(f"{paper}: the settled map or the closure is not this paper's")
        if [r["reading"] for r in s["closure"][paper]["readings"]] != order:
            raise SystemExit(f"{paper}: the closure's readings are not the readings file's")
        texts = {r["r"]: r["text"].strip() for r in rkey["readings"] if r["paper"] == paper}
        if texts != dict(s["readings"][paper]):
            raise SystemExit(f"{paper}: the readings file and T37's key differ")
        if any(r not in order for r in rkey["sample"][paper]):
            raise SystemExit(f"{paper}: a reading of the PI's sample not in the file")
        for _, coder in CODERS:
            for r in order:
                c = codes[coder].get(f"{paper}|{r}")
                if c is None:
                    raise SystemExit(f"{paper} {r}: no code of {coder}")
                for cl in [c["main_claim"]] + [x[0] for x in others_of(c)]:
                    if cl != "none" and cl not in claims:
                        raise SystemExit(f"{paper} {r}: {coder} codes a claim not the paper's")
    if "## The six types\n" not in s["RELATION_CODING.md"] or "## An example" not in s["RELATION_CODING.md"]:
        raise SystemExit("RELATION_CODING.md: no section on the six types")


def others_of(c):
    out = []
    for x in c.get("other_claims") or []:
        out.append((x["claim"], x["type"]) if isinstance(x, dict) else (x[0], x[1]))
    return out


# ---------------------------------------------------------------- the elements

def construct_groups(settled, index):
    cmap = bms.canon_map(settled, index)
    names, claims = collections.defaultdict(list), collections.defaultdict(list)
    for c in sorted(index["claims"], key=lambda c: num(c["claim_id"])):
        for n in c.get("constructs") or []:
            k = cmap[n]
            if n not in names[k]:
                names[k].append(n)
            if c["claim_id"] not in claims[k]:
                claims[k].append(c["claim_id"])
    return [{"canonical": k, "names": sorted(names[k], key=bms.name_key), "claims": claims[k]}
            for k in sorted(names, key=bms.name_key)]


def elements(index, settled):
    claims, _, _, _, _ = bmt.structure(index)
    out = collections.OrderedDict()

    def put(i, v):
        if i in out:
            raise SystemExit(f"element id {i} twice")
        out[i] = v
    for c in sorted(claims, key=num):
        put(c, {"kind": "claim", "what": claims[c]["statement"], "core": claims[c].get("core") is True})
    edges = sorted(settled["edges"], key=lambda e: (num(e["claim"]), num(e["base"])))
    if len({(e["base"], e["claim"]) for e in edges}) != len(edges):
        raise SystemExit("two settled edges on one pair")
    for i, e in enumerate(edges, 1):
        put("E%02d" % i, {"kind": "edge", "claim": e["claim"], "base": e["base"], "type": e["type"],
                          "strength": e["strength"], "what": f"{e['claim']} rests on {e['base']}"})
    for i, g in enumerate(construct_groups(settled, index), 1):
        put("K%02d" % i, {"kind": "construct", "what": g["canonical"], "names": g["names"], "claims": g["claims"]})
    for f in settled["factors"]:
        put(f["id"], {"kind": "factor", "what": f["name"], "claim": f["claim"]})
    for i, u in enumerate(settled["units"], 1):
        put("U%d" % i, {"kind": "unit", "what": u["unit"], "claim": u["claim"], "outcome": u["outcome"],
                        "values": u["values"]})
    ctl = {c["contrast"]: c["controls"] for c in settled["contrasts"]}
    for p in index.get("contrast_pairs", []):
        put(p["pair_id"], {"kind": "contrast pair", "what": p["differs_in"], "controls": ctl.get(p["pair_id"], {})})
    for g in index.get("declared_gaps", []):
        put(g["gap_id"], {"kind": "declared gap", "what": g["statement"]})
    for a in index.get("absences", []):
        put(a["absence_id"], {"kind": "absence", "what": a["basis"]})
    return out


def of_kind(E, kind):
    return [k for k, v in E.items() if v["kind"] == kind]


# ---------------------------------------------------------------- the placement chats' files

MAP_INTRO = ("The settled map of the paper's argument: its claims and the edges between them, the constructs the claims "
             "name, the paper's explanatory factors, and its passages, contrast pairs, declared gaps and absences, as "
             "the paper's index records them. Every record is a summary, not the paper's words: check the page. "
             "Every element has an id; place each reading on these ids.")
EDGE_NOTE = ("\"Cj rests on Ci\": with the paper's support for Ci set aside, the paper as printed no longer gives a "
             "reason to hold Cj as stated. Each edge has a type and a strength, as the placement task defines them.")
CODES_INTRO = ("For each reading of the readings file, in its order, two coders' codes, coder A's and then coder B's: "
               "the claim the reading bears on most directly (its main claim) and what the reading does to it (its "
               "main type), up to two more claims with what it does to each, the page the main type rests on, and the "
               "coder's reason. The claims are numbered as in the map. A code is a coder's reading of the reading, not "
               "the paper's words.")


def settled_md(index, settled, E):
    claims, frags, _, _, _ = bmt.structure(index)
    kid = {n: k for k in of_kind(E, "construct") for n in E[k]["names"]}
    into, outof = collections.defaultdict(list), collections.defaultdict(list)
    for k in of_kind(E, "edge"):
        into[E[k]["claim"]].append(k)
        outof[E[k]["base"]].append(k)
    links = collections.defaultdict(lambda: collections.defaultdict(list))
    for l in index["claim_links"]:
        links[l["claim_id"]][l["direction"]].append(l["fragment_id"])
    L = [f"# {index['meta']['citation']}", "", MAP_INTRO, "", f"## The claims ({len(claims)})", ""]
    for c in sorted(claims, key=num):
        cl = claims[c]
        tags = (["core"] if cl.get("core") is True else []) + ([cl["stated_as"]] if cl.get("stated_as") else [])
        L.append(f"- **{c}** · {cl.get('location') or 'no page'} · {', '.join(tags)}")
        L.append(f"  {cl['statement']}")
        L.append(f"  *outcome:* {cl.get('outcome') or 'none'}")
        L.append(f"  *constructs:* {'; '.join(f'{n} ({kid[n]})' for n in cl.get('constructs') or []) or 'none'}")
        by = links[c]
        parts = [f"{k}: {', '.join(sorted(set(by[k]), key=num))}" for k in bmt.DIRECTIONS if by.get(k)]
        parts += [f"{k}: {', '.join(sorted(set(by[k]), key=num))}" for k in sorted(by) if k not in bmt.DIRECTIONS]
        L.append(f"  *passages:* {' · '.join(parts) if parts else 'none'}")
        rests = ", ".join(f"{e} (on {E[e]['base']})" for e in into[c]) or "none"
        rested = ", ".join(f"{e} ({E[e]['claim']})" for e in outof[c]) or "none"
        L.append(f"  *edges:* it rests on: {rests} · rests on it: {rested}")
    edges = of_kind(E, "edge")
    L += ["", f"## The edges ({len(edges)})", "", EDGE_NOTE, ""]
    for k in edges:
        L.append(f"- **{k}** · {E[k]['claim']} rests on {E[k]['base']} · {E[k]['type']}, {E[k]['strength']}")
    if not edges:
        L.append("None.")
    ks = of_kind(E, "construct")
    L += ["", f"## The constructs ({len(ks)})", "",
          "Each construct the claims name: its name, the names the claims use for it, and the claims that use them.", ""]
    for k in ks:
        L.append(f"- **{k}** · {E[k]['what']} · names: {'; '.join(E[k]['names'])} · claims: {', '.join(E[k]['claims'])}")
    xs = of_kind(E, "factor")
    L += ["", f"## The factors ({len(xs)})", "", "The paper's own explanatory factors, each with the claim that states it.",
          ""]
    for k in xs:
        L.append(f"- **{k}** · {E[k]['what']} · stated in {E[k]['claim']}")
    if not xs:
        L.append("None.")
    us = of_kind(E, "unit")
    if us:
        L += ["", f"## The units ({len(us)})", "",
              "The units of the paper's comparative table: each with the claim that names it, whether it shows the "
              "outcome the paper explains across its units (shows it, does not, or not said), and each factor's value "
              "in it (present, absent or not said).", ""]
        for k in us:
            v = E[k]
            vals = "; ".join(f"{x} {bms.VALUE_WORD[v['values'].get(x, '?')]}" for x in xs)
            L.append(f"- **{k}** · {v['what']} · {v['claim']} · outcome: {bms.OUTCOME_WORD[v['outcome']]}"
                     + (f" · {vals}" if vals else ""))
    L += ["", f"## The passages ({len(frags)})", ""]
    for f in sorted(frags, key=num):
        fr = frags[f]
        head = f"- **{f}** · {fr.get('location') or 'no page'} · {fr.get('unit') or 'no unit'} · " \
               f"{fr.get('data_type') or 'no data type'}"
        if fr.get("speaker_name"):
            head += f" · speaker: {fr['speaker_name']}"
        L.append(head)
        L.append(f"  {fr.get('stub') or ''}")
        if fr.get("author_gloss"):
            L.append(f"  *gloss:* {fr['author_gloss']}")
    L += ["", "## The contrast pairs", ""]
    for p in index.get("contrast_pairs", []):
        L.append(f"- **{p['pair_id']}** · drawn by the paper: {'yes' if p.get('drawn_by_authors') else 'no'}")
        L.append(f"  *a:* {p['episode_a']} ({', '.join(p.get('fragments_a') or []) or 'no passage'})")
        L.append(f"  *b:* {p['episode_b']} ({', '.join(p.get('fragments_b') or []) or 'no passage'})")
        L.append(f"  *held comparable:* {p['held_comparable']}")
        L.append(f"  *differs in:* {p['differs_in']}")
        ctl = E[p["pair_id"]]["controls"]
        if xs:
            L.append(f"  *factors:* {'; '.join(f'{x} {ctl.get(x, chr(63))}' for x in xs)}")
    if not index.get("contrast_pairs"):
        L.append("None.")
    L += ["", "## The declared gaps", ""]
    for g in index.get("declared_gaps", []):
        L.append(f"- **{g['gap_id']}** · {g.get('location') or 'no page'} · {g.get('kind') or ''} · bears on "
                 f"{', '.join(g.get('bears_on') or []) or 'no claim'}")
        L.append(f"  {g['statement']}")
    if not index.get("declared_gaps"):
        L.append("None.")
    L += ["", "## The absences", ""]
    for a in index.get("absences", []):
        L.append(f"- **{a['absence_id']}** · {a.get('kind') or ''} · {a.get('evidence_class') or ''} · bears on "
                 f"{', '.join(a.get('bears_on') or []) or 'no claim'}")
        L.append(f"  {a['basis']}")
    if not index.get("absences"):
        L.append("None.")
    L.append("")
    return "\n".join(L).encode("utf-8")


def types_section(coding):
    a = coding.index("## The six types\n") + len("## The six types\n")
    return coding[a:coding.index("## An example")].strip()


def code_line(c):
    s = "main claim none" if c["main_claim"] == "none" else f"main claim {c['main_claim']}, {c['main_type']}"
    others = [f"{cl}, {ty}" for cl, ty in others_of(c)]
    if others:
        s += "; also " + "; ".join(others)
    if c.get("page"):
        s += f". Page: {c['page']}"
    return s + f". Reason: {c['reason']}"


def codes_md(paper, citation, order, codes, coding):
    L = [f"# Two coders' codes of the readings of {citation}", "", CODES_INTRO, "", "## The six types", "",
         types_section(coding), "", "## The codes", ""]
    for r in order:
        L += [f"### {r}", ""]
        for label, coder in CODERS:
            L.append(f"- **Coder {label}:** {code_line(codes[coder][f'{paper}|{r}'])}")
        L.append("")
    return "\n".join(L).encode("utf-8")


def builder_sha():
    return sha(os.path.abspath(__file__))


def files_cmd(outdir, folder=None):
    s = load_sources(folder)
    names = [f"SETTLED_MAP_{p}.md" for p in PAPERS] + [f"T37_CODES_{p}.md" for p in PAPERS] + ["map_round2_key.json"]
    bms.refuse_existing(outdir, names)
    files = collections.OrderedDict()
    key = {"what": "The key of the map test's round 2 (T39, MAP_TEST_PLAN.md, part E): the map's elements with their "
                   "ids, the passages, the readings and the PI's sample, for the checks and the PI's comparison sheet. Not "
                   "for any chat.",
           "builder_sha256": builder_sha(), "inputs": s["inputs"], "coders": dict(CODERS), "papers": {}}
    for paper in PAPERS:
        index, settled = s["index"][paper], s["settled"][paper]
        E = elements(index, settled)
        order = list(s["readings"][paper])
        citation = index["meta"]["citation"]
        m = settled_md(index, settled, E)
        c = codes_md(paper, citation, order, s["relation_results.json"]["codes"], s["RELATION_CODING.md"])
        files[f"SETTLED_MAP_{paper}.md"], files[f"T37_CODES_{paper}.md"] = m, c
        claims, frags, _, _, core = bmt.structure(index)
        key["papers"][paper] = {
            "citation": citation, "core": core, "readings": order,
            "sample": s["relation_key.json"]["sample"][paper], "elements": E,
            "counts": dict(collections.Counter(v["kind"] for v in E.values())),
            "passages": {f: frags[f].get("location") or None for f in sorted(frags, key=num)},
            "files": {f"SETTLED_MAP_{paper}.md": sha_bytes(m), f"T37_CODES_{paper}.md": sha_bytes(c)}}
    files["map_round2_key.json"] = dumps(key)
    for name, data in files.items():
        bmt.write_new(os.path.join(outdir, name), data)
    return files


# ---------------------------------------------------------------- the check of a placement reply

def norm_id(x):
    return re.sub(r"\s+", "", x).upper() if isinstance(x, str) else x


def is_none(s):
    return isinstance(s, str) and plain(s).strip(" .") == "none"


def normalized_placement(rep, kp):
    """(a copy of the reply with forms read leniently, remarks)."""
    if not isinstance(rep, dict):
        return rep, []
    r = copy.deepcopy(rep)
    remarks = []
    known = set(kp["elements"]) | set(kp["passages"]) | set(kp["readings"])

    def fix(obj, field, where, tie=False):
        v = obj.get(field)
        if not isinstance(v, str) or v in known or (tie and v == OFF):
            return
        if tie and plain(v) in OFF_WORDS:
            obj[field] = OFF
            remarks.append(f"{where}: tie {v!r} read as {OFF!r}")
            return
        w = norm_id(v)
        if w in known:
            obj[field] = w
            remarks.append(f"{where}: {field} {v!r} read as {w!r}")
    pl = r.get("placements")
    for i, p in enumerate(pl if isinstance(pl, list) else [], 1):
        if not isinstance(p, dict):
            continue
        fix(p, "reading", f"placement {i}")
        where = st(p.get("reading"))
        for b in p["bears_on"] if isinstance(p.get("bears_on"), list) else []:
            if isinstance(b, dict):
                fix(b, "element", where)
        ev = p.get("evidence")
        for j, x in enumerate(ev if isinstance(ev, list) else []):
            if isinstance(x, str) and x not in known and norm_id(x) in known:
                ev[j] = norm_id(x)
                remarks.append(f"{where}: evidence {x!r} read as {ev[j]!r}")
        for o in p["opening"] if isinstance(p.get("opening"), list) else []:
            if isinstance(o, dict):
                fix(o, "tie", where, tie=True)
    return r, remarks


def placement_faults(rep, paper, kp):
    """(faults, remarks) of one placement reply, by the docstring's rules."""
    E, passages, readings = kp["elements"], kp["passages"], kp["readings"]
    if not isinstance(rep, dict):
        return ["not a JSON object"], []
    missing = [k for k in ("paper", "placements") if k not in rep]
    if missing:
        return [f"no key `{k}`" for k in missing], []
    faults, remarks = [], []
    if not (isinstance(rep["paper"], str) and AUTHOR[paper] in rep["paper"].lower()):
        remarks.append(f"`paper` does not name {AUTHOR[paper].capitalize()}")
    pl = rep["placements"]
    if not isinstance(pl, list) or not all(isinstance(p, dict) for p in pl):
        return ["`placements` is not a list of objects"], remarks
    count = collections.Counter(st(p.get("reading")) for p in pl)
    other = sorted(k for k in count if k not in readings)
    if other:
        faults.append(f"readings not the paper's: {', '.join(other[:8])}{' …' if len(other) > 8 else ''}")
    twice_ = sorted(k for k, n in count.items() if k in readings and n > 1)
    if twice_:
        faults.append(f"readings placed more than once: {', '.join(twice_)}")
    left = [r for r in readings if r not in count]
    if left:
        faults.append(f"readings not placed: {', '.join(left[:8])}{' …' if len(left) > 8 else ''} ({len(left)})")
    for p in pl:
        r = st(p.get("reading"))
        if r not in readings:
            continue
        bo = p.get("bears_on")
        if not isinstance(bo, list) or not all(isinstance(b, dict) for b in bo):
            faults.append(f"{r}: `bears_on` is not a list of objects")
            bo = None
        else:
            seen = collections.Counter()
            for j, b in enumerate(bo, 1):
                el = b.get("element")
                if el in E:
                    seen[el] += 1
                elif el in passages:
                    seen[el] += 1
                    remarks.append(f"{r}: a passage given as an element ({el}), kept")
                else:
                    faults.append(f"{r}: element {j} is not an element of the map")
                if not nonempty(b.get("how")):
                    faults.append(f"{r}: element {j} has no line on how")
            if any(n > 1 for n in seen.values()):
                remarks.append(f"{r}: an element given twice")
        ev = p.get("evidence")
        if not isinstance(ev, list) or not all(isinstance(x, str) for x in ev):
            faults.append(f"{r}: `evidence` is not a list of passage ids")
        elif any(x not in passages for x in ev):
            remarks.append(f"{r}: an evidence id not a passage of the paper")
        om = p.get("off_map")
        if not nonempty(om):
            faults.append(f"{r}: `off_map` is not a non-empty string")
        elif bo == [] and is_none(om):
            faults.append(f"{r}: no element and nothing off the map")
        op = p.get("opening")
        if not isinstance(op, list) or len(op) != 2 or not all(isinstance(o, dict) for o in op):
            faults.append(f"{r}: `opening` is not two objects")
            continue
        for j, o in enumerate(op, 1):
            if not nonempty(o.get("sentence")):
                faults.append(f"{r}: opening sentence {j} is empty")
            t = o.get("tie")
            if t == OFF or t in E:
                continue
            if t in passages:
                remarks.append(f"{r}: opening sentence {j} tied to a passage ({t}), kept")
            else:
                faults.append(f"{r}: opening sentence {j} is tied to no element and not marked off-map")
    return faults, remarks


def kind_of(x, kp):
    return kp["elements"][x]["kind"] if x in kp["elements"] else "passage" if x in kp["passages"] else "?"


def report_spans(rep, cite):
    spans = bcs.long_spans(rep, cite)
    for where, words, unclosed, inside in spans:
        print(f"  QUOTED SPAN of more than {bcs.MAX_QUOTE} words: {where}: {words} words"
              f"{', unclosed' if unclosed else ''}{', within the citation' if inside else ''}")
    if spans:
        print("  COUNTS; the spans above go to the PI before the reply is filed, and are not filed")
        return 2
    print(f"  COUNTS; no quoted span of more than {bcs.MAX_QUOTE} words")
    return 0


def load_reply(path, normalize, faults_fn):
    rep, faults = bcs.load(path)
    remarks = []
    if not faults:
        dup = bmt.twice(path)
        faults = [f"a key given twice in one object: {', '.join(sorted(set(dup)))}"] if dup else []
        rep, remarks = normalize(rep)
        f2, r2 = faults_fn(rep)
        faults += f2
        remarks += r2
    return rep, faults, remarks


def check_placement(paper, key_path, path):
    """Print the check of one placement reply; return (exit code, the reply as read)."""
    kp = json.load(open(key_path, encoding="utf-8"))["papers"][paper]
    rep, faults, remarks = load_reply(path, lambda r: normalized_placement(r, kp),
                                      lambda r: placement_faults(r, paper, kp))
    print(f"{paper} placement reply: {os.path.basename(path)}, sha256 {sha(path)}")
    for f in faults:
        print(f"  FAULT: {f}")
    if faults:
        print("  DOES NOT COUNT")
        return 1, None
    for r in sorted(set(remarks)):
        print(f"  remark: {r}")
    pl = rep["placements"]
    kinds = collections.Counter(kind_of(b["element"], kp) for p in pl for b in p["bears_on"])
    print(f"  placements {len(pl)}; elements placed {sum(kinds.values())}, by kind "
          f"{dict(sorted(kinds.items(), key=lambda kv: (KIND_ORDER + ('passage',)).index(kv[0])))}")
    print(f"  readings with something off the map {sum(not is_none(p['off_map']) for p in pl)}; with evidence "
          f"{sum(bool(p['evidence']) for p in pl)}; opening sentences tied off the map "
          f"{sum(o['tie'] == OFF for p in pl for o in p['opening'])} of {2 * len(pl)}")
    return report_spans(rep, kp["citation"]), rep


def counted(paper, key_path, path):
    code, rep = check_placement(paper, key_path, path)
    if code != 0:
        raise SystemExit(f"{path}: not used, since it does not count with no long quoted span (exit {code})")
    return rep


def same_inputs(key, folder=None):
    s = load_sources(folder)
    if key["inputs"] != s["inputs"]:
        raise SystemExit("the key's inputs are not the sources as they are")
    if folder is None and key["builder_sha256"] != builder_sha():
        raise SystemExit("the key was written by another version of this script")
    return s


# ---------------------------------------------------------------- the openings and their check

OPEN_INTRO = ("The readings of this paper, in their order, each followed by its opening: two sentences another reader "
              "wrote, the first on what the reading opens in the paper's argument, the second on the next "
              "investigation that would settle it.")


def openings_md(citation, order, texts, by):
    L = [f"# Openings of the readings of {citation}", "", OPEN_INTRO, ""]
    for r in order:
        L += [f"## Reading {r}", "", f"**The reading.** {texts[r]}", "",
              f"**Its opening.** {' '.join(o['sentence'].strip() for o in by[r]['opening'])}", ""]
    return "\n".join(L).encode("utf-8")


def openings_cmd(key_path, paths, outdir, folder=None):
    key = json.load(open(key_path, encoding="utf-8"))
    s = same_inputs(key, folder)
    names = [f"OPENINGS_{p}.md" for p in PAPERS] + ["map_openings_key.json"]
    bms.refuse_existing(outdir, names)
    files = collections.OrderedDict()
    okey = {"what": "The key of the openings' check (T39, round 2, chats 79 and 80): the readings, and the hashes of "
                    "the placement replies the openings come from. Not for any chat.",
            "builder_sha256": builder_sha(), "round2_key_sha256": sha(key_path), "papers": {}}
    for paper, path in zip(PAPERS, paths):
        rep = counted(paper, key_path, path)
        kp = key["papers"][paper]
        by = {p["reading"]: p for p in rep["placements"]}
        data = openings_md(kp["citation"], kp["readings"], s["readings"][paper], by)
        files[f"OPENINGS_{paper}.md"] = data
        okey["papers"][paper] = {"citation": kp["citation"], "readings": kp["readings"], "reply_sha256": sha(path),
                                 "file_sha256": sha_bytes(data)}
    files["map_openings_key.json"] = dumps(okey)
    for name, data in files.items():
        bmt.write_new(os.path.join(outdir, name), data)
    return files


def normalized_openings(rep, okp):
    if not isinstance(rep, dict):
        return rep, []
    r = copy.deepcopy(rep)
    remarks = []
    ch = r.get("checks")
    for i, c in enumerate(ch if isinstance(ch, list) else [], 1):
        if not isinstance(c, dict):
            continue
        v = c.get("reading")
        if isinstance(v, str) and v not in okp["readings"] and norm_id(v) in okp["readings"]:
            c["reading"] = norm_id(v)
            remarks.append(f"check {i}: reading {v!r} read as {c['reading']!r}")
        f = c.get("faithful")
        if isinstance(f, str) and plain(f) in YES_NO:
            c["faithful"] = YES_NO[plain(f)]
            remarks.append(f"{st(c.get('reading'))}: faithful {f!r} read as {c['faithful']!r}")
    return r, remarks


def openings_faults(rep, paper, okp):
    if not isinstance(rep, dict):
        return ["not a JSON object"], []
    missing = [k for k in ("paper", "checks") if k not in rep]
    if missing:
        return [f"no key `{k}`" for k in missing], []
    faults, remarks = [], []
    if not (isinstance(rep["paper"], str) and AUTHOR[paper] in rep["paper"].lower()):
        remarks.append(f"`paper` does not name {AUTHOR[paper].capitalize()}")
    ch = rep["checks"]
    if not isinstance(ch, list) or not all(isinstance(c, dict) for c in ch):
        return ["`checks` is not a list of objects"], remarks
    readings = okp["readings"]
    count = collections.Counter(st(c.get("reading")) for c in ch)
    other = sorted(k for k in count if k not in readings)
    if other:
        faults.append(f"readings not the paper's: {', '.join(other[:8])}")
    twice_ = sorted(k for k, n in count.items() if k in readings and n > 1)
    if twice_:
        faults.append(f"readings checked more than once: {', '.join(twice_)}")
    left = [r for r in readings if r not in count]
    if left:
        faults.append(f"readings not checked: {', '.join(left[:8])}{' …' if len(left) > 8 else ''} ({len(left)})")
    for c in ch:
        r = st(c.get("reading"))
        if r not in readings:
            continue
        if not isinstance(c.get("faithful"), bool):
            faults.append(f"{r}: `faithful` not true or false")
        if not nonempty(c.get("reason")):
            faults.append(f"{r}: no reason")
    return faults, remarks


def check_openings(paper, okey_path, path):
    okp = json.load(open(okey_path, encoding="utf-8"))["papers"][paper]
    rep, faults, remarks = load_reply(path, lambda r: normalized_openings(r, okp),
                                      lambda r: openings_faults(r, paper, okp))
    print(f"{paper} openings' check: {os.path.basename(path)}, sha256 {sha(path)}")
    for f in faults:
        print(f"  FAULT: {f}")
    if faults:
        print("  DOES NOT COUNT")
        return 1, None
    for r in sorted(set(remarks)):
        print(f"  remark: {r}")
    yes = sum(c["faithful"] for c in rep["checks"])
    print(f"  checks {len(rep['checks'])}; faithful {yes}; not {len(rep['checks']) - yes}")
    return report_spans(rep, okp["citation"]), rep


# ---------------------------------------------------------------- the comparison sheet

def what_line(eid, v):
    kind = v["kind"]
    if kind == "edge":
        return f"{eid} (edge): {v['claim']} rests on {v['base']}, {v['type']}, {v['strength']}"
    if kind == "factor":
        return f"{eid} (factor): {v['what']}, stated in {v['claim']}"
    if kind == "unit":
        return f"{eid} (unit): {v['what']}, named in {v['claim']}"
    if kind == "contrast pair":
        return f"{eid} (contrast pair): differs in {v['what']}"
    return f"{eid} ({kind}): {v['what']}"


def t37_account(paper, r, codes, E):
    L = []
    for _, coder in CODERS:
        c = codes[coder][f"{paper}|{r}"]
        L.append(ORDINAL[coder][0].upper() + ORDINAL[coder][1:] + ":")
        if c["main_claim"] == "none":
            L.append("- main claim: none")
        else:
            L.append(f"- main claim {c['main_claim']}, {c['main_type']}: {E[c['main_claim']]['what']}")
        for cl, ty in others_of(c):
            L.append(f"- also {cl}, {ty}: {E[cl]['what']}")
        L.append(f"- page: {c.get('page') or 'none given'}")
        L.append(f"- reason: {c['reason']}")
        L.append("")
    return "\n".join(L).strip()


def map_account(p, closure_r, kp):
    E, passages, core = kp["elements"], kp["passages"], set(kp["core"])
    L = ["Where it bears on the map:"]
    for b in p["bears_on"]:
        el = b["element"]
        L.append("- " + (what_line(el, E[el]) if el in E else f"{el} (passage): {passages[el] or 'no page'}"))
        L.append(f"  how: {b['how'].strip()}")
    if not p["bears_on"]:
        L.append("- on no element")
    L += ["", "The labels its codes give core claims through the map's edges:"]
    for _, coder in CODERS:
        labels = closure_r["coders"][coder]["labels"]
        eff = [f"{c} {labels[c]}" for c in sorted(labels, key=num) if c in core]
        L.append(f"- under {ORDINAL[coder]}'s codes: {'; '.join(eff) if eff else 'none'}")
    ev = [f"{f} ({passages[f] or 'no page'})" if f in passages else f"{f} (not a passage of the paper)"
          for f in p["evidence"]]
    L += ["", "Where its evidence sits:", "- " + ("; ".join(ev) if ev else "no passage found")]
    L += ["", "Off the map:", "- " + ("nothing" if is_none(p["off_map"]) else p["off_map"].strip())]
    L += ["", "Its opening:"]
    for o in p["opening"]:
        L.append(f"- {o['sentence'].strip()} (tied to {o['tie']})")
    return "\n".join(L)


INSTRUCTIONS = [
    ("What it is", "For each of {n} readings, one tab: the reading, as its file gives it, and two accounts of how it "
                   "relates to the paper's argument. One is built from two coders' codes of the reading on the paper's "
                   "claims; the other places the reading on the settled map of the paper's argument. Which is account "
                   "1 was drawn for each tab; the accounts are not labelled, though their form shows which is which."),
    ("Question 1", QUESTIONS[0]),
    ("Question 2", QUESTIONS[1]),
    ("Fill in", "Only the yellow cells: on each tab, your answer to each question from the list (account 1, account 2, "
                "the same, or neither), and a note if you like, which is welcome and never needed."),
    ("Time", "If you can, type the time you start and the time you finish, as hh:mm, in the two yellow cells below; "
             "or just tell me roughly how long it took."),
    ("Started at", ""),
    ("Finished at", ""),
    ("Save", "Save it as MAP_COMPARISON_filled.xlsx, with no other change, and send it in our conversation."),
]


def comparison_book(tabs, n):
    wb = Workbook()
    ins = wb.active
    ins.title = "Instructions"
    ins.column_dimensions["A"].width = 18
    ins.column_dimensions["B"].width = 110
    bws.cell(ins, "A1", "The map test: your comparison of two accounts of each reading", bold=True, size=12, wrap=False)
    for r, (label, t) in enumerate(INSTRUCTIONS, 2):
        bws.cell(ins, f"A{r}", label, bold=True)
        c = bws.cell(ins, f"B{r}", t.format(n=n), fill=bws.YELLOW if r in TIME_ROWS else None)
        ins.row_dimensions[r].height = bws.height(t, 110)
        if r in TIME_ROWS:
            c.border = bws.BOX
    bws.printable(ins)
    for t in tabs:
        ws = wb.create_sheet(t["tab"])
        ws.column_dimensions["A"].width = 16
        ws.column_dimensions["B"].width = 120
        bws.cell(ws, "A1", f"{t['tab']}: {t['citation']}", bold=True, size=11, wrap=False)
        rows = [("The reading", t["text"], False), ("Account 1", t["accounts"][0], False),
                ("Account 2", t["accounts"][1], False), ("Question 1", QUESTIONS[0], False),
                ("Your answer", None, True), ("Question 2", QUESTIONS[1], False), ("Your answer", None, True),
                ("Note", None, True)]
        dv = bms.validation(ANSWERS, "Your answer")
        ws.add_data_validation(dv)
        for r, (label, v, yellow) in enumerate(rows, 2):
            bws.cell(ws, f"A{r}", label, bold=True)
            c = bws.cell(ws, f"B{r}", v, fill=bws.YELLOW if yellow else None)
            if yellow:
                c.border = bws.BOX
                if label == "Your answer":
                    dv.add(f"B{r}")
            ws.row_dimensions[r].height = bws.height(v or "", 115, minimum=30.0 if yellow else 15.0)
        bws.printable(ws)
    return wb


def save_xlsx(wb, path):
    wb.properties.created = wb.properties.modified = datetime.datetime(*bms.STAMP)
    wb.properties.creator = "build_map_round2.py"
    wb.save(path)
    bms.steady(path)


def comparison_cmd(seed, key_path, paths, outdir, folder=None, take=None):
    key = json.load(open(key_path, encoding="utf-8"))
    s = same_inputs(key, folder)
    names = ["MAP_COMPARISON.xlsx", "map_comparison_key.json"]
    bms.refuse_existing(outdir, names)
    reps = {paper: counted(paper, key_path, path) for paper, path in zip(PAPERS, paths)}
    rng = random.Random(seed)
    tabs = []
    for paper in PAPERS:
        kp = key["papers"][paper]
        by = {p["reading"]: p for p in reps[paper]["placements"]}
        clo = {x["reading"]: x for x in s["closure"][paper]["readings"]}
        for r in kp["sample"][:take]:
            first = rng.choice(ACCOUNTS)
            acc = {"T37": t37_account(paper, r, s["relation_results.json"]["codes"], kp["elements"]),
                   "map": map_account(by[r], clo[r], kp)}
            order = (first, [a for a in ACCOUNTS if a != first][0])
            tabs.append({"tab": f"{paper.capitalize()} {r}", "paper": paper, "reading": r,
                         "citation": kp["citation"], "text": s["readings"][paper][r], "order": order,
                         "accounts": [acc[a] for a in order]})
    wb = comparison_book(tabs, len(tabs))
    os.makedirs(outdir, exist_ok=True)
    path = os.path.join(outdir, "MAP_COMPARISON.xlsx")
    save_xlsx(wb, path)
    print(f"written: {path} ({os.path.getsize(path)} bytes, sha256 {sha(path)})")
    ckey = {"what": "The key of the PI's comparison sheet (T39, round 2): for each tab its reading and which account is "
                    "account 1, drawn as the docstring says. Not for any chat; not for the PI until the PI's sheet is back.",
            "builder_sha256": builder_sha(), "seed": seed, "take": take, "round2_key_sha256": sha(key_path),
            "replies": {paper: sha(path_) for paper, path_ in zip(PAPERS, paths)}, "sheet_sha256": sha(path),
            "tabs": [{"tab": t["tab"], "paper": t["paper"], "reading": t["reading"], "account_1": t["order"][0],
                      "account_2": t["order"][1], "accounts_sha256": [sha_bytes(a.encode("utf-8"))
                                                                      for a in t["accounts"]]} for t in tabs]}
    bmt.write_new(os.path.join(outdir, "map_comparison_key.json"), dumps(ckey))
    return path


def yellow_of(ckey):
    free = {"Instructions": (set(TIME_ROWS), {2})}
    for t in ckey["tabs"]:
        free[t["tab"]] = (set(ANSWER_ROWS), {2})
    return free


def read_comparison(ckey, filled, blank):
    """(answers, problems) from the PI's filled sheet, by the docstring's rules."""
    faults, wf = bms.compare_cells(filled, blank, yellow_of(ckey))
    if faults:
        return None, faults
    out = {"start": bms.text(wf["Instructions"].cell(TIME_ROWS[0], 2).value),
           "end": bms.text(wf["Instructions"].cell(TIME_ROWS[1], 2).value), "tabs": []}
    problems = []
    for t in ckey["tabs"]:
        ws = wf[t["tab"]]
        row = {"tab": t["tab"], "paper": t["paper"], "reading": t["reading"],
               "note": bms.text(ws.cell(ANSWER_ROWS[2], 2).value)}
        for q, rr in (("q1", ANSWER_ROWS[0]), ("q2", ANSWER_ROWS[1])):
            v = bms.text(ws.cell(rr, 2).value)
            a = plain(v) if v is not None else None
            if a is not None and a not in ANSWERS:
                problems.append(f"{t['tab']}: the answer to {q} does not read")
                continue
            row[q] = a
            row[q + "_pick"] = None if a is None else {"account 1": t["account_1"], "account 2": t["account_2"],
                                                       "the same": "the same", "neither": "neither"}[a]
        out["tabs"].append(row)
    return (None if problems else out), problems


def read_comparison_cmd(ckey_path, filled, outdir, folder=None):
    ckey = json.load(open(ckey_path, encoding="utf-8"))
    blank = os.path.join(os.path.dirname(os.path.abspath(ckey_path)), "MAP_COMPARISON.xlsx")
    if sha(blank) != ckey["sheet_sha256"]:
        raise SystemExit("the blank sheet beside the key is not the one the key was written with")
    bms.refuse_existing(outdir, ["map_comparison_answers.json"])
    out, problems = read_comparison(ckey, filled, blank)
    for p in problems:
        print(f"  PROBLEM: {p}")
    if problems:
        raise SystemExit("the filled sheet is refused; nothing written")
    out.update({"what": "The PI's answers on the PI's comparison sheet (T39, round 2), read by build_map_round2.py "
                        "read-comparison.", "builder_sha256": builder_sha(), "comparison_key_sha256": sha(ckey_path),
                "filled_sha256": sha(filled)})
    bmt.write_new(os.path.join(outdir, "map_comparison_answers.json"), dumps(out))
    return out


# ---------------------------------------------------------------- self-test

def made_up_sources(folder):
    """Made-up sources in the real ones' form, written to folder: two indexes (build_map_test's), settled maps,
    closures, readings files, T37's key, results and coding task."""
    texts = {"bechky": {"R01": "Wards adopt because nurses push, not because of endorsement (p. 12).",
                        "R02": "The budget, not endorsement, decides the speed (p. 14).",
                        "R03": "Ward B never resisted; the paper misreads its delay (p. 18)."},
             "pratt": {"R01": "Identity drives adoption more than endorsement does (p. 15).",
                       "R02": "The contrast pair confounds endorsement with budget (p. 20)."}}
    codes = {"gemini": {}, "astra": {}}
    readings, closure = [], {}
    for paper in PAPERS:
        d = bmt.made_up_index(paper)
        d["claims"][5]["constructs"].append("Night shift")
        open(os.path.join(folder, bmt.INDEX[paper]), "w", encoding="utf-8").write(json.dumps(d))
        settled = {"paper": paper,
                   "edges": [{"base": "C06", "claim": "C04", "type": "comparative", "strength": "necessary"},
                             {"base": "C03", "claim": "C01", "type": "evidential", "strength": "contributing"},
                             {"base": "C01", "claim": "C05", "type": "significance", "strength": "contributing"},
                             {"base": "C02", "claim": "C01", "type": "evidential", "strength": "contributing"}],
                   "registry": [{"canonical": "ward", "names": ["ward A", "ward B"]}],
                   "factors": [{"id": "X1", "name": "endorsement", "claim": "C01"},
                               {"id": "X2", "name": "budget", "claim": "C06"}],
                   "units": [{"unit": "ward A", "claim": "C02", "outcome": "+", "values": {"X1": "+", "X2": "?"}},
                             {"unit": "ward B", "claim": "C03", "outcome": "-", "values": {"X1": "-", "X2": "?"}}]
                   if paper == "bechky" else [],
                   "contrasts": [{"contrast": "CP01", "controls": {"X1": "varies", "X2": "held"}}]}
        open(os.path.join(folder, f"map_settled_{paper}.json"), "w", encoding="utf-8").write(json.dumps(settled))
        rs = []
        for i, r in enumerate(texts[paper], 1):
            codes["gemini"][f"{paper}|{r}"] = {"main_claim": "C01", "main_type": "other_mechanism",
                                               "other_claims": [["C06", "complements"]],
                                               "reason": f"FIRST-REASON-{paper}-{r}", "page": "p. 12"}
            codes["astra"][f"{paper}|{r}"] = {"main_claim": "C02" if i != 2 else "none",
                                              "main_type": "rereads_evidence" if i != 2 else "none",
                                              "other_claims": [], "reason": f"SECOND-REASON-{paper}-{r}", "page": "p. 13"}
            rs.append({"reading": r, "coders": {"gemini": {"labels": {"C01": "contested", "C05": "idle",
                                                                      "C06": "contested"}},
                                                "astra": {"labels": {"C02": "contested", "C01": "weakened"}}}})
            readings.append({"paper": paper, "r": r, "text": texts[paper][r]})
        closure[paper] = {"paper": paper, "readings": rs}
        open(os.path.join(folder, f"map_closure_{paper}.json"), "w", encoding="utf-8").write(json.dumps(closure[paper]))
        body = "".join(f"## Reading {r}\n\n{t}\n\n" for r, t in texts[paper].items())
        open(os.path.join(folder, f"RELATION_READINGS_{paper}.md"), "w", encoding="utf-8").write(
            f"# Readings of Doe 2024\n\n{len(texts[paper])} readings of this paper.\n\n" + body)
    key = {"readings": readings, "sample": {"bechky": ["R01", "R03"], "pratt": ["R02"]}, "planted": []}
    open(os.path.join(folder, "relation_key.json"), "w", encoding="utf-8").write(json.dumps(key))
    open(os.path.join(folder, "relation_results.json"), "w", encoding="utf-8").write(json.dumps({"codes": codes}))
    open(os.path.join(folder, "RELATION_CODING.md"), "w", encoding="utf-8").write(
        "# How\n\nTask.\n\n## The six types\n\nFor one reading and one claim, one type:\n\n- `inverts`: TYPE-INVERTS.\n"
        "- `restates`: TYPE-RESTATES.\n\n**If more than one type fits, give the first.**\n\n## An example\n\n"
        "EXAMPLE-TEXT\n\n## Rules\n\n- Rule.\n")


def good_placement(paper, readings):
    pl = []
    for i, r in enumerate(readings):
        pl.append({"reading": r,
                   "bears_on": [{"element": "C01", "how": "it gives another cause for the adoption"},
                                {"element": "E01", "how": "it cuts the tie from ward A's case"},
                                {"element": "K03", "how": "it splits the construct"}] if i != 1 else [],
                   "evidence": ["F01", "F13"] if i != 1 else [],
                   "off_map": "none" if i != 1 else "the nurses' role, which the map does not hold",
                   "opening": [{"sentence": f"OPENING-ONE-{paper}-{r}.", "tie": "C01" if i != 1 else OFF},
                               {"sentence": f"OPENING-TWO-{paper}-{r}.", "tie": OFF}]})
    return {"paper": "Bechky 2020" if paper == "bechky" else "Pratt 2000", "placements": pl}


def selftest():
    import contextlib
    import io
    import shutil
    import tempfile
    n_ok = n_bad = 0

    def expect(label, cond):
        nonlocal n_ok, n_bad
        if cond:
            n_ok += 1
        else:
            n_bad += 1
            print(f"FAIL: {label}")

    def quiet(fn, *a):
        buf = io.StringIO()
        with contextlib.redirect_stdout(buf):
            try:
                return fn(*a), buf.getvalue()
            except SystemExit as e:
                return ("EXIT", str(e)), buf.getvalue()

    def refused(r, words):
        return isinstance(r, tuple) and r[0] == "EXIT" and words in r[1]

    tmp = tempfile.mkdtemp()
    try:
        src = os.path.join(tmp, "src")
        os.makedirs(src)
        made_up_sources(src)
        out1, out2 = os.path.join(tmp, "o1"), os.path.join(tmp, "o2")
        f1, _ = quiet(files_cmd, out1, src)
        f2, _ = quiet(files_cmd, out2, src)
        expect("files: five files", isinstance(f1, dict) and len(f1) == 5)
        expect("files: the same bytes twice", isinstance(f1, dict) and f1 == f2)
        again, _ = quiet(files_cmd, out1, src)
        expect("files: refuses to write over", refused(again, "refusing"))
        key_path = os.path.join(out1, "map_round2_key.json")
        key = json.load(open(key_path, encoding="utf-8"))
        kb = key["papers"]["bechky"]
        E = kb["elements"]
        expect("elements: the edges sorted by claim, then base",
               [(E[k]["claim"], E[k]["base"]) for k in ("E01", "E02", "E03", "E04")] ==
               [("C01", "C02"), ("C01", "C03"), ("C04", "C06"), ("C05", "C01")])
        expect("elements: the constructs, a group and the names in none, by name",
               [E[k]["what"] for k in ("K01", "K02", "K03", "K04", "K05", "K06")] ==
               ["adoption", "budget", "endorsement", "identity", "Night shift", "ward"] and "K07" not in E)
        expect("elements: a construct's names and claims", E["K06"]["names"] == ["ward A", "ward B"] and
               E["K06"]["claims"] == ["C02", "C03"] and E["K03"]["claims"] == ["C01", "C02", "C03", "C04", "C05"])
        expect("elements: factors, units, contrast pair, gap, absence",
               E["X2"]["claim"] == "C06" and E["U1"]["what"] == "ward A" and E["U2"]["claim"] == "C03" and
               E["CP01"]["controls"] == {"X1": "varies", "X2": "held"} and E["DG01"]["kind"] == "declared gap" and
               E["A01"]["kind"] == "absence")
        expect("elements: no unit where the map has none", "U1" not in key["papers"]["pratt"]["elements"])
        expect("key: the counts", kb["counts"] == {"claim": 6, "edge": 4, "construct": 6, "factor": 2, "unit": 2,
                                                   "contrast pair": 1, "declared gap": 1, "absence": 1})
        expect("key: passages with pages, readings, sample", kb["passages"]["F01"] == "p. 11" and
               kb["readings"] == ["R01", "R02", "R03"] and kb["sample"] == ["R01", "R03"] and
               kb["core"] == ["C01", "C04", "C05"])
        m = f1["SETTLED_MAP_bechky.md"].decode("utf-8")
        expect("map: an edge line", "- **E01** · C01 rests on C02 · evidential, contributing" in m)
        expect("map: a claim's constructs with their ids", "*constructs:* endorsement (K03); ward A (K06)" in m)
        expect("map: a claim's edges", "*edges:* it rests on: E01 (on C02), E02 (on C03) · rests on it: E04 (C05)" in m)
        expect("map: a construct line", "- **K06** · ward · names: ward A; ward B · claims: C02, C03" in m)
        expect("map: a factor line", "- **X2** · budget · stated in C06" in m)
        expect("map: a unit line", "- **U1** · ward A · C02 · outcome: shows it · X1 present; X2 not said" in m)
        expect("map: a contrast pair's controls", "  *factors:* X1 varies; X2 held" in m)
        expect("map: a passage with its gloss and speaker", "- **F01** · p. 11 · ward A · interview · speaker: Ann" in m
               and "*gloss:* a gloss" in m)
        expect("map: gaps and absences", "- **DG01** · p. 20 · scope_limit · bears on C01" in m and
               "- **A01** · not_collected · costs · bears on C06" in m)
        expect("map: no note, code or basis of the index", all(w not in m for w in ("NOTE-X", "NOTE-Y", "CODE-X",
                                                                                     "BASIS-Z")))
        mp = f1["SETTLED_MAP_pratt.md"].decode("utf-8")
        expect("map: no units section where the map has none", "## The units" not in mp and "## The units" in m)
        c = f1["T37_CODES_bechky.md"].decode("utf-8")
        expect("codes: coder A is gemini, coder B astra",
               c.index("FIRST-REASON-bechky-R01") < c.index("SECOND-REASON-bechky-R01") < c.index("### R02") and
               "- **Coder A:** main claim C01, other_mechanism; also C06, complements. Page: p. 12. Reason: "
               "FIRST-REASON-bechky-R01" in c)
        expect("codes: a main claim none", "- **Coder B:** main claim none. Page: p. 13. Reason: SECOND-REASON-bechky-R02"
               in c)
        expect("codes: the six types as the task gives them, not its example",
               "TYPE-INVERTS" in c and "**If more than one type fits, give the first.**" in c and
               "EXAMPLE-TEXT" not in c and "Rule." not in c)
        expect("codes: no model's or chat's name", all(w not in c.lower() for w in ("gemini", "astra", "chat 6")))
        # the sources are checked
        bad = os.path.join(tmp, "bad")
        shutil.copytree(src, bad)
        t = open(os.path.join(bad, "RELATION_READINGS_pratt.md"), encoding="utf-8").read()
        open(os.path.join(bad, "RELATION_READINGS_pratt.md"), "w", encoding="utf-8").write(t.replace("Identity", "Id"))
        r, _ = quiet(files_cmd, os.path.join(tmp, "o3"), bad)
        expect("sources: a readings file unlike T37's key is refused", refused(r, "differ"))
        shutil.rmtree(bad)
        shutil.copytree(src, bad)
        s0 = json.load(open(os.path.join(bad, "map_settled_bechky.json"), encoding="utf-8"))
        s0["edges"].append(dict(s0["edges"][0]))
        open(os.path.join(bad, "map_settled_bechky.json"), "w", encoding="utf-8").write(json.dumps(s0))
        r, _ = quiet(files_cmd, os.path.join(tmp, "o4"), bad)
        expect("sources: two settled edges on one pair are refused", refused(r, "two settled edges"))
        shutil.rmtree(bad)
        shutil.copytree(src, bad)
        k0 = json.load(open(os.path.join(bad, "relation_results.json"), encoding="utf-8"))
        k0["codes"]["astra"]["pratt|R01"]["main_claim"] = "C99"
        open(os.path.join(bad, "relation_results.json"), "w", encoding="utf-8").write(json.dumps(k0))
        r, _ = quiet(files_cmd, os.path.join(tmp, "o5"), bad)
        expect("sources: a code on a claim not the paper's is refused", refused(r, "not the paper's"))
        shutil.rmtree(bad)

        # the check of a placement reply
        def reply(name, obj, raw=None):
            p = os.path.join(tmp, name)
            open(p, "w", encoding="utf-8").write(raw if raw is not None else "Here it is:\n" + json.dumps(obj) + "\nend")
            return p

        def placed(name, obj, paper="bechky", raw=None):
            (code, rep), text = quiet(check_placement, paper, key_path, reply(name, obj, raw))
            return code, rep, text
        g = good_placement("bechky", ["R01", "R02", "R03"])
        code, rep, text = placed("g.json", g)
        expect("placement: a good reply counts", code == 0 and "FAULT" not in text and "remark" not in text)
        expect("placement: its counts", "placements 3; elements placed 6" in text and
               "with something off the map 1; with evidence 2" in text and "tied off the map 4 of 6" in text)
        x = copy.deepcopy(g)
        x["placements"][0]["bears_on"][0]["element"] = "c 01"
        x["placements"][0]["opening"][0]["tie"] = "Off Map"
        x["placements"][2]["opening"][1]["tie"] = "NONE"
        x["placements"][2]["reading"] = "r03"
        x["placements"][0]["evidence"] = ["f01"]
        code, rep, text = placed("l.json", x)
        expect("placement: forms read leniently, each a remark", code == 0 and text.count("remark:") == 5 and
               rep["placements"][0]["bears_on"][0]["element"] == "C01" and rep["placements"][2]["reading"] == "R03"
               and rep["placements"][0]["opening"][0]["tie"] == OFF and rep["placements"][0]["evidence"] == ["F01"])

        def faulty(label, change, want, paper="bechky"):
            y = copy.deepcopy(g if paper == "bechky" else good_placement("pratt", ["R01", "R02"]))
            change(y)
            code, _, text = placed("f.json", y, paper)
            expect(f"placement fault: {label}", code == 1 and want in text)
        faulty("a reading not placed", lambda y: y["placements"].pop(), "readings not placed: R03")
        faulty("a reading twice", lambda y: y["placements"].append(copy.deepcopy(y["placements"][0])),
               "placed more than once: R01")
        faulty("a reading not the paper's", lambda y: y["placements"][2].update(reading="R09"), "not the paper's: R09")
        faulty("an element not the map's", lambda y: y["placements"][0]["bears_on"][1].update(element="E09"),
               "element 2 is not an element of the map")
        faulty("an element with no line on how", lambda y: y["placements"][0]["bears_on"][2].update(how=" "),
               "element 3 has no line on how")
        faulty("bears_on not a list", lambda y: y["placements"][0].update(bears_on="C01"), "`bears_on` is not a list")
        faulty("evidence not a list", lambda y: y["placements"][0].update(evidence="F01"), "`evidence` is not a list")
        faulty("off_map empty", lambda y: y["placements"][0].update(off_map=""), "`off_map` is not a non-empty")
        faulty("nothing placed", lambda y: y["placements"][1].update(off_map="None."), "no element and nothing off")
        faulty("one opening sentence", lambda y: y["placements"][0]["opening"].pop(), "`opening` is not two objects")
        faulty("an empty opening sentence", lambda y: y["placements"][0]["opening"][1].update(sentence=""),
               "opening sentence 2 is empty")
        faulty("a tie to nothing", lambda y: y["placements"][0]["opening"][0].update(tie="Z9"),
               "opening sentence 1 is tied to no element")
        faulty("no placements", lambda y: y.pop("placements"), "no key `placements`")
        faulty("placements not objects", lambda y: y.update(placements=["R01"]), "not a list of objects")
        faulty("the units of another map", lambda y: y["placements"][0]["bears_on"][0].update(element="U1"),
               "element 1 is not an element of the map", paper="pratt")
        code, _, text = placed("d.json", None, raw='{"paper": "Bechky", "paper": "Bechky", "placements": []}')
        expect("placement fault: a key twice", code == 1 and "a key given twice in one object: paper" in text)
        code, _, text = placed("n.json", None, raw="no json here")
        expect("placement fault: no JSON", code == 1 and "no JSON object found" in text)
        y = copy.deepcopy(g)
        y["paper"] = "Doe 2024"
        y["placements"][0]["bears_on"].append({"element": "F02", "how": "a passage"})
        y["placements"][0]["bears_on"].append({"element": "C01", "how": "again"})
        y["placements"][2]["evidence"] = ["F99"]
        y["placements"][2]["opening"][0]["tie"] = "F03"
        code, _, text = placed("r.json", y)
        expect("placement: remarks that do not stop it", code == 0 and "`paper` does not name Bechky" in text and
               "a passage given as an element (F02), kept" in text and "R01: an element given twice" in text and
               "an evidence id not a passage" in text and "tied to a passage (F03), kept" in text)
        y = copy.deepcopy(g)
        y["placements"][0]["bears_on"][0]["how"] = 'it says "one two three four five six seven eight nine ten"'
        code, _, text = placed("s.json", y)
        expect("placement: a long quoted span is named, not printed", code == 2 and "10 words" in text and
               "seven eight" not in text)

        # the openings
        p77 = reply("p77.json", g)
        p78 = reply("p78.json", good_placement("pratt", ["R01", "R02"]))
        oo = os.path.join(tmp, "oo")
        fo, _ = quiet(openings_cmd, key_path, (p77, p78), oo, src)
        expect("openings: three files", isinstance(fo, dict) and len(fo) == 3)
        ob = fo["OPENINGS_bechky.md"].decode("utf-8") if isinstance(fo, dict) else ""
        expect("openings: each reading's text and its opening, nothing else",
               "## Reading R03\n\n**The reading.** Ward B never resisted; the paper misreads its delay (p. 18).\n\n"
               "**Its opening.** OPENING-ONE-bechky-R03. OPENING-TWO-bechky-R03." in ob and "C01" not in ob
               and "FIRST-REASON" not in ob and "how" not in ob)
        fo2, _ = quiet(openings_cmd, key_path, (p77, p78), os.path.join(tmp, "oo2"), src)
        expect("openings: the same bytes twice", fo == fo2)
        ps = reply("ps.json", json.loads(open(os.path.join(tmp, "s.json"), encoding="utf-8").read()
                                         .split("\n", 1)[1].rsplit("\n", 1)[0]))
        r, _ = quiet(openings_cmd, key_path, (ps, p78), os.path.join(tmp, "oo3"), src)
        expect("openings: a reply with a long span is not used", refused(r, "exit 2"))
        pf = reply("pf.json", {"paper": "Bechky", "placements": []})
        r, _ = quiet(openings_cmd, key_path, (pf, p78), os.path.join(tmp, "oo4"), src)
        expect("openings: a reply that does not count is not used", refused(r, "exit 1"))
        okey = os.path.join(oo, "map_openings_key.json")

        def checked(obj, paper="bechky"):
            (code, rep), text = quiet(check_openings, paper, okey, reply("c.json", obj))
            return code, rep, text
        oc = {"paper": "Bechky 2020", "checks": [{"reading": r, "faithful": r != "R02", "reason": "it says so"}
                                                 for r in ("R01", "R02", "R03")]}
        code, _, text = checked(oc)
        expect("openings' check: a good reply counts", code == 0 and "checks 3; faithful 2; not 1" in text and
               "remark" not in text)
        y = copy.deepcopy(oc)
        y["checks"][0]["faithful"] = "Yes"
        y["checks"][1]["reading"] = "r02"
        code, rep, text = checked(y)
        expect("openings' check: yes and ids read leniently", code == 0 and text.count("remark:") == 2 and
               rep["checks"][0]["faithful"] is True)
        for label, change, want in (
                ("a reading not checked", lambda y: y["checks"].pop(), "readings not checked: R03"),
                ("faithful not a truth", lambda y: y["checks"][0].update(faithful="maybe"), "`faithful` not true"),
                ("no reason", lambda y: y["checks"][1].update(reason=""), "R02: no reason"),
                ("a reading twice", lambda y: y["checks"].append(dict(y["checks"][0])), "more than once: R01")):
            y = copy.deepcopy(oc)
            change(y)
            code, _, text = checked(y)
            expect(f"openings' check fault: {label}", code == 1 and want in text)

        # the comparison sheet
        co = os.path.join(tmp, "co")
        path, _ = quiet(comparison_cmd, 7, key_path, (p77, p78), co, src)
        ck = json.load(open(os.path.join(co, "map_comparison_key.json"), encoding="utf-8"))
        rng = random.Random(7)
        want = [rng.choice(("T37", "map")) for _ in range(3)]
        expect("comparison: the tabs, Bechky's sample then Pratt's", [t["tab"] for t in ck["tabs"]] ==
               ["Bechky R01", "Bechky R03", "Pratt R02"])
        expect("comparison: account 1 drawn per tab, in order", [t["account_1"] for t in ck["tabs"]] == want and
               all(t["account_2"] != t["account_1"] for t in ck["tabs"]))
        wb = load_workbook(os.path.join(co, "MAP_COMPARISON.xlsx"))
        expect("comparison: the instructions and one tab per reading", wb.sheetnames ==
               ["Instructions", "Bechky R01", "Bechky R03", "Pratt R02"])
        ws = wb["Bechky R03"]
        accs = {ck["tabs"][1]["account_1"]: ws["B3"].value, ck["tabs"][1]["account_2"]: ws["B4"].value}
        expect("comparison: the reading as its file gives it", ws["B2"].value ==
               "Ward B never resisted; the paper misreads its delay (p. 18).")
        t37, mp_ = accs["T37"], accs["map"]
        expect("comparison: T37's account, the first coder first, with the claims' statements",
               t37.index("The first coder:") < t37.index("FIRST-REASON-bechky-R03") < t37.index("The second coder:")
               < t37.index("SECOND-REASON-bechky-R03") and
               "- main claim C01, other_mechanism: Endorsement leads wards to adopt checklists." in t37 and
               "- also C06, complements: Budgets shape checklist adoption speed." in t37)
        expect("comparison: the map's account, elements with what each is and how",
               "- E01 (edge): C01 rests on C02, evidential, contributing\n  how: it cuts the tie from ward A's case"
               in mp_ and "- K03 (construct): endorsement" in mp_)
        expect("comparison: the labels on core claims only, by coder",
               "- under the first coder's codes: C01 contested; C05 idle" in mp_ and
               "- under the second coder's codes: C01 weakened" in mp_ and "C06 contested" not in mp_)
        expect("comparison: the evidence with pages, off the map, the opening with its ties",
               "- F01 (p. 11); F13 (p. 23)" in mp_ and "Off the map:\n- nothing" in mp_ and
               "- OPENING-ONE-bechky-R03. (tied to C01)" in mp_)
        expect("comparison: no model, chat or coder name", all(w not in (t37 + mp_).lower() for w in
                                                               ("gemini", "astra", "opus", "gpt", "argon")))
        expect("comparison: the answer cells yellow, with the list",
               all(ws[f"B{r}"].fill.start_color.rgb.endswith("FFFF00") for r in ANSWER_ROWS) and
               not ws["B3"].fill.start_color.rgb.endswith("FFFF00") and
               any("B6" in str(dv.sqref) and "account 1" in dv.formula1 for dv in ws.data_validations.dataValidation))
        path2, _ = quiet(comparison_cmd, 7, key_path, (p77, p78), os.path.join(tmp, "co2"), src)
        expect("comparison: the same bytes twice", isinstance(path2, str) and sha(path) == sha(path2))
        r, _ = quiet(comparison_cmd, 7, key_path, (p77, p78), co, src)
        expect("comparison: refuses to write over", refused(r, "refusing"))
        c3 = os.path.join(tmp, "c3")
        quiet(comparison_cmd, 7, key_path, (p77, p78), c3, src, 1)
        k3 = json.load(open(os.path.join(c3, "map_comparison_key.json"), encoding="utf-8"))
        expect("comparison: --take keeps the first of each paper's sample, drawn afresh",
               [t["tab"] for t in k3["tabs"]] == ["Bechky R01", "Pratt R02"] and k3["take"] == 1 and
               [t["account_1"] for t in k3["tabs"]] == want[:2])
        bad = os.path.join(tmp, "bad")
        shutil.copytree(src, bad)
        k0 = json.load(open(os.path.join(bad, "relation_results.json"), encoding="utf-8"))
        k0["codes"]["gemini"]["bechky|R01"]["reason"] = "changed"
        open(os.path.join(bad, "relation_results.json"), "w", encoding="utf-8").write(json.dumps(k0))
        r, _ = quiet(comparison_cmd, 7, key_path, (p77, p78), os.path.join(tmp, "co3"), bad)
        expect("comparison: refuses sources not the key's", refused(r, "not the sources"))
        shutil.rmtree(bad)

        # the PI's filled sheet
        ckey_path = os.path.join(co, "map_comparison_key.json")

        def fill(answers, tamper=None, start="20:00"):
            wbf = load_workbook(os.path.join(co, "MAP_COMPARISON.xlsx"))
            wbf["Instructions"].cell(TIME_ROWS[0], 2).value = start
            for t, (a1, a2, note) in zip(ck["tabs"], answers):
                wsf = wbf[t["tab"]]
                wsf.cell(ANSWER_ROWS[0], 2).value = a1
                wsf.cell(ANSWER_ROWS[1], 2).value = a2
                wsf.cell(ANSWER_ROWS[2], 2).value = note
            if tamper:
                tamper(wbf)
            p = os.path.join(tmp, f"filled{n_ok + n_bad}.xlsx")
            wbf.save(p)
            return p
        ans = [("account 1", "the same", "a note"), ("Account 2 ", "neither", None), ("account 2", None, None)]
        out, problems = read_comparison(ck, fill(ans), os.path.join(co, "MAP_COMPARISON.xlsx"))
        expect("the PI's sheet: read", out is not None and not problems and out["start"] == "20:00" and out["end"] is None)
        if out:
            t0, t1, t2 = out["tabs"]
            expect("the PI's sheet: each answer mapped by the key",
                   t0["q1_pick"] == ck["tabs"][0]["account_1"] and t0["q2_pick"] == "the same" and
                   t1["q1"] == "account 2" and t1["q1_pick"] == ck["tabs"][1]["account_2"] and
                   t1["q2_pick"] == "neither" and t0["note"] == "a note")
            expect("the PI's sheet: an empty answer is unanswered", t2["q2"] is None and t2["q2_pick"] is None)
        _, problems = read_comparison(ck, fill(ans, tamper=lambda w: setattr(w["Bechky R01"]["B3"], "value", "x")),
                                      os.path.join(co, "MAP_COMPARISON.xlsx"))
        expect("the PI's sheet: a change outside the yellow cells is refused", any("B3 differs" in p for p in problems))
        _, problems = read_comparison(ck, fill([("maybe", "neither", None)] + ans[1:]),
                                      os.path.join(co, "MAP_COMPARISON.xlsx"))
        expect("the PI's sheet: an answer not in the list is refused", problems == ["Bechky R01: the answer to q1 does not "
                                                                               "read"])
        oa = os.path.join(tmp, "oa")
        r, _ = quiet(read_comparison_cmd, ckey_path, fill(ans), oa)
        expect("the PI's sheet: the answers written", isinstance(r, dict) and
               os.path.exists(os.path.join(oa, "map_comparison_answers.json")))
    finally:
        shutil.rmtree(tmp)
    print(f"self-test: {n_ok} checks pass, {n_bad} fail")
    return 1 if n_bad else 0


def main(argv):
    if argv[:1] == ["--selftest"]:
        return selftest()
    if argv[:1] == ["files"] and len(argv) == 2:
        files_cmd(argv[1])
        return 0
    if argv[:1] == ["check-placement"] and len(argv) == 4 and argv[1] in PAPERS:
        return check_placement(argv[1], argv[2], argv[3])[0]
    if argv[:1] == ["openings"] and len(argv) == 5:
        openings_cmd(argv[1], (argv[2], argv[3]), argv[4])
        return 0
    if argv[:1] == ["check-openings"] and len(argv) == 4 and argv[1] in PAPERS:
        return check_openings(argv[1], argv[2], argv[3])[0]
    if argv[:1] == ["comparison"] and argv[1:2] == [f"--seed={SEED}"]:
        take = 3 if argv[2:3] == ["--take=3"] else None
        rest = argv[3:] if take else argv[2:]
        if len(rest) == 4:
            comparison_cmd(SEED, rest[0], (rest[1], rest[2]), rest[3], take=take)
            return 0
    if argv[:1] == ["read-comparison"] and len(argv) == 4:
        read_comparison_cmd(argv[1], argv[2], argv[3])
        return 0
    raise SystemExit(__doc__)


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
