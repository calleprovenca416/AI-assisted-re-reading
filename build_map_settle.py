#!/usr/bin/env python3
"""
The map test (T39): the PI's settling sheets, the settled maps, the closure, the modules and the network.

WHY
---
`MAP_TEST_PLAN.md`, parts C and D: where the two maps of a paper differ, the PI chooses; what both maps agree on, with
the PI's choices, is the settled map; from it, by script, the closure of each reading's T37 codes through the settled
edges (the typed effects), the mechanical modules, and the network of the claims' constructs. The plan's order asks
that this script be fixed by its sha256 before any reply of round 1 is opened. It judges nothing.

USAGE
-----
    python3 -B build_map_settle.py sheet <paper> <map_key.json> <reply A> <reply B> <out dir>
        MAP_SETTLE_<paper>.xlsx (the PI's sheet) and map_settle_key_<paper>.json
    python3 -B build_map_settle.py settle <paper> <map_settle_key_<paper>.json> <filled sheet> <out dir>
        map_settled_<paper>.json (the blank sheet is read from beside the settle key)
    python3 -B build_map_settle.py closure <paper> <map_settled_<paper>.json> <relation_results.json> <out dir>
        map_closure_<paper>.json
    python3 -B build_map_settle.py modules <paper> <map_settled_<paper>.json> <out dir>
        map_modules_<paper>.json and MAP_MODULES_<paper>.md
    python3 -B build_map_settle.py network <paper> <map_settled_<paper>.json> <out dir>
        MAP_NETWORK_<paper>.html
    python3 -B build_map_settle.py --selftest

Map A is the Opus map (chats 73 and 75), map B the GPT 6 map (chats 74 and 76); the sheet names no model. Each reply is
read with `build_map_test.py`'s check and refused unless it counts with no quoted span of more than eight words (a
reply with such a span is filed without it, and the filed copy is the one given here). The index and the voice file
are checked against their hashes. It refuses to write over a file, and the same inputs always give the same bytes.

THE SETTLING SHEET (the plan's part C)
--------------------------------------
Every item on which the two maps differ, numbered S001 onward in this order, each with both maps' answers and reasons,
a yellow cell for the PI's choice (A, B, or the PI's own where an own answer can differ from both) and one for a note:
1. a listed pair that one map calls an edge and the other not; a listed pair both call an edge, with another type or
   strength (tab Pairs, in the order of the pairs);
2. a further edge (a pair not listed) that one map gives and the other not, or both give with another type or
   strength (tab Pairs, by the number in the base's id, then the claim's); a further edge on a listed pair is left
   aside, the listed pair's answers standing, and a map's second further edge on one pair is left aside too;
3. two construct names that one map's registry puts in one group and the other's does not (tab Constructs), each
   map's groups taken as the transitive closure of its registry over the sheet's construct names;
4. a factor one map gives and the other does not (tab Factors), the two maps' factors paired in order by the claim
   each cites; then, where part 5 applies, a unit's outcome, and a paired factor's value in a unit, on which the maps
   differ; and, where part 6 applies, a paired factor's control in a contrast pair on which they differ.
Choosing A or B takes that map's answer, which may be no edge, apart, or no such factor. Two yellow cells take the PI's
start and end times. Values are shown as words: present, absent and not said for a factor's value; shows it, does
not and not said for a unit's outcome; held, varies and unclear for a control.

THE SETTLED MAP
---------------
Its edges: the listed pairs both maps call edges of one type and strength, the further edges both give alike, and
every item of kind 1 or 2 as the PI chose it. Its registry: the name pairs both maps group, and those the PI chose to group,
closed transitively; a group's canonical name is map A's, then map B's, where the group is that map's group exactly,
and otherwise the name the most claims use (ties by alphabet). Its factors: the paired factors (named as in map A) and
those of one map the PI kept, numbered X1 onward, map A's order first; each value, outcome and control as both maps give
it, or as the PI chose, or as the one map that has the factor gives it. The settle step compares the filled sheet with
the blank, cell by cell outside the yellow cells, and refuses any other change, an item without a choice, and an own
answer that does not read.

THE CLOSURE (the plan's part D, as its foot reads it from 18 Oct)
------------------------------------------------------------------
A reading hits a claim when a T37 code gives it inverts, other_mechanism, rereads_evidence or challenges_premise, as
its main claim or another; each coder's codes (gemini, astra) are taken on their own. A hit claim, core or not, is
contested. Through each settled edge whose base is contested or unsupported: evidential, the claim is unsupported if
the edge is necessary or every evidential edge into the claim comes from a contested or unsupported claim, otherwise
weakened; mechanistic, unsupported if necessary, weakened if contributing; comparative, restated; significance, idle.
Only contested and unsupported claims carry an effect onward; a claim keeps the strongest of its labels (contested,
unsupported, weakened, restated, idle); the rules are applied until nothing changes. Effects are reported on core
claims, each with the coder or coders whose codes give it; every claim's label is kept in the file.

THE MODULES AND THE NETWORK
---------------------------
For each core claim, from the index, its voice file and the settled map: evidence (its passages by direction, voice
code, data type and unit); timing (the paper's periods and resolution, and its passages' periods); scope (its
passages' units, and the declared scope limits bearing on it); constructs (its canonical constructs, each with the
other claims that name it); openings (whether it is ventured, the ventured claims with a settled edge into it, the
declared gaps and absences bearing on it, its qualifying and against links); and, where part 5 applies, the factors
whose value equals the unit's outcome in every unit, none unknown. The network: a page that draws each claim beside
the canonical constructs that two or more claims name, a line for each claim naming one, and lists every canonical
construct below; labelled as shared constructs, not dependence.
"""
import collections
import datetime
import hashlib
import html
import itertools
import json
import os
import re
import sys
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import build_map_test as bmt  # noqa: E402
import build_worth_sheet as bws  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402
from openpyxl.styles import Alignment  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

bcs = bmt.bcs
PAPERS = bmt.PAPERS
CHATS = {"bechky": {"A": 73, "B": 74}, "pratt": {"A": 75, "B": 76}}
VOICE = {"bechky": "bechky_2020_v1_voice_gloss.json", "pratt": "pratt_2000_v1_voice_gloss.json"}
VOICE_SHA = {"bechky_2020_v1_voice_gloss.json": "e4b6d8c974a6661440e9366f24f85e8ebfc51d0defdceea99281f2eb567bea8a",
             "pratt_2000_v1_voice_gloss.json": "7e26bffb2e5cc8cdb9f359bae8472dbaa5a9a667e6a4dd6d9911bd74e9d89956"}
RESULTS_SHA = "5dff925eb5d23484ae44642628f29fd12e28366ef4db1e5fc907c4ad97a1fc47"
CODERS = ("gemini", "astra")
HIT = ("inverts", "other_mechanism", "rereads_evidence", "challenges_premise")
LABELS = ("contested", "unsupported", "weakened", "restated", "idle")
VALUE_WORD = {"+": "present", "-": "absent", "?": "not said"}
OUTCOME_WORD = {"+": "shows it", "-": "does not", "?": "not said"}
CONTROLS = bmt.CONTROLS
OWN_TYPES = ("none",) + bmt.TYPES
OWN_WORDS = ("present", "absent", "not said", "shows it", "does not", "held", "varies", "unclear")
STAMP = (2026, 10, 18, 0, 0, 0)
EDGE_KINDS = ("pair-edge", "pair-type", "further-one", "further-type")
FACTOR_KINDS = ("factor-one", "unit-outcome", "unit-value", "contrast")
KINDS = {
    "pair-edge": "listed pair: an edge in one map only",
    "pair-type": "listed pair: an edge in both, another type or strength",
    "further-one": "further edge: in one map only",
    "further-type": "further edge: in both, another type or strength",
    "constructs": "two names: one group in one map only",
    "factor-one": "factor: in one map only",
    "unit-outcome": "unit: its outcome",
    "unit-value": "unit: a factor's value",
    "contrast": "contrast pair: a factor's control",
}


def sha_bytes(data):
    return hashlib.sha256(data).hexdigest()


def sha(path):
    return sha_bytes(open(path, "rb").read())


num = bmt.num


def refuse_existing(outdir, names):
    os.makedirs(outdir, exist_ok=True)
    for n in names:
        if os.path.exists(os.path.join(outdir, n)):
            raise SystemExit(f"refusing to write over {os.path.join(outdir, n)}")


def dumps(obj):
    return (json.dumps(obj, ensure_ascii=False, indent=1, sort_keys=True) + "\n").encode("utf-8")


def name_key(s):
    return (s.casefold(), s)


# ---------------------------------------------------------------- the replies and their differences

def read_reply(paper, key, path, index):
    """The reply's object as build_map_test's check reads it (its lenient readings applied), refused unless it
    counts, with no long quoted span."""
    rep, faults = bcs.load(path)
    if not faults:
        dup = bmt.twice(path)
        faults = [f"a key given twice: {', '.join(sorted(set(dup)))}"] if dup else []
        rep, _ = bmt.normalized(rep, paper, key, index)
        faults += bmt.faults_of(rep, paper, key, index)[0]
    if faults:
        raise SystemExit(f"{path}: does not count: {'; '.join(faults[:5])}")
    if bcs.long_spans(rep, index["meta"]["citation"]):
        raise SystemExit(f"{path}: a quoted span of more than eight words; give the filed copy, without it")
    return rep


def edge_text(e):
    if not e:
        return "not given"
    if e.get("rests_on") is False:
        return "no"
    return f"yes: {e['type']}, {e['strength']}; basis: {e['basis']}; {e['page']}"


def further_of(rep, listed):
    out = {}
    for e in rep["further_edges"]:
        k = (e["base"], e["claim"])
        if k not in listed and k not in out:
            out[k] = dict(e, rests_on=True)
    return out


def closure_groups(pairs, names):
    """The groups (two names or more) of the transitive closure of name pairs."""
    parent = {n: n for n in names}

    def find(x):
        while parent[x] != x:
            parent[x] = parent[parent[x]]
            x = parent[x]
        return x
    for x, y in pairs:
        rx, ry = find(x), find(y)
        if rx != ry:
            parent[max(rx, ry)] = min(rx, ry)
    members = collections.defaultdict(set)
    for n in names:
        members[find(n)].add(n)
    return [frozenset(m) for m in members.values() if len(m) > 1]


def groups_of(reg, names):
    """A map's groups over the sheet's names: ({name: its group}, {group: canonical name})."""
    pairs = []
    canon_of = {}
    for g in reg:
        ns = [n for n in g["names"] if n in names]
        pairs += list(zip(ns, ns[1:]))
        for n in ns:
            canon_of.setdefault(n, g["canonical"])
    groups, canon = {}, {}
    for g in closure_groups(pairs, names):
        for n in g:
            groups[n] = g
        canon[g] = next((canon_of[n] for n in sorted(g, key=name_key) if n in canon_of), sorted(g, key=name_key)[0])
    return groups, canon


def together(groups):
    return {frozenset(p) for g in set(groups.values()) for p in itertools.combinations(sorted(g), 2)}


def align(fa, fb):
    """Pairs (i, j) of factors citing one claim, in order; then A's unpaired, then B's."""
    used, pairs = set(), []
    for i, x in enumerate(fa):
        j = next((j for j, y in enumerate(fb) if j not in used and y["claim"] == x["claim"]), None)
        if j is not None:
            used.add(j)
            pairs.append((i, j))
    paired = {i for i, _ in pairs}
    return pairs, [i for i in range(len(fa)) if i not in paired], [j for j in range(len(fb)) if j not in used]


def compare(paper, key, index, ra, rb):
    """(items, agreed): every difference between the two maps, and each map's answers where they agree or one alone
    answers."""
    kp = key["papers"][paper]
    ap = kp["applies"]
    items = []

    def item(kind, **kw):
        items.append(dict(kw, kind=kind, id=None))
    agreed = {"edges": []}
    A = {p["pair"]: p for p in ra["pairs"]}
    B = {p["pair"]: p for p in rb["pairs"]}
    for o in kp["order"]:
        a, b = A[o["id"]], B[o["id"]]
        base = {"pair": o["id"], "base": o["base"], "claim": o["claim"], "a": a, "b": b}
        if a["rests_on"] != b["rests_on"]:
            item("pair-edge", **base)
        elif a["rests_on"] and (a["type"], a["strength"]) != (b["type"], b["strength"]):
            item("pair-type", **base)
        elif a["rests_on"]:
            agreed["edges"].append(dict(base, type=a["type"], strength=a["strength"], source="listed", by="both"))
    listed = {(o["base"], o["claim"]) for o in kp["order"]}
    fa, fb = further_of(ra, listed), further_of(rb, listed)
    for k in sorted(set(fa) | set(fb), key=lambda k: (num(k[0]), num(k[1]))):
        a, b = fa.get(k), fb.get(k)
        base = {"pair": None, "base": k[0], "claim": k[1], "a": a, "b": b}
        if a and b and (a["type"], a["strength"]) == (b["type"], b["strength"]):
            agreed["edges"].append(dict(base, type=a["type"], strength=a["strength"], source="further", by="both"))
        elif a and b:
            item("further-type", **base)
        else:
            item("further-one", **base)
    names = sorted(bmt.construct_names(index), key=name_key)
    ga, ca = groups_of(ra["registry"], names)
    gb, cb = groups_of(rb["registry"], names)
    ta, tb = together(ga), together(gb)
    for p in sorted(ta ^ tb, key=lambda p: sorted(name_key(s) for s in p)):
        x, y = sorted(p, key=name_key)
        item("constructs", names=[x, y], a=f"one group: {ca[ga[x]]}" if p in ta else "apart",
             b=f"one group: {cb[gb[x]]}" if p in tb else "apart")
    agreed["registry"] = {
        "names": names, "both": sorted(sorted(p, key=name_key) for p in ta & tb),
        "a": [{"canonical": c, "names": sorted(g, key=name_key)} for g, c in sorted(ca.items(), key=lambda gc: sorted(gc[0]))],
        "b": [{"canonical": c, "names": sorted(g, key=name_key)} for g, c in sorted(cb.items(), key=lambda gc: sorted(gc[0]))]}
    pairs, a_only, b_only = align(ra["factors"], rb["factors"])
    pair_of = dict(pairs)
    order = [("A", i) for i in sorted([i for i, _ in pairs] + a_only)] + [("B", j) for j in b_only]
    factors = []
    for n, (m, i) in enumerate(order, 1):
        fx = {"id": "F%d" % n, "a": ra["factors"][i] if m == "A" else None, "b": None}
        if m == "A" and i in pair_of:
            fx["b"] = rb["factors"][pair_of[i]]
        if m == "B":
            fx["b"] = rb["factors"][i]
        one = fx["a"] or fx["b"]
        fx["name"], fx["claim"] = one["name"], one["claim"]
        factors.append(fx)
    agreed["factors"] = factors
    for fx in factors:
        if not (fx["a"] and fx["b"]):
            item("factor-one", factor=fx["id"], a=fx["a"], b=fx["b"])
    agreed["units"], agreed["contrasts"] = [], []
    if ap["table"]:
        UA = {u["unit"]: u for u in ra["units"]}
        UB = {u["unit"]: u for u in rb["units"]}
        for x in ap["units"]:
            u = x["unit"]
            oa, ob = bmt.SIGNS[UA[u]["outcome"]], bmt.SIGNS[UB[u]["outcome"]]
            row = {"unit": u, "claim": x["claim"], "outcome": {"a": oa, "b": ob}, "values": {},
                   "ids": {"a": UA[u]["ids"], "b": UB[u]["ids"]}}
            if oa != ob:
                item("unit-outcome", unit=u, a=oa, b=ob)
            for fx in factors:
                va = bmt.SIGNS[UA[u]["values"][fx["a"]["id"]]] if fx["a"] else None
                vb = bmt.SIGNS[UB[u]["values"][fx["b"]["id"]]] if fx["b"] else None
                row["values"][fx["id"]] = {"a": va, "b": vb}
                if va and vb and va != vb:
                    item("unit-value", unit=u, factor=fx["id"], a=va, b=vb)
            agreed["units"].append(row)
    if ap["contrasts"]:
        CA = {c["contrast"]: c for c in ra["contrasts"]}
        CB = {c["contrast"]: c for c in rb["contrasts"]}
        for cp in ap["contrasts"]:
            row = {"contrast": cp, "controls": {}, "ids": {"a": CA[cp]["ids"], "b": CB[cp]["ids"]}}
            for fx in factors:
                va = CA[cp]["controls"][fx["a"]["id"]] if fx["a"] else None
                vb = CB[cp]["controls"][fx["b"]["id"]] if fx["b"] else None
                row["controls"][fx["id"]] = {"a": va, "b": vb}
                if va and vb and va != vb:
                    item("contrast", contrast=cp, factor=fx["id"], a=va, b=vb)
            agreed["contrasts"].append(row)
    rank = {k: i for i, k in enumerate(EDGE_KINDS + ("constructs",) + FACTOR_KINDS)}
    items.sort(key=lambda it: (0 if it["kind"] in EDGE_KINDS[:2] else 1 if it["kind"] in EDGE_KINDS[2:] else
                               rank[it["kind"]]))
    for n, it in enumerate(items, 1):
        it["id"] = "S%03d" % n
    return items, agreed


# ---------------------------------------------------------------- the sheet

def steady(path):
    stamp = b"2026-10-18T00:00:00Z"
    with zipfile.ZipFile(path) as z:
        entries = [(i, z.read(i.filename)) for i in z.infolist()]
    with zipfile.ZipFile(path, "w") as z:
        for i, data in entries:
            if i.filename == "docProps/core.xml":
                data, n = re.subn(rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)",
                                  rb"\g<1>" + stamp + rb"\g<2>", data)
                if n != 1:
                    raise SystemExit("docProps/core.xml: no single modified date")
            j = zipfile.ZipInfo(i.filename, date_time=STAMP)
            j.compress_type = i.compress_type
            j.external_attr = i.external_attr
            z.writestr(j, data)


def save_xlsx(wb, path):
    wb.properties.created = wb.properties.modified = datetime.datetime(*STAMP)
    wb.properties.creator = "build_map_settle.py"
    wb.save(path)
    steady(path)


def validation(options, title):
    formula = '"' + ",".join(options) + '"'
    if len(formula) > 255:
        raise SystemExit(f"{title}: the list of choices is too long for a cell")
    return DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=True,
                          errorTitle=title, error="Choose one of the list: " + ", ".join(options),
                          showInputMessage=True, promptTitle=title, prompt=", ".join(options))


def table(wb, title, heads, widths, rows, yellow=(), lists=()):
    """A tab: a grey head row, then the rows; the columns in `yellow` filled yellow; each list (options, column)."""
    ws = wb.create_sheet(title)
    cols = [chr(ord("A") + i) for i in range(len(heads))]
    for col, w in zip(cols, widths):
        ws.column_dimensions[col].width = w
    for i, h in enumerate(heads):
        bws.cell(ws, f"{cols[i]}1", h, bold=True, fill=bws.GREY).border = bws.BOX
    dvs = []
    for options, col in lists:
        dv = validation(options, heads[cols.index(col)])
        ws.add_data_validation(dv)
        dvs.append((dv, col))
    for r, row in enumerate(rows, start=2):
        live = row[0] != "(none)"
        for i, v in enumerate(row):
            c = bws.cell(ws, f"{cols[i]}{r}", v, bold=(i == 0), fill=(bws.YELLOW if i in yellow and live else None))
            c.border = bws.BOX
            if i in yellow:
                c.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
        if live:
            for dv, col in dvs:
                dv.add(f"{col}{r}")
        ws.row_dimensions[r].height = max(bws.height(v or "", max(8, w - 2)) for v, w in zip(row, widths))
    ws.freeze_panes = "B2"
    bws.printable(ws)
    return ws


def factor_text(f):
    return "no such factor" if not f else f"{f['name']} (stated in {f['claim']}; the map's {f['id']})"


INSTRUCTIONS = [
    (None, None),
    ("What it is", "The two maps of this paper, built blind to the readings by two coders, differ on the items below. "
                   "For each, choose which map's answer stands, or give your own. What both maps agree on stands "
                   "without you. Map A and map B are the two coders' maps; the sheet does not say which is which."),
    ("Fill in", "Only the yellow cells. Your choice: A or B takes that map's answer as it is, which may be no edge, "
                "apart, or no such factor; own takes the answer you give in the cells after it. A note is welcome and "
                "never needed."),
    ("Pairs", "Each asks whether the claim rests on the base, as the map task defines it. If own, give a type (none "
              "for no edge) and, for an edge, a strength. The claims are on the tab Claims."),
    ("Constructs", "Each asks whether two construct names name one construct in this paper: A or B. Each map's groups "
                   "are on the tab Groups."),
    ("Factors", "Each asks whether a factor stands (A or B), or for a unit's outcome, a factor's value in a unit, or a "
                "factor's control in a contrast pair (A, B, or own with a word from the list: present, absent or not "
                "said for a value; shows it, does not or not said for an outcome; held, varies or unclear for a "
                "control). Each map's factors are on the tab Map factors."),
    ("Time", "Type the time you start and the time you finish, as hh:mm, in the two yellow cells below."),
    ("Started at", ""),
    ("Finished at", ""),
    ("Save", "Save it as {filled}, with no other change, and send it in our conversation or put it in replies/."),
]
TIME_ROWS = (8, 9)


def build_sheet(paper, index, ra, rb, items, agreed):
    claims = bmt.structure(index)[0]
    wb = Workbook()
    ins = wb.active
    ins.title = "Instructions"
    ins.column_dimensions["A"].width = 18
    ins.column_dimensions["B"].width = 110
    filled = f"MAP_SETTLE_{paper}_filled.xlsx"
    for r, (label, t) in enumerate(INSTRUCTIONS, 1):
        if r == 1:
            bws.cell(ins, "A1", f"{index['meta']['citation']}: the settling sheet of the map test", bold=True,
                     size=12, wrap=False)
            continue
        bws.cell(ins, f"A{r}", label, bold=True)
        c = bws.cell(ins, f"B{r}", t.format(filled=filled), fill=bws.YELLOW if r in TIME_ROWS else None)
        ins.row_dimensions[r].height = bws.height(t, 120)
        if r in TIME_ROWS:
            c.border = bws.BOX
    bws.printable(ins)
    rows = [[it["id"], KINDS[it["kind"]], it["pair"] or "not listed", f"does {it['claim']} rest on {it['base']}?",
             edge_text(it["a"]), (it["a"] or {}).get("reason", ""), edge_text(it["b"]), (it["b"] or {}).get("reason", ""),
             None, None, None, None] for it in items if it["kind"] in EDGE_KINDS]
    table(wb, "Pairs", ["item", "kind", "pair", "question", "map A", "map A's reason", "map B", "map B's reason",
                        "your choice", "own type", "own strength", "note"],
          [7, 18, 9, 16, 34, 40, 34, 40, 10, 14, 13, 24], rows or [["(none)"] + [""] * 11],
          yellow=(8, 9, 10, 11), lists=((("A", "B", "own"), "I"), (OWN_TYPES, "J"), (bmt.STRENGTHS, "K")))
    rows = [[it["id"], it["names"][0], it["names"][1], it["a"], it["b"], None, None]
            for it in items if it["kind"] == "constructs"]
    table(wb, "Constructs", ["item", "name 1", "name 2", "map A", "map B", "your choice", "note"],
          [7, 34, 34, 34, 34, 11, 30], rows or [["(none)"] + [""] * 6], yellow=(5, 6), lists=((("A", "B"), "F"),))
    fxs = {fx["id"]: fx for fx in agreed["factors"]}
    rows = []
    for it in items:
        if it["kind"] == "factor-one":
            what, a, b = f"factor {it['factor']}", factor_text(it["a"]), factor_text(it["b"])
        elif it["kind"] == "unit-outcome":
            what, a, b = f"{it['unit']}: its outcome", OUTCOME_WORD[it["a"]], OUTCOME_WORD[it["b"]]
        elif it["kind"] == "unit-value":
            what = f"{it['unit']}: factor {it['factor']}, {fxs[it['factor']]['name']}"
            a, b = VALUE_WORD[it["a"]], VALUE_WORD[it["b"]]
        elif it["kind"] == "contrast":
            what, a, b = f"{it['contrast']}: factor {it['factor']}, {fxs[it['factor']]['name']}", it["a"], it["b"]
        else:
            continue
        rows.append([it["id"], KINDS[it["kind"]], what, a, b, None, None, None])
    table(wb, "Factors", ["item", "kind", "what", "map A", "map B", "your choice", "own value", "note"],
          [7, 22, 40, 34, 34, 11, 13, 30], rows or [["(none)"] + [""] * 7], yellow=(5, 6, 7),
          lists=((("A", "B", "own"), "F"), (OWN_WORDS, "G")))
    rows = [[c, "core" if claims[c].get("core") is True else "", claims[c].get("location") or "", claims[c]["statement"]]
            for c in sorted(claims, key=num)]
    table(wb, "Claims", ["claim", "core", "pages", "statement"], [8, 7, 30, 110], rows)
    rows = [[m, g["canonical"], "; ".join(g["names"])] for m, rep in (("A", ra), ("B", rb)) for g in rep["registry"]]
    table(wb, "Groups", ["map", "canonical name", "names"], [7, 34, 110], rows or [["(none)", "", ""]])
    rows = []
    for fx in agreed["factors"]:
        for m in ("a", "b"):
            if fx[m]:
                rows.append([m.upper(), fx[m]["id"], fx[m]["name"], fx[m]["claim"], fx["id"]])
    rows.sort(key=lambda r: (r[0], [f[m]["id"] for f in agreed["factors"] for m in ("a", "b")
                                    if f[m] and m.upper() == r[0]].index(r[1])))
    table(wb, "Map factors", ["map", "the map's id", "name", "claim", "id on this sheet"], [7, 12, 60, 8, 16],
          rows or [["(none)", "", "", "", ""]])
    return wb


def sheet_cmd(paper, key_path, path_a, path_b, outdir, folder=None):
    key = json.load(open(key_path, encoding="utf-8"))
    index, index_sha = bmt.load_index(paper, folder)
    ra, rb = read_reply(paper, key, path_a, index), read_reply(paper, key, path_b, index)
    items, agreed = compare(paper, key, index, ra, rb)
    names = [f"MAP_SETTLE_{paper}.xlsx", f"map_settle_key_{paper}.json"]
    refuse_existing(outdir, names)
    xpath = os.path.join(outdir, names[0])
    save_xlsx(build_sheet(paper, index, ra, rb, items, agreed), xpath)
    skey = {"what": f"the settle key of T39 for {paper}: the two maps' differences as items, and each map's answers "
                    "where they agree or one alone answers; with the hashes of the inputs and of the blank sheet",
            "paper": paper, "chats": CHATS[paper], "builder_sha256": sha(os.path.abspath(__file__)),
            "inputs": {"map_key": sha(key_path), "index": index_sha, "reply_a": sha(path_a), "reply_b": sha(path_b)},
            "reply_files": {"a": os.path.basename(path_a), "b": os.path.basename(path_b)},
            "sheet": names[0], "sheet_sha256": sha(xpath), "items": items, "agreed": agreed,
            "counts": dict(sorted(collections.Counter(it["kind"] for it in items).items()))}
    bmt.write_new(os.path.join(outdir, names[1]), dumps(skey))
    print(f"written: {xpath} ({os.path.getsize(xpath)} bytes, sha256 {sha(xpath)}); {len(items)} items "
          f"{skey['counts']}")
    return skey


# ---------------------------------------------------------------- the settled map

def tab_of(kind):
    return "Pairs" if kind in EDGE_KINDS else "Constructs" if kind == "constructs" else "Factors"


def yellow_cells(skey):
    """{tab: (rows, cols)} of the cells the PI may change, as build_sheet lays them out."""
    n = collections.Counter(tab_of(it["kind"]) for it in skey["items"])
    return {"Instructions": (set(TIME_ROWS), {2}),
            "Pairs": (set(range(2, 2 + n["Pairs"])), {9, 10, 11, 12}),
            "Constructs": (set(range(2, 2 + n["Constructs"])), {6, 7}),
            "Factors": (set(range(2, 2 + n["Factors"])), {6, 7, 8})}


def compare_cells(filled, blank, free):
    wf, wb = load_workbook(filled), load_workbook(blank)
    if wf.sheetnames != wb.sheetnames:
        return [f"tabs {wf.sheetnames} are not {wb.sheetnames}"], None
    faults = []
    for name in wb.sheetnames:
        a, b = wf[name], wb[name]
        rows, cols = free.get(name, ((), ()))
        for r in range(1, max(a.max_row, b.max_row) + 1):
            for c in range(1, max(a.max_column, b.max_column) + 1):
                if r in rows and c in cols:
                    continue
                if a.cell(r, c).value != b.cell(r, c).value:
                    faults.append(f"{name}!{a.cell(r, c).coordinate} differs from the blank sheet")
    return faults, wf


def text(v):
    return None if v is None or not str(v).strip() else str(v).strip()


def read_choices(skey, wf):
    """({item: choice}, problems) from the filled workbook."""
    choices, problems = {}, []
    rows = collections.Counter()
    for it in skey["items"]:
        tab = tab_of(it["kind"])
        rows[tab] += 1
        r = rows[tab] + 1
        ws = wf[tab]
        if ws.cell(r, 1).value != it["id"]:
            problems.append(f"{tab} row {r}: not item {it['id']}")
            continue
        if tab == "Pairs":
            ch, ot, st, note = (text(ws.cell(r, c).value) for c in (9, 10, 11, 12))
            ok = ch in ("A", "B") or (ch == "own" and ot in OWN_TYPES and (ot == "none" or st in bmt.STRENGTHS))
            c = {"choice": ch, "type": ot if ch == "own" else None,
                 "strength": st if ch == "own" and ot != "none" else None, "note": note}
        elif tab == "Constructs":
            ch, note = text(ws.cell(r, 6).value), text(ws.cell(r, 7).value)
            ok, c = ch in ("A", "B"), {"choice": ch, "note": note}
        else:
            ch, own, note = (text(ws.cell(r, c).value) for c in (6, 7, 8))
            allowed = {"unit-outcome": {w: s for s, w in OUTCOME_WORD.items()},
                       "unit-value": {w: s for s, w in VALUE_WORD.items()},
                       "contrast": {w: w for w in CONTROLS}}.get(it["kind"], {})
            ok = ch in ("A", "B") or (ch == "own" and it["kind"] != "factor-one" and own in allowed)
            c = {"choice": ch, "own": allowed.get(own) if ch == "own" else None, "note": note}
        if not ok:
            problems.append(f"{it['id']}: the choice does not read")
        else:
            choices[it["id"]] = c
    return choices, problems


def settle(skey, index, choices):
    """The settled map from the settle key and the PI's choices."""
    agreed = skey["agreed"]
    edges = [dict(e) for e in agreed["edges"]]
    for it in skey["items"]:
        if it["kind"] not in EDGE_KINDS:
            continue
        c = choices[it["id"]]
        base = {"pair": it["pair"], "base": it["base"], "claim": it["claim"], "a": it["a"], "b": it["b"],
                "source": "listed" if it["pair"] else "further", "item": it["id"]}
        if c["choice"] in ("A", "B"):
            ans = it[c["choice"].lower()]
            if ans and ans.get("rests_on"):
                edges.append(dict(base, type=ans["type"], strength=ans["strength"], by=c["choice"]))
        elif c["type"] != "none":
            edges.append(dict(base, type=c["type"], strength=c["strength"], by="own"))
    edges.sort(key=lambda e: (num(e["claim"]), num(e["base"])))
    reg = agreed["registry"]
    pairs = [tuple(p) for p in reg["both"]]
    for it in skey["items"]:
        if it["kind"] == "constructs" and it[choices[it["id"]]["choice"].lower()] != "apart":
            pairs.append(tuple(it["names"]))
    use = collections.Counter(n for c in index["claims"] for n in set(c.get("constructs") or []))
    exact = {}
    for m in ("a", "b"):
        for g in reg[m]:
            exact.setdefault(frozenset(g["names"]), g["canonical"])
    registry = []
    for g in closure_groups(sorted(pairs), reg["names"]):
        canon = exact.get(g) or sorted(g, key=lambda n: (-use[n], n.casefold(), n))[0]
        registry.append({"canonical": canon, "names": sorted(g, key=name_key)})
    registry.sort(key=lambda g: name_key(g["canonical"]))
    dropped = set()
    for it in skey["items"]:
        if it["kind"] == "factor-one" and not it[choices[it["id"]]["choice"].lower()]:
            dropped.add(it["factor"])
    kept = [fx for fx in agreed["factors"] if fx["id"] not in dropped]
    renum = {fx["id"]: "X%d" % n for n, fx in enumerate(kept, 1)}
    factors = [{"id": renum[fx["id"]], "name": fx["name"], "claim": fx["claim"], "a": fx["a"], "b": fx["b"],
                "by": "both" if fx["a"] and fx["b"] else "A" if fx["a"] else "B"} for fx in kept]
    chosen = {}
    for it in skey["items"]:
        if it["kind"] in ("unit-outcome", "unit-value", "contrast"):
            c = choices[it["id"]]
            v = c["own"] if c["choice"] == "own" else it[c["choice"].lower()]
            chosen[(it["kind"], it.get("unit") or it.get("contrast"), it.get("factor"))] = v

    def pick(ab, k):
        a, b = ab["a"], ab["b"]
        if a and b:
            return a if a == b else chosen[k]
        return a or b
    units = []
    for u in agreed["units"]:
        units.append({"unit": u["unit"], "claim": u["claim"],
                      "outcome": pick(u["outcome"], ("unit-outcome", u["unit"], None)),
                      "values": {renum[f]: pick(v, ("unit-value", u["unit"], f)) for f, v in u["values"].items()
                                 if f in renum}, "ids": u["ids"]})
    contrasts = []
    for cp in agreed["contrasts"]:
        contrasts.append({"contrast": cp["contrast"], "ids": cp["ids"],
                          "controls": {renum[f]: pick(v, ("contrast", cp["contrast"], f))
                                       for f, v in cp["controls"].items() if f in renum}})
    return {"what": f"the settled map of T39 for {skey['paper']}: what both maps agree on, with the PI's choices",
            "paper": skey["paper"], "edges": edges, "registry": registry, "factors": factors, "units": units,
            "contrasts": contrasts, "choices": {i: choices[i] for i in sorted(choices)},
            "counts": {"edges": len(edges), "by": dict(sorted(collections.Counter(e["by"] for e in edges).items())),
                       "items": len(skey["items"])}}


def settle_cmd(paper, skey_path, filled, outdir, folder=None):
    skey = json.load(open(skey_path, encoding="utf-8"))
    if skey["paper"] != paper:
        raise SystemExit(f"{skey_path}: a settle key for {skey['paper']}, not {paper}")
    blank = os.path.join(os.path.dirname(os.path.abspath(skey_path)), skey["sheet"])
    if sha(blank) != skey["sheet_sha256"]:
        raise SystemExit(f"{blank}: not the blank sheet the settle key names")
    index, index_sha = bmt.load_index(paper, folder)
    if index_sha != skey["inputs"]["index"]:
        raise SystemExit("the index is not the one the settle key names")
    faults, wf = compare_cells(filled, blank, yellow_cells(skey))
    if faults:
        raise SystemExit(f"{filled}: " + "; ".join(faults[:6]))
    choices, problems = read_choices(skey, wf)
    if problems:
        raise SystemExit(f"{filled}: " + "; ".join(problems[:8]))
    settled = settle(skey, index, choices)
    ins = wf["Instructions"]
    settled["times"] = {"start": text(ins.cell(TIME_ROWS[0], 2).value), "end": text(ins.cell(TIME_ROWS[1], 2).value)}
    settled["inputs"] = {"settle_key": sha(skey_path), "filled_sheet": sha(filled), "blank_sheet": skey["sheet_sha256"],
                         **{k: v for k, v in skey["inputs"].items()}}
    settled["builder_sha256"] = sha(os.path.abspath(__file__))
    name = f"map_settled_{paper}.json"
    refuse_existing(outdir, [name])
    bmt.write_new(os.path.join(outdir, name), dumps(settled))
    print(f"  {settled['counts']}; times {settled['times']}")
    return settled


# ---------------------------------------------------------------- the closure

def labels_of(edges, hit):
    """{claim: label} under the plan's part D, as its foot reads it: every hit claim contested, then the edges."""
    label = {c: "contested" for c in hit}
    into = collections.defaultdict(list)
    for e in edges:
        if e["type"] == "evidential":
            into[e["claim"]].append(e["base"])
    changed = True
    while changed:
        changed = False
        carry = {c for c, l in label.items() if l in ("contested", "unsupported")}
        for e in edges:
            if e["base"] not in carry:
                continue
            t, s, c = e["type"], e["strength"], e["claim"]
            if t == "evidential":
                new = "unsupported" if s == "necessary" or all(b in carry for b in into[c]) else "weakened"
            elif t == "mechanistic":
                new = "unsupported" if s == "necessary" else "weakened"
            elif t == "comparative":
                new = "restated"
            else:
                new = "idle"
            if c not in label or LABELS.index(new) < LABELS.index(label[c]):
                label[c] = new
                changed = True
    return label


def hits_of(code):
    out = []
    for c, t in [(code["main_claim"], code["main_type"])] + [tuple(x) for x in code.get("other_claims") or []]:
        if t in HIT and c not in out:
            out.append(c)
    return out


def closure_cmd(paper, settled_path, results_path, outdir, folder=None, results_sha=RESULTS_SHA):
    settled = json.load(open(settled_path, encoding="utf-8"))
    if settled["paper"] != paper:
        raise SystemExit(f"{settled_path}: a settled map of {settled['paper']}, not {paper}")
    if results_sha and sha(results_path) != results_sha:
        raise SystemExit(f"{results_path}: not the T37 results of 17 Oct")
    index, _ = bmt.load_index(paper, folder)
    core = bmt.structure(index)[4]
    res = json.load(open(results_path, encoding="utf-8"))
    readings = sorted({k for coder in CODERS for k in res["codes"][coder] if k.split("|")[0] == paper},
                      key=lambda k: num(k.split("|")[1]))
    out = []
    for k in readings:
        entry = {"reading": k.split("|")[1], "coders": {}, "core_effects": []}
        for coder in CODERS:
            code = res["codes"][coder].get(k)
            if code is None:
                continue
            hit = hits_of(code)
            lab = labels_of(settled["edges"], hit)
            entry["coders"][coder] = {"main_claim": code["main_claim"], "main_type": code["main_type"], "hits": hit,
                                      "labels": dict(sorted(lab.items(), key=lambda kv: num(kv[0])))}
        effects = collections.defaultdict(list)
        for coder, v in entry["coders"].items():
            for c, l in v["labels"].items():
                if c in core:
                    effects[(c, l)].append(coder)
        entry["core_effects"] = [{"claim": c, "label": l, "coders": cs}
                                 for (c, l), cs in sorted(effects.items(), key=lambda kv: (num(kv[0][0]),
                                                                                          LABELS.index(kv[0][1])))]
        out.append(entry)
    doc = {"what": f"the closure of T37's codes through the settled map of {paper} (T39, part D, as the plan's foot "
                   "reads it from 18 Oct)", "paper": paper, "core": core, "readings": out,
           "inputs": {"settled": sha(settled_path), "results": sha(results_path)},
           "builder_sha256": sha(os.path.abspath(__file__))}
    name = f"map_closure_{paper}.json"
    refuse_existing(outdir, [name])
    bmt.write_new(os.path.join(outdir, name), dumps(doc))
    return doc


# ---------------------------------------------------------------- the modules

def canon_map(settled, index):
    """{construct name: canonical name} by the settled registry; a name in no group is its own."""
    m = {}
    for g in settled["registry"]:
        for n in g["names"]:
            m[n] = g["canonical"]
    for c in index["claims"]:
        for n in c.get("constructs") or []:
            m.setdefault(n, n)
    return m


def modules(paper, settled, index, voice):
    claims, frags, _, _, core = bmt.structure(index)
    vcode = {v["fragment_id"]: v["code"] for v in voice["voice"]}
    cmap = canon_map(settled, index)
    uses = collections.defaultdict(set)
    for c in claims.values():
        for n in c.get("constructs") or []:
            uses[cmap[n]].add(c["claim_id"])
    meta = index["meta"]
    out = {}
    for c in core:
        links = [l for l in index["claim_links"] if l["claim_id"] == c]
        fr = [frags[l["fragment_id"]] for l in links]
        ev = {"passages": len({l["fragment_id"] for l in links}),
              "by_direction": dict(sorted(collections.Counter(l["direction"] for l in links).items())),
              "by_direction_and_voice": {
                  d: dict(sorted(collections.Counter(vcode.get(l["fragment_id"], "no code") for l in links
                                                     if l["direction"] == d).items()))
                  for d in sorted({l["direction"] for l in links})},
              "by_data_type": dict(sorted(collections.Counter(f["data_type"] for f in fr).items())),
              "by_unit": dict(sorted(collections.Counter(f["unit"] for f in fr).items(), key=lambda kv: (-kv[1], kv[0])))}
        timing = {"periods": meta.get("periods"), "resolution_period": meta.get("resolution_period"),
                  "resolution_basis": meta.get("resolution_basis"),
                  "reports_outcome_after_resolution": meta.get("reports_outcome_after_resolution"),
                  "passages_by_period": dict(sorted(collections.Counter(f.get("period") for f in fr).items()))}
        scope = {"units": sorted({f["unit"] for f in fr}),
                 "scope_limits": [g["gap_id"] for g in index.get("declared_gaps", [])
                                  if g.get("kind") == "scope_limit" and c in (g.get("bears_on") or [])]}
        cons = []
        for n in claims[c].get("constructs") or []:
            k = cmap[n]
            if k not in [x["canonical"] for x in cons]:
                cons.append({"canonical": k, "other_claims": sorted(uses[k] - {c}, key=num)})
        into = [e for e in settled["edges"] if e["claim"] == c]
        openings = {"ventured": claims[c].get("stated_as") == "ventured",
                    "ventured_bases": sorted({e["base"] for e in into if claims[e["base"]].get("stated_as") == "ventured"},
                                             key=num),
                    "declared_gaps": [g["gap_id"] for g in index.get("declared_gaps", []) if c in (g.get("bears_on") or [])],
                    "absences": [a["absence_id"] for a in index.get("absences", []) if c in (a.get("bears_on") or [])],
                    "qualifying": sorted({l["fragment_id"] for l in links if l["direction"] == "qualifies"}, key=num),
                    "against": sorted({l["fragment_id"] for l in links if l["direction"] == "against"}, key=num)}
        out[c] = {"statement": claims[c]["statement"], "edges_in": [
            {"base": e["base"], "type": e["type"], "strength": e["strength"]} for e in into],
            "evidence": ev, "timing": timing, "scope": scope, "constructs": cons, "openings": openings}
    follows = []
    if settled.get("units"):
        for fx in settled["factors"]:
            vals = [(u["outcome"], u["values"].get(fx["id"])) for u in settled["units"]]
            if all(o in "+-" and v == o for o, v in vals):
                follows.append(fx["id"])
    paper_level = {"ventured_claims": sorted([c for c in claims if claims[c].get("stated_as") == "ventured"], key=num),
                   "factors_following_the_outcome": follows,
                   "canonical_constructs": len(uses)}
    return {"core": out, "paper": paper_level}


def modules_md(index, mods, settled):
    L = [f"# The mechanical modules of the map: {index['meta']['citation']}", "",
         "*Made by `build_map_settle.py modules` from the frozen index, its voice file and the settled map. Counts and "
         "ids only.*", ""]
    p = mods["paper"]
    L += [f"Ventured claims: {', '.join(p['ventured_claims']) or 'none'}. Canonical constructs: "
          f"{p['canonical_constructs']}. Factors whose value follows the outcome in every unit: "
          f"{', '.join(p['factors_following_the_outcome']) or 'none'}.", ""]
    for c, m in mods["core"].items():
        L.append(f"## {c}")
        L.append("")
        L.append(m["statement"])
        L.append("")
        e = m["evidence"]
        ins = "; ".join(f"{x['base']} ({x['type']}, {x['strength']})" for x in m["edges_in"]) or "none"
        L.append(f"- **Rests on:** {ins}")
        L.append(f"- **Evidence:** {e['passages']} passages; by direction and voice: " + "; ".join(
            f"{d} {', '.join(f'{k} {n}' for k, n in v.items())}" for d, v in e["by_direction_and_voice"].items()) +
            f"; by data type: {', '.join(f'{k} {n}' for k, n in e['by_data_type'].items())}; by unit: "
            f"{', '.join(f'{k} {n}' for k, n in e['by_unit'].items())}")
        t = m["timing"]
        L.append(f"- **Timing:** periods {', '.join(t['periods'] or []) or 'none'}; resolution "
                 f"{t['resolution_period'] or 'none'}; passages by period "
                 f"{', '.join(f'{k} {n}' for k, n in t['passages_by_period'].items())}")
        s = m["scope"]
        L.append(f"- **Scope:** units {', '.join(s['units']) or 'none'}; declared scope limits "
                 f"{', '.join(s['scope_limits']) or 'none'}")
        L.append("- **Constructs:** " + "; ".join(
            f"{x['canonical']} (also {', '.join(x['other_claims']) or 'no other claim'})" for x in m["constructs"]))
        o = m["openings"]
        L.append(f"- **Openings:** {'ventured' if o['ventured'] else 'established'}; ventured bases "
                 f"{', '.join(o['ventured_bases']) or 'none'}; declared gaps {', '.join(o['declared_gaps']) or 'none'}; "
                 f"absences {', '.join(o['absences']) or 'none'}; qualifying {', '.join(o['qualifying']) or 'none'}; "
                 f"against {', '.join(o['against']) or 'none'}")
        L.append("")
    return "\n".join(L).encode("utf-8")


def modules_cmd(paper, settled_path, outdir, folder=None):
    settled = json.load(open(settled_path, encoding="utf-8"))
    if settled["paper"] != paper:
        raise SystemExit(f"{settled_path}: a settled map of {settled['paper']}, not {paper}")
    index, index_sha = bmt.load_index(paper, folder)
    vpath = os.path.join(folder or HERE, VOICE[paper])
    if folder is None and sha(vpath) != VOICE_SHA[VOICE[paper]]:
        raise SystemExit(f"{vpath}: not the frozen voice file")
    mods = modules(paper, settled, index, json.load(open(vpath, encoding="utf-8")))
    names = [f"map_modules_{paper}.json", f"MAP_MODULES_{paper}.md"]
    refuse_existing(outdir, names)
    doc = {"what": f"the mechanical modules of the settled map of {paper} (T39, part C)", "paper": paper, **mods,
           "inputs": {"settled": sha(settled_path), "index": index_sha, "voice": sha(vpath)},
           "builder_sha256": sha(os.path.abspath(__file__))}
    bmt.write_new(os.path.join(outdir, names[0]), dumps(doc))
    bmt.write_new(os.path.join(outdir, names[1]), modules_md(index, mods, settled))
    return doc


# ---------------------------------------------------------------- the network

def network_html(index, settled):
    claims, _, _, _, core = bmt.structure(index)
    cmap = canon_map(settled, index)
    uses = collections.defaultdict(set)
    for c in claims.values():
        for n in c.get("constructs") or []:
            uses[cmap[n]].add(c["claim_id"])
    aliases = collections.defaultdict(set)
    for n, k in cmap.items():
        aliases[k].add(n)
    ids = sorted(claims, key=num)
    pos = {c: i for i, c in enumerate(ids)}
    shared = [k for k in uses if len(uses[k]) >= 2]
    shared.sort(key=lambda k: (sum(pos[c] for c in uses[k]) / len(uses[k]), k.casefold(), k))
    row, top, left, right = 26, 70, 40, 520
    height = top + row * max(len(ids), len(shared), 1) + 40
    y_claim = {c: top + row * i for i, c in enumerate(ids)}
    off = (len(ids) - len(shared)) * row / 2 if len(shared) < len(ids) else 0
    y_con = {k: top + off + row * i for i, k in enumerate(shared)}
    S = [f'<svg xmlns="http://www.w3.org/2000/svg" width="980" height="{int(height)}" role="img" '
         f'aria-label="claims and the constructs they share">']
    for k in shared:
        for c in sorted(uses[k], key=num):
            S.append(f'<line x1="{left + 70}" y1="{y_claim[c] - 4}" x2="{right - 8}" y2="{y_con[k] - 4:.1f}" '
                     f'class="{"core" if c in core else "other"}"><title>{html.escape(c)}: '
                     f'{html.escape(k)}</title></line>')
    for c in ids:
        cls = "claim core" if c in core else "claim"
        S.append(f'<text x="{left}" y="{y_claim[c]}" class="{cls}"><title>{html.escape(claims[c]["statement"])}'
                 f'</title>{html.escape(c)}{" (core)" if c in core else ""}</text>')
    for k in shared:
        al = sorted(aliases[k] - {k}, key=name_key)
        tip = k + (" (also: " + "; ".join(al) + ")" if al else "")
        S.append(f'<text x="{right}" y="{y_con[k]:.1f}" class="con"><title>{html.escape(tip)}</title>'
                 f'{html.escape(k)} ({len(uses[k])})</text>')
    S.append("</svg>")
    rows = "".join(f"<tr><td>{html.escape(k)}</td><td>{html.escape('; '.join(sorted(aliases[k], key=name_key)))}</td>"
                   f"<td>{', '.join(sorted(uses[k], key=num))}</td></tr>"
                   for k in sorted(uses, key=lambda k: (-len(uses[k]), k.casefold(), k)))
    alone = [c for c in ids if not any(c in uses[k] for k in shared)]
    cite = html.escape(index["meta"]["citation"])
    return ("<!doctype html>\n<html lang=\"en\"><head><meta charset=\"utf-8\">"
            "<meta name=\"viewport\" content=\"width=device-width, initial-scale=1\">"
            f"<title>Shared constructs</title><style>"
            ":root{--bg:#ffffff;--fg:#1d1d1f;--mute:#6e6e73;--core:#b5462c;--line:#9aa5b1;--rule:#e3e3e8}"
            "@media (prefers-color-scheme: dark){:root{--bg:#161618;--fg:#ececf0;--mute:#a1a1aa;--core:#ff8a6b;"
            "--line:#5b6673;--rule:#2c2c31}}"
            "body{background:var(--bg);color:var(--fg);font:14px/1.45 -apple-system,Segoe UI,Helvetica,Arial,"
            "sans-serif;margin:0 auto;max-width:1000px;padding:16px}"
            "h1{font-size:18px;margin:8px 0}p{color:var(--mute);margin:4px 0 12px}"
            ".wrap{overflow-x:auto}svg text{fill:var(--fg);font-size:12px}svg text.core{fill:var(--core);"
            "font-weight:600}svg line{stroke:var(--line);stroke-width:1}svg line.core{stroke:var(--core);"
            "stroke-opacity:.6}table{border-collapse:collapse;width:100%;margin-top:16px}td,th{border-top:1px solid "
            "var(--rule);padding:4px 6px;text-align:left;vertical-align:top}th{color:var(--mute);font-weight:600}"
            "</style></head><body>"
            f"<h1>Shared constructs: {cite}</h1>"
            "<p>Each line joins a claim (left) to a construct it names (right), for the constructs two or more claims "
            "name, after the settled map's merges of names. <strong>Shared constructs, not dependence:</strong> a "
            "line says two claims name the same construct, not that one rests on the other. Core claims in colour; "
            "hover a claim for its statement, a construct for its other names.</p>"
            f"<div class=\"wrap\">{''.join(S)}</div>"
            f"<p>Claims that share no construct: {', '.join(alone) or 'none'}.</p>"
            "<table><thead><tr><th>construct</th><th>names in the index</th><th>claims</th></tr></thead>"
            f"<tbody>{rows}</tbody></table>"
            "<p>Made by build_map_settle.py network from the frozen index and the settled map's registry.</p>"
            "</body></html>\n").encode("utf-8")


def network_cmd(paper, settled_path, outdir, folder=None):
    settled = json.load(open(settled_path, encoding="utf-8"))
    if settled["paper"] != paper:
        raise SystemExit(f"{settled_path}: a settled map of {settled['paper']}, not {paper}")
    index, _ = bmt.load_index(paper, folder)
    name = f"MAP_NETWORK_{paper}.html"
    refuse_existing(outdir, [name])
    bmt.write_new(os.path.join(outdir, name), network_html(index, settled))


# ---------------------------------------------------------------- self-test

def selftest():
    import contextlib
    import io
    import shutil
    import tempfile
    n_ok = n_bad = 0

    def expect(label, cond):
        nonlocal n_ok, n_bad
        if cond:
            n_ok += 1
            print(f"ok    {label}")
        else:
            n_bad += 1
            print(f"FAIL  {label}")

    tmp = tempfile.mkdtemp()
    quiet = contextlib.redirect_stdout(io.StringIO())
    try:
        src = os.path.join(tmp, "src")
        os.makedirs(src)
        for paper in PAPERS:
            with open(os.path.join(src, bmt.INDEX[paper]), "w") as f:
                json.dump(bmt.made_up_index(paper), f)
            frs = bmt.made_up_index(paper)["fragments"]
            with open(os.path.join(src, VOICE[paper]), "w") as f:
                json.dump({"paper_id": paper, "voice": [{"fragment_id": x["fragment_id"],
                                                         "code": "author_insight" if x["fragment_id"] in ("F01", "F02")
                                                         else "reported", "reason": "r"} for x in frs],
                           "gloss_relation": []}, f)
        kdir = os.path.join(tmp, "k")
        with quiet:
            bmt.build(kdir, 7, folder=src, sample={"bechky": 3, "pratt": 4})
        key = json.load(open(os.path.join(kdir, "map_key.json")))
        order = key["papers"]["bechky"]["order"]
        pid = {(o["base"], o["claim"]): o["id"] for o in order}
        d = bmt.made_up_index("bechky")
        listed = set(pid)
        free = [(b, c) for b, c in bmt.all_pairs(d) if (b, c) not in listed]

        def reply(edges, further, registry, factors, units, contrasts):
            pairs = []
            for o in order:
                e = edges.get((o["base"], o["claim"]))
                p = {"pair": o["id"], "claim": o["claim"], "base": o["base"], "rests_on": bool(e), "reason": "r."}
                if e:
                    p.update(type=e[0], strength=e[1], basis="F01", page="p. 12")
                pairs.append(p)
            fe = [{"base": b, "claim": c, "type": t, "strength": s, "basis": "b", "page": "p. 3", "reason": "r."}
                  for (b, c), (t, s) in further.items()]
            return {"paper": "Bechky (made up)", "pairs": pairs, "further_edges": fe, "registry": registry,
                    "factors": factors, "units": units, "contrasts": contrasts}

        e_both = ("C02", "C01")
        e_one = ("C03", "C01") if ("C03", "C01") in listed else None
        listed_list = sorted(listed, key=lambda p: (num(p[0]), num(p[1])))
        other = [p for p in listed_list if p != e_both][:2]
        f1, f2 = free[0], free[1]
        ea = {e_both: ("evidential", "contributing"), other[0]: ("mechanistic", "necessary"),
              other[1]: ("comparative", "necessary")}
        eb = {e_both: ("evidential", "contributing"), other[0]: ("evidential", "necessary")}
        fa_ = {f1: ("significance", "contributing"), f2: ("evidential", "necessary")}
        fb_ = {f2: ("evidential", "necessary")}
        ra = reply(ea, fa_, [{"canonical": "endorsement", "names": ["endorsement", "adoption"]},
                             {"canonical": "budget", "names": ["budget", "identity"]}],
                   [{"id": "X1", "name": "endorsement", "claim": "C01"}, {"id": "X2", "name": "budget", "claim": "C04"}],
                   [{"unit": "ward A", "outcome": "+", "values": {"X1": "+", "X2": "?"}, "ids": ["C02"]},
                    {"unit": "ward B", "outcome": "-", "values": {"X1": "-", "X2": "+"}, "ids": ["C03"]}],
                   [{"contrast": "CP01", "controls": {"X1": "varies", "X2": "held"}, "ids": ["F01"]}])
        rb = reply(eb, fb_, [{"canonical": "adoption", "names": ["adoption", "endorsement"]},
                             {"canonical": "identity", "names": ["identity", "ward A"]}],
                   [{"id": "K1", "name": "senior endorsement", "claim": "C01"},
                    {"id": "K2", "name": "identity", "claim": "C05"}],
                   [{"unit": "ward A", "outcome": "+", "values": {"K1": "+", "K2": "+"}, "ids": ["C02"]},
                    {"unit": "ward B", "outcome": "+", "values": {"K1": "?", "K2": "-"}, "ids": ["C03"]}],
                   [{"contrast": "CP01", "controls": {"K1": "unclear", "K2": "held"}, "ids": ["F07"]}])
        pa, pb = os.path.join(tmp, "ra.json"), os.path.join(tmp, "rb.json")
        json.dump(ra, open(pa, "w"))
        json.dump(rb, open(pb, "w"))
        kpath = os.path.join(kdir, "map_key.json")
        with quiet:
            expect("the made-up replies count by the map test's check",
                   bmt.check_map("bechky", kpath, pa, folder=src) == 0 and bmt.check_map("bechky", kpath, pb, folder=src) == 0)
        items, agreed = compare("bechky", key, d, ra, rb)
        kinds = collections.Counter(it["kind"] for it in items)
        expect("an edge in both alike is agreed, not an item",
               any((e["base"], e["claim"]) == e_both for e in agreed["edges"])
               and not any(it.get("base") == e_both[0] and it.get("claim") == e_both[1] for it in items))
        expect("another type on a listed pair: one item", kinds["pair-type"] == 1)
        expect("an edge in one map only on a listed pair: one item", kinds["pair-edge"] == 1)
        expect("further edges: one in A only an item, one in both alike agreed",
               kinds["further-one"] == 1 and any((e["base"], e["claim"]) == f2 and e["source"] == "further"
                                                 for e in agreed["edges"]))
        cons = sorted(tuple(it["names"]) for it in items if it["kind"] == "constructs")
        expect("registries: the name pairs one map groups and the other not: budget-identity (A), identity-ward A (B)",
               kinds["constructs"] == 2 and cons == [("budget", "identity"), ("identity", "ward A")])
        expect("factors: C01 paired, C04 in A only, C05 in B only",
               [(fx["a"] or {}).get("id") for fx in agreed["factors"]] == ["X1", "X2", None]
               and kinds["factor-one"] == 2)
        expect("units: ward B's outcome differs; values differ for the paired factor in ward B only",
               kinds["unit-outcome"] == 1 and kinds["unit-value"] == 1)
        expect("contrasts: the paired factor's control differs", kinds["contrast"] == 1)
        expect("items numbered S001 onward, pairs first",
               [it["id"] for it in items] == ["S%03d" % i for i in range(1, len(items) + 1)]
               and items[0]["kind"] in ("pair-edge", "pair-type"))
        out1, out2 = os.path.join(tmp, "s1"), os.path.join(tmp, "s2")
        with quiet:
            sk1 = sheet_cmd("bechky", kpath, pa, pb, out1, folder=src)
            sk2 = sheet_cmd("bechky", kpath, pa, pb, out2, folder=src)
        b1 = open(os.path.join(out1, "MAP_SETTLE_bechky.xlsx"), "rb").read()
        b2 = open(os.path.join(out2, "MAP_SETTLE_bechky.xlsx"), "rb").read()
        expect("two builds of the sheet give the same bytes", b1 == b2 and sk1 == sk2)
        try:
            with quiet:
                sheet_cmd("bechky", kpath, pa, pb, out1, folder=src)
            expect("refuses to write over a file", False)
        except SystemExit:
            expect("refuses to write over a file", True)
        bad = dict(ra)
        bad["pairs"] = ra["pairs"][1:]
        pbad = os.path.join(tmp, "bad.json")
        json.dump(bad, open(pbad, "w"))
        try:
            with quiet:
                sheet_cmd("bechky", kpath, pbad, pb, os.path.join(tmp, "s3"), folder=src)
            expect("refuses a reply that does not count", False)
        except SystemExit:
            expect("refuses a reply that does not count", True)
        span = json.loads(json.dumps(ra))
        span["pairs"][0]["reason"] = 'It says "' + " ".join(["word"] * 10) + '".'
        pspan = os.path.join(tmp, "span.json")
        json.dump(span, open(pspan, "w"))
        try:
            with quiet:
                sheet_cmd("bechky", kpath, pspan, pb, os.path.join(tmp, "s4"), folder=src)
            expect("refuses a reply with a long quoted span", False)
        except SystemExit:
            expect("refuses a reply with a long quoted span", True)
        wb = load_workbook(os.path.join(out1, "MAP_SETTLE_bechky.xlsx"))
        expect("the sheet's tabs", wb.sheetnames == ["Instructions", "Pairs", "Constructs", "Factors", "Claims",
                                                     "Groups", "Map factors"])
        allcells = " ".join(str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value is not None)
        expect("the sheet names no model and no chat", not re.search(r"opus|gpt|claude|chatgpt|\b7[3-6]\b", allcells,
                                                                       re.I))
        sk = json.load(open(os.path.join(out1, "map_settle_key_bechky.json")))

        def fill(path, choices, start="10:00", end="10:40", tamper=None):
            wbk = load_workbook(os.path.join(out1, "MAP_SETTLE_bechky.xlsx"))
            wbk["Instructions"].cell(TIME_ROWS[0], 2).value = start
            wbk["Instructions"].cell(TIME_ROWS[1], 2).value = end
            rows = collections.Counter()
            for it in sk["items"]:
                tab = tab_of(it["kind"])
                rows[tab] += 1
                r = rows[tab] + 1
                for col, v in choices.get(it["id"], {}).items():
                    wbk[tab].cell(r, col).value = v
            if tamper:
                tamper(wbk)
            wbk.save(path)

        ids = {it["kind"]: [] for it in sk["items"]}
        for it in sk["items"]:
            ids[it["kind"]].append(it)
        ch = {}
        for it in sk["items"]:
            if it["kind"] == "pair-edge":
                ch[it["id"]] = {9: "own", 10: "significance", 11: "contributing", 12: "the PI's note"}
            elif it["kind"] == "pair-type":
                ch[it["id"]] = {9: "B"}
            elif it["kind"] in ("further-one", "further-type"):
                ch[it["id"]] = {9: "A"}
            elif it["kind"] == "constructs":
                ch[it["id"]] = {6: "A"}
            elif it["kind"] == "factor-one":
                ch[it["id"]] = {6: "A"}
            elif it["kind"] == "unit-outcome":
                ch[it["id"]] = {6: "own", 7: "does not"}
            elif it["kind"] == "unit-value":
                ch[it["id"]] = {6: "B"}
            elif it["kind"] == "contrast":
                ch[it["id"]] = {6: "own", 7: "unclear"}
        fpath = os.path.join(tmp, "filled.xlsx")
        fill(fpath, ch)
        sdir = os.path.join(tmp, "settled")
        with quiet:
            st = settle_cmd("bechky", os.path.join(out1, "map_settle_key_bechky.json"), fpath, sdir, folder=src)
        E = {(e["base"], e["claim"]): (e["type"], e["strength"], e["by"]) for e in st["edges"]}
        pe = next(it for it in sk["items"] if it["kind"] == "pair-edge")
        pt = next(it for it in sk["items"] if it["kind"] == "pair-type")
        expect("settled: the agreed edge stands", E.get(e_both) == ("evidential", "contributing", "both"))
        expect("settled: the PI's own answer on an edge in one map", E.get((pe["base"], pe["claim"])) ==
               ("significance", "contributing", "own"))
        expect("settled: map B's type where the PI chose B", E.get((pt["base"], pt["claim"])) ==
               ("evidential", "necessary", "B"))
        expect("settled: a further edge in A only, kept as A", E.get(f1) == ("significance", "contributing", "A"))
        expect("settled: a further edge in both alike", E.get(f2) == ("evidential", "necessary", "both"))
        reg = {tuple(g["names"]): g["canonical"] for g in st["registry"]}
        expect("settled registry: adoption-endorsement (both), budget-identity (A, chosen); not identity-ward A",
               reg == {("adoption", "endorsement"): "endorsement", ("budget", "identity"): "budget"})
        expect("settled factors: the paired, A's kept, B's dropped (A chosen)",
               [(f["id"], f["claim"], f["by"]) for f in st["factors"]] == [("X1", "C01", "both"), ("X2", "C04", "A")])
        U = {u["unit"]: u for u in st["units"]}
        expect("settled units: the PI's own outcome for ward B; B's value where chosen; A's value of A's own factor",
               U["ward B"]["outcome"] == "-" and U["ward B"]["values"] == {"X1": "?", "X2": "+"}
               and U["ward A"]["values"] == {"X1": "+", "X2": "?"} and U["ward A"]["outcome"] == "+")
        expect("settled contrasts: the PI's own control, and A's own factor's",
               st["contrasts"][0]["controls"] == {"X1": "unclear", "X2": "held"})
        expect("settled: the PI's times and the PI's note kept", st["times"] == {"start": "10:00", "end": "10:40"}
               and st["choices"][pe["id"]]["note"] == "the PI's note")
        try:
            fill(os.path.join(tmp, "t1.xlsx"), ch, tamper=lambda w: setattr(w["Claims"].cell(2, 4), "value", "x"))
            with quiet:
                settle_cmd("bechky", os.path.join(out1, "map_settle_key_bechky.json"), os.path.join(tmp, "t1.xlsx"),
                           os.path.join(tmp, "x1"), folder=src)
            expect("refuses a sheet changed outside the yellow cells", False)
        except SystemExit:
            expect("refuses a sheet changed outside the yellow cells", True)
        ch2 = dict(ch)
        ch2.pop(pt["id"])
        try:
            fill(os.path.join(tmp, "t2.xlsx"), ch2)
            with quiet:
                settle_cmd("bechky", os.path.join(out1, "map_settle_key_bechky.json"), os.path.join(tmp, "t2.xlsx"),
                           os.path.join(tmp, "x2"), folder=src)
            expect("refuses an item without a choice", False)
        except SystemExit:
            expect("refuses an item without a choice", True)
        ch3 = dict(ch)
        ch3[pe["id"]] = {9: "own", 10: "evidential", 11: None}
        try:
            fill(os.path.join(tmp, "t3.xlsx"), ch3)
            with quiet:
                settle_cmd("bechky", os.path.join(out1, "map_settle_key_bechky.json"), os.path.join(tmp, "t3.xlsx"),
                           os.path.join(tmp, "x3"), folder=src)
            expect("refuses an own edge without a strength", False)
        except SystemExit:
            expect("refuses an own edge without a strength", True)
        ch4 = dict(ch)
        uo = next(it for it in sk["items"] if it["kind"] == "unit-outcome")
        ch4[uo["id"]] = {6: "own", 7: "held"}
        try:
            fill(os.path.join(tmp, "t4.xlsx"), ch4)
            with quiet:
                settle_cmd("bechky", os.path.join(out1, "map_settle_key_bechky.json"), os.path.join(tmp, "t4.xlsx"),
                           os.path.join(tmp, "x4"), folder=src)
            expect("refuses an own value of the wrong kind", False)
        except SystemExit:
            expect("refuses an own value of the wrong kind", True)
        # the closure
        E1 = [{"base": "C02", "claim": "C01", "type": "evidential", "strength": "contributing"},
              {"base": "C03", "claim": "C01", "type": "evidential", "strength": "contributing"},
              {"base": "C01", "claim": "C04", "type": "mechanistic", "strength": "necessary"},
              {"base": "C04", "claim": "C05", "type": "comparative", "strength": "necessary"},
              {"base": "C06", "claim": "C05", "type": "significance", "strength": "contributing"},
              {"base": "C06", "claim": "C04", "type": "evidential", "strength": "necessary"}]
        expect("closure: one of two contributing evidential bases hit: weakened",
               labels_of(E1, ["C02"]) == {"C02": "contested", "C01": "weakened"})
        expect("closure: both hit: unsupported, and on through a necessary mechanistic edge, then comparative",
               labels_of(E1, ["C02", "C03"]) == {"C02": "contested", "C03": "contested", "C01": "unsupported",
                                                 "C04": "unsupported", "C05": "restated"})
        expect("closure: a weakened claim carries nothing", labels_of(E1, ["C02"]).get("C04") is None)
        expect("closure: a necessary evidential edge: unsupported; significance: idle",
               labels_of(E1, ["C06"]) == {"C06": "contested", "C04": "unsupported", "C05": "restated"})
        expect("closure: the strongest label kept (contested over unsupported)",
               labels_of(E1, ["C02", "C03", "C04"])["C04"] == "contested")
        E4 = [{"base": "C02", "claim": "C01", "type": "mechanistic", "strength": "contributing"}]
        expect("closure: a contributing mechanistic edge: weakened", labels_of(E4, ["C02"])["C01"] == "weakened")
        E2 = [{"base": "C06", "claim": "C05", "type": "significance", "strength": "contributing"}]
        expect("closure: significance alone gives idle", labels_of(E2, ["C06"]) == {"C06": "contested", "C05": "idle"})
        E3 = [{"base": "C02", "claim": "C01", "type": "evidential", "strength": "contributing"},
              {"base": "C03", "claim": "C01", "type": "evidential", "strength": "contributing"},
              {"base": "C06", "claim": "C03", "type": "evidential", "strength": "necessary"}]
        expect("closure: weakened becomes unsupported once its other base is unsupported",
               labels_of(E3, ["C02", "C06"])["C01"] == "unsupported")
        res = {"codes": {"gemini": {"bechky|R01": {"main_claim": "C02", "main_type": "other_mechanism",
                                                   "other_claims": [["C03", "rereads_evidence"]]},
                                    "bechky|R02": {"main_claim": "C01", "main_type": "complements", "other_claims": []}},
                         "astra": {"bechky|R01": {"main_claim": "C02", "main_type": "inverts", "other_claims": []},
                                   "bechky|R02": {"main_claim": "C01", "main_type": "restates", "other_claims": []}}}}
        rpath = os.path.join(tmp, "results.json")
        json.dump(res, open(rpath, "w"))
        spath = os.path.join(tmp, "settled_e1.json")
        json.dump({"paper": "bechky", "edges": E1, "registry": [], "factors": [], "units": [], "contrasts": []},
                  open(spath, "w"))
        with quiet:
            cl = closure_cmd("bechky", spath, rpath, os.path.join(tmp, "cl"), folder=src, results_sha=None)
        r1 = cl["readings"][0]
        expect("closure file: effects on core claims by coder",
               r1["reading"] == "R01" and {(e["claim"], e["label"], tuple(e["coders"])) for e in r1["core_effects"]} ==
               {("C01", "unsupported", ("gemini",)), ("C04", "unsupported", ("gemini",)), ("C05", "restated", ("gemini",)),
                ("C01", "weakened", ("astra",))})
        expect("closure file: complements and restates hit nothing", cl["readings"][1]["core_effects"] == [])
        try:
            with quiet:
                closure_cmd("bechky", spath, rpath, os.path.join(tmp, "cl2"), folder=src)
            expect("closure refuses results that are not T37's", False)
        except SystemExit:
            expect("closure refuses results that are not T37's", True)
        # the modules and the network
        with quiet:
            mo = modules_cmd("bechky", os.path.join(sdir, "map_settled_bechky.json"), os.path.join(tmp, "mo"),
                             folder=src)
        m1 = mo["core"]["C01"]
        expect("modules: C01's evidence by direction and voice",
               m1["evidence"]["passages"] == 5 and m1["evidence"]["by_direction_and_voice"] ==
               {"illustrates": {"author_insight": 2, "reported": 3}})
        expect("modules: C01's units and the scope limit bearing on it",
               m1["scope"]["units"] == ["ward A", "ward B"] and m1["scope"]["scope_limits"] == ["DG01"])
        expect("modules: constructs by the settled registry",
               any(x["canonical"] == "endorsement" and x["other_claims"] == ["C02", "C03", "C04", "C05"]
                   for x in m1["constructs"]))
        expect("modules: no factor follows the outcome where a value is not said (X1 in ward B)",
               mo["paper"]["factors_following_the_outcome"] == [])
        fake = {"edges": [], "registry": [], "factors": [{"id": "X1"}, {"id": "X2"}, {"id": "X3"}],
                "units": [{"outcome": "+", "values": {"X1": "+", "X2": "+", "X3": "?"}},
                          {"outcome": "-", "values": {"X1": "-", "X2": "+", "X3": "-"}}]}
        mf = modules("bechky", fake, d, {"voice": []})
        fake2 = {"edges": [], "registry": [], "factors": [{"id": "X1"}],
                 "units": [{"outcome": "?", "values": {"X1": "?"}}, {"outcome": "?", "values": {"X1": "?"}}]}
        expect("modules: an outcome and a value both not said do not make a factor follow",
               modules("bechky", fake2, d, {"voice": []})["paper"]["factors_following_the_outcome"] == [])
        expect("modules: a factor whose value equals the outcome in every unit follows it; one off or unknown does not",
               mf["paper"]["factors_following_the_outcome"] == ["X1"])
        md = open(os.path.join(tmp, "mo", "MAP_MODULES_bechky.md"), encoding="utf-8").read()
        expect("modules: the page has a section per core claim", all(f"## {c}" in md for c in ("C01", "C04", "C05")))
        with quiet:
            network_cmd("bechky", os.path.join(sdir, "map_settled_bechky.json"), os.path.join(tmp, "nw"), folder=src)
        page = open(os.path.join(tmp, "nw", "MAP_NETWORK_bechky.html"), encoding="utf-8").read()
        expect("network: labelled shared constructs, not dependence, with the merged construct",
               "Shared constructs, not dependence" in page and ">endorsement (5)<" in page)
        expect("network: one page, no script, no outside resource",
               "<script" not in page and "http" not in page.replace('xmlns="http://www.w3.org/2000/svg"', ""))
        with quiet:
            network_cmd("bechky", os.path.join(sdir, "map_settled_bechky.json"), os.path.join(tmp, "nw2"), folder=src)
        expect("network: the same bytes twice", page == open(os.path.join(tmp, "nw2", "MAP_NETWORK_bechky.html"),
                                                             encoding="utf-8").read())
    finally:
        shutil.rmtree(tmp)
    print(f"{n_ok} checks passed, {n_bad} failed")
    print("SELFTEST PASSES" if n_bad == 0 else "SELFTEST FAILS")
    return 0 if n_bad == 0 else 1


def main(argv):
    a = argv[1:]
    if a == ["--selftest"]:
        return selftest()
    if len(a) == 6 and a[0] == "sheet" and a[1] in PAPERS:
        sheet_cmd(a[1], a[2], a[3], a[4], a[5])
        return 0
    if len(a) == 5 and a[0] == "settle" and a[1] in PAPERS:
        settle_cmd(a[1], a[2], a[3], a[4])
        return 0
    if len(a) == 5 and a[0] == "closure" and a[1] in PAPERS:
        closure_cmd(a[1], a[2], a[3], a[4])
        return 0
    if len(a) == 4 and a[0] == "modules" and a[1] in PAPERS:
        modules_cmd(a[1], a[2], a[3])
        return 0
    if len(a) == 4 and a[0] == "network" and a[1] in PAPERS:
        network_cmd(a[1], a[2], a[3])
        return 0
    print(__doc__)
    return 2


if __name__ == "__main__":
    sys.exit(main(sys.argv))
